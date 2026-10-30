"""Output-tool invariants without executing an input repository or launching Word."""

import importlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
from types import SimpleNamespace

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
exporter = importlib.import_module("export_manuscript")
converter = importlib.import_module("_publication_convert")
templates = importlib.import_module("_publication_templates")
checks = importlib.import_module("_publication_checks")
author_facts = importlib.import_module("_author_profile")

ARTICLE = """# Shared configuration cleanup

## Abstract

A conditional static example explains a maintenance decision.

Keywords: maintenance; configuration; ownership

## A practical decision

A shared requirement may survive the command that first wrote it. [1]

## Actionable insights

- Record the continuing consumer.
- Preserve the original setting.
- Refuse uncertain restoration.

## Acknowledgments

OpenAI Codex assisted with the analysis and prose throughout this article.

## References

[1] Example Authors, A retained source, 2025. https://example.org/source
"""


@pytest.fixture
def manuscript(tmp_path):
    source = tmp_path / "article.md"
    source.write_text(ARTICLE, encoding="utf-8")
    return source


def test_preflight_uses_explicit_facts_and_preserves_input(manuscript, monkeypatch):
    monkeypatch.setenv("PF_AUTHOR_PROFILE_JSON", '{"display_name":"Ambient person"}')
    before = manuscript.read_bytes()
    result = exporter.preflight(manuscript, "ieee-software")
    assert result["author_metadata"] == {}
    assert result["prefilter_passed"]
    assert not result["submission_ready"]
    assert not result["native_word_measured"]
    assert manuscript.read_bytes() == before
    assert list(manuscript.parent.iterdir()) == [manuscript]


@pytest.mark.parametrize("separator", ["\n", "\n\n"])
def test_preflight_counts_multiline_references_with_supported_separators(manuscript, separator):
    references = [
        "[1] First Authors,\n  A retained first title, 2024. https://example.org/first",
        "[2] Second Authors,\n  A retained second title, 2025. https://example.org/second",
    ]
    manuscript.write_text(ARTICLE.split("## References", 1)[0] + "## References\n\n" +
                          separator.join(references) + "\n", encoding="utf-8")
    original = manuscript.read_bytes()
    result = exporter.preflight(manuscript, "ieee-software")
    statuses = {row["id"]: row["status"] for row in result["items"]}
    assert result["metrics"]["reference_count"] == 2
    assert statuses["reference_count"] == "passed"
    assert statuses["reference_metadata"] == "passed"
    assert result["prefilter_passed"]
    assert not result["native_word_measured"]
    assert manuscript.read_bytes() == original


@pytest.mark.parametrize("case, count_status, metadata_status", [
    ("duplicate", "failed", "passed"),
    ("gap", "failed", "passed"),
    ("over_limit", "failed", "passed"),
    ("raw_url", "passed", "failed"),
])
def test_preflight_rejects_invalid_consecutive_references(manuscript, case, count_status, metadata_status):
    numbers = {
        "duplicate": [1, 1, 2, 3],
        "gap": [1, 3],
        "over_limit": list(range(1, checks.VENUE_LIMITS["ieee-software"]["max_references"] + 2)),
        "raw_url": [1, 2],
    }[case]
    references = [f"[{number}] Example Authors, Retained source {index}, 2025."
                  for index, number in enumerate(numbers, 1)]
    if case == "raw_url":
        references[1] = "[2] https://example.org/second"
    manuscript.write_text(ARTICLE.split("## References", 1)[0] + "## References\n\n" +
                          "\n".join(references) + "\n", encoding="utf-8")
    original = manuscript.read_bytes()
    result = exporter.preflight(manuscript, "ieee-software")
    statuses = {row["id"]: row["status"] for row in result["items"]}
    assert result["metrics"]["reference_count"] == len(numbers)
    assert statuses["reference_count"] == count_status
    assert statuses["reference_metadata"] == metadata_status
    assert not result["prefilter_passed"]
    assert not result["native_word_measured"]
    assert manuscript.read_bytes() == original


@pytest.mark.parametrize("citation", ["[1]", "[1, 3]", "[1-4]", "[1; 3]", "[1, 3–4]"])
def test_preflight_rejects_supported_numeric_citations_in_abstract(manuscript, citation):
    text = ARTICLE.replace("A conditional static example explains a maintenance decision.",
                           "A conditional static example explains a maintenance decision " + citation + ".", 1)
    manuscript.write_text(text, encoding="utf-8")
    original = manuscript.read_bytes()
    result = exporter.preflight(manuscript, "ieee-software")
    statuses = {row["id"]: row["status"] for row in result["items"]}
    assert statuses["abstract_structure"] == "failed"
    assert not result["prefilter_passed"]
    assert not result["native_word_measured"]
    assert manuscript.read_bytes() == original


def test_preflight_body_grouped_citation_does_not_fail_abstract_structure(manuscript):
    text = ARTICLE.replace("A shared requirement may survive the command that first wrote it. [1]",
                           "A shared requirement may survive the command that first wrote it. [1, 3–4]", 1)
    manuscript.write_text(text, encoding="utf-8")
    original = manuscript.read_bytes()
    result = exporter.preflight(manuscript, "ieee-software")
    statuses = {row["id"]: row["status"] for row in result["items"]}
    assert statuses["abstract_structure"] == "passed"
    assert result["prefilter_passed"]
    assert not result["native_word_measured"]
    assert manuscript.read_bytes() == original


def test_author_input_requires_real_typed_fields(tmp_path):
    assert author_facts._explicit_author({"given_name": "A", "family_name": "B"})["display_name"] == "A B"
    assert author_facts._explicit_author({"orcid": "0009-0009-8671-9967"})["orcid"] == "0009-0009-8671-9967"
    for value in [[], {"email": "invalid"}, {"orcid": "0009-0009-8671-9968"},
                  {"approved": True}, {"display_name": 123}, {"display_name": "A\u2028"},
                  {"biography_details": {"role": 4}}]:
        with pytest.raises(ValueError):
            author_facts._explicit_author(value)


@pytest.mark.parametrize("value", [
    {"display_name": "A\x00B"},
    {"given_name": "A\vB"},
    {"affiliation": "Lab\nTown"},
    {"email": "person@example.org\t"},
    {"biography": "First\fSecond"},
    {"biography_details": {"role": "A\x00B"}},
    {"biography_details": {"role": "A\vB"}},
    {"biography_details": {"role": "A\x7fB"}},
    {"biography_details": {"role\x00": "Researcher"}},
])
def test_author_control_characters_are_rejected_before_trimming(value):
    with pytest.raises(ValueError, match="control character"):
        author_facts._explicit_author(value)


def test_biography_allows_only_normal_multiline_control_characters():
    result = author_facts._explicit_author({
        "given_name": " A ", "family_name": " B ",
        "biography": " First line\nSecond line\r\nThird\tpart ",
        "biography_details": {"role": " First\nSecond\tpart "},
    })
    assert result["display_name"] == "A B"
    assert result["biography"] == "First line\nSecond line\r\nThird\tpart"
    assert result["biography_details"]["role"] == "First\nSecond\tpart"


def test_same_multi_author_profile_reaches_preflight_and_export_without_losing_coauthors(manuscript, tmp_path, monkeypatch):
    import submission_support
    profile = tmp_path / "profile.json"
    alice = {"display_name": "Alice Research", "affiliation": "Synthetic Lab A", "email": "alice@example.org",
             "orcid": "0009-0009-8671-9967", "biography": "Alice studies software maintenance.",
             "biography_details": {key: "Synthetic supplied fact for " + key for key in ("role", "location", "research_interests", "highest_degree", "degree_institution", "professional_membership")}}
    profile.write_text(json.dumps({"authors": [{"id": "alice", "author": alice},
                                               {"id": "bob", "author": {"display_name": "Bob Scientist"}}],
                                   "order": ["bob", "alice"], "corresponding_author": "alice", "funding": None, "conflicts": None}), encoding="utf-8")
    prepared = tmp_path / "prepared"
    result = submission_support.prepare_authors(profile, manuscript, prepared)
    assert "bob: biography_details.highest_degree" in result["missing_supplied_fields"]
    check = exporter.preflight(prepared / "manuscript.md", "ieee-software", profile=prepared / "authors.json")
    assert check["author_metadata"]["order"] == ["bob", "alice"]
    bob, alice = check["metrics"]["authors"]
    assert bob["id"] == "bob" and "affiliation" in bob["missing_fields"]
    assert "biography_details.highest_degree" in bob["missing_fields"]
    assert alice["id"] == "alice" and alice["supplied_fields_complete"]
    assert alice["supplied_biography_in_manuscript"]
    assert not bob["missing_byline_values"] and not alice["missing_byline_values"]
    assert not check["submission_ready"]
    _mock_native_export(monkeypatch)
    original_report = exporter.compliance_report
    def report(*args, **kwargs):
        assert kwargs["author_metadata"]["order"] == ["bob", "alice"]
        return original_report(*args, **kwargs)
    monkeypatch.setattr(exporter, "compliance_report", report)
    output = tmp_path / "export"
    exporter.export(prepared / "manuscript.md", "ieee-software", output, profile=prepared / "authors.json")
    assert (output / "authors.json").read_bytes() == (prepared / "authors.json").read_bytes()
    provenance = json.loads((output / "provenance.json").read_text())
    assert [entry["id"] for entry in provenance["author_metadata"]["authors"]] == ["bob", "alice"]


def test_coauthor_metadata_does_not_hide_a_missing_rendered_byline(manuscript, tmp_path):
    profile = tmp_path / "profile.json"
    profile.write_text(json.dumps({"authors": [{"id": "bob", "author": {"display_name": "Bob Scientist", "email": "bob@example.org"}}],
                                   "order": ["bob"]}), encoding="utf-8")
    check = exporter.preflight(manuscript, "ieee-software", profile=profile)
    row = check["metrics"]["authors"][0]
    assert set(row["missing_byline_values"]) == {"display_name", "email"}
    assert next(item for item in check["items"] if item["id"] == "author_profile_binding")["status"] == "pending"


def _real_pandoc():
    try:
        binary = converter.pandoc_binary()
        converter.pandoc_version(binary)
        return binary
    except ValueError as exc:
        if os.environ.get("CI", "").casefold() == "true":
            pytest.fail("CI requires the actual Pandoc conversion regressions: " + str(exc))
        pytest.skip("An installed Pandoc is required for the actual native conversion regression")


def test_ci_cannot_skip_a_broken_pandoc_installation(monkeypatch):
    monkeypatch.setenv("CI", "true")
    def unavailable():
        raise ValueError("Broken discovery")
    monkeypatch.setattr(converter, "pandoc_binary", unavailable)
    with pytest.raises(pytest.fail.Exception, match="CI requires"):
        _real_pandoc()


def test_conversion_rejects_a_python_executable_without_creating_docx(manuscript, tmp_path):
    target = tmp_path / "wrong.docx"
    with pytest.raises(ValueError, match="identify itself as Pandoc"):
        converter.convert_docx(manuscript, target, venue_id="ieee-software", pandoc=sys.executable)
    assert not target.exists()


def test_prepared_author_manuscript_exports_real_docx_with_explicit_resource_dir(manuscript, tmp_path, monkeypatch, capsys):
    import submission_support
    import zipfile
    from docx import Document
    from PIL import Image
    binary = _real_pandoc()
    manuscript.write_text(ARTICLE.replace("## References", "![A relative image](figure.png)\n\n## References", 1),
                          encoding="utf-8")
    figure = tmp_path / "figure.png"
    Image.new("RGB", (12, 12), "navy").save(figure)
    profile = tmp_path / "profile.json"
    profile.write_text(json.dumps({"authors": [{"id": "alice", "author": {
        "display_name": "Alice Research", "affiliation": "Lab A", "email": "alice@example.org",
        "biography": "Studies software."}, "contributions": ["Writing"]}], "order": ["alice"]}), encoding="utf-8")
    original = {path: path.read_bytes() for path in (manuscript, profile, figure)}
    prepared = tmp_path / "prepared"
    submission_support.prepare_authors(profile, manuscript, prepared)
    prepared_manuscript, prepared_profile = prepared / "manuscript.md", prepared / "authors.json"
    original.update({path: path.read_bytes() for path in (prepared_manuscript, prepared_profile)})
    assert not (prepared / "figure.png").exists()
    sentinel = "Native Word intentionally not run"
    word_calls = []
    def stop_before_word(docx, pdf):
        word_calls.append((docx, pdf))
        raise ValueError(sentinel)
    monkeypatch.setattr(exporter, "word_pdf", stop_before_word)
    output = tmp_path / "export"
    result = exporter.main(["export", "--manuscript", str(prepared_manuscript), "--profile", str(prepared_profile),
                            "--venue", "ieee-software", "--output-dir", str(output), "--pandoc", str(binary),
                            "--resource-dir", str(tmp_path)])
    assert result == 2
    assert json.loads(capsys.readouterr().err)["error"] == sentinel
    docx, pdf = output / "manuscript-ieee-software.docx", output / "manuscript-ieee-software.pdf"
    assert word_calls == [(docx, pdf)]
    document = Document(docx)
    assert len(document.inline_shapes) == 1
    assert "Alice Research | Lab A | alice@example.org" in [paragraph.text for paragraph in document.paragraphs]
    with zipfile.ZipFile(docx) as archive:
        images = [archive.read(name) for name in archive.namelist()
                  if name.startswith("word/media/") and name.lower().endswith(".png")]
    assert original[figure] in images
    assert all(path.read_bytes() == contents for path, contents in original.items())
    assert not pdf.exists()


@pytest.mark.parametrize("kind", ["relative", "missing", "file"])
def test_conversion_rejects_invalid_resource_directory_before_docx(manuscript, tmp_path, kind):
    resource_dir = {"relative": Path("relative-media"), "missing": tmp_path / "missing", "file": manuscript}[kind]
    output = tmp_path / "invalid-resource.docx"
    with pytest.raises(ValueError, match="Resource directory must be an existing absolute directory"):
        converter.convert_docx(manuscript, output, venue_id="ieee-software", resource_dir=resource_dir)
    assert not output.exists()


def test_real_pandoc_hides_only_author_controls_and_preserves_code_and_relative_media(tmp_path):
    from docx import Document
    from PIL import Image
    binary = _real_pandoc()
    source = tmp_path / "article.md"
    start, end = "<!-- paper-factory:authors:start -->", "<!-- paper-factory:authors:end -->"
    front = "# Shared configuration cleanup\n\n" + start + "\n\nAlice Example | alice@example.org\n\n" + end + "\n\n"
    body = ARTICLE.split("## Abstract", 1)[1]
    examples = ("\n\n## Control marker examples\n\nInline example: `" + start + "` is code.\n\n"
                "```html\n" + start + "\n" + end + "\n```\n\n"
                "Arbitrary <!-- unrelated comment --> prose survives.\n\n"
                "A paragraph containing " + end + " remains literal prose.\n\n![A relative image](figure.png)\n")
    source.write_text(front + "## Abstract" + body + examples, encoding="utf-8")
    Image.new("RGB", (12, 12), "navy").save(tmp_path / "figure.png")
    original = source.read_bytes()
    output = tmp_path / "actual.docx"
    receipt = converter.convert_docx(source, output, venue_id="ieee-software", pandoc=str(binary))
    document = Document(output)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert "Alice Example | alice@example.org" in paragraphs
    assert start not in paragraphs and end not in paragraphs
    assert "Inline example: " + start + " is code." in paragraphs
    assert start + "\n" + end in paragraphs
    assert "Arbitrary <!-- unrelated comment --> prose survives." in paragraphs
    assert "A paragraph containing " + end + " remains literal prose." in paragraphs
    assert len(document.inline_shapes) == 1
    assert source.read_bytes() == original
    assert receipt["input_sha256"] == exporter.digest_file(source)
    assert receipt["working_directory"] == str(tmp_path.resolve())
    import hashlib
    marker_receipt = receipt["author_control_markers"]
    assert marker_receipt["filter_sha256"] == hashlib.sha256(marker_receipt["filter_source"].encode()).hexdigest()


def test_real_pandoc_does_not_hide_unpaired_or_inline_code_only_marker(tmp_path):
    from docx import Document
    binary = _real_pandoc()
    start = "<!-- paper-factory:authors:start -->"
    source = tmp_path / "unpaired.md"
    source.write_text("# Title\n\n" + start + "\n\n`<!-- paper-factory:authors:end -->`\n\n## Abstract\n\nA bounded example.\n\n## Analysis\n\nAn example is preserved.\n", encoding="utf-8")
    output = tmp_path / "unpaired.docx"
    converter.convert_docx(source, output, venue_id="ieee-software", pandoc=str(binary))
    paragraphs = [paragraph.text for paragraph in Document(output).paragraphs]
    assert start in paragraphs
    assert "<!-- paper-factory:authors:end -->" in paragraphs


@pytest.mark.parametrize("missing", [True, False])
def test_real_pandoc_cannot_silently_deliver_a_missing_or_invalid_figure(tmp_path, missing):
    binary = _real_pandoc()
    source = tmp_path / "article.md"
    source.write_text(ARTICLE + "\n![A necessary figure](figure.png)\n", encoding="utf-8")
    if not missing:
        (tmp_path / "figure.png").write_bytes(b"This is not an image.")
    original = source.read_bytes()
    with pytest.raises(ValueError, match="Pandoc failed" if missing else "unreadable raster image") as failure:
        converter.convert_docx(source, tmp_path / "incomplete.docx", venue_id="ieee-software", pandoc=str(binary))
    if missing:
        assert "figure.png" in str(failure.value)
    assert source.read_bytes() == original


def test_pandoc_resolves_only_an_existing_installation(tmp_path, monkeypatch):
    binary = tmp_path / "pandoc.exe"
    binary.write_bytes(b"test executable")
    assert converter.pandoc_binary(str(binary)) == binary.resolve()
    with pytest.raises(ValueError):
        converter.pandoc_binary("relative/pandoc")
    monkeypatch.delenv("PYPANDOC_PANDOC", raising=False)
    monkeypatch.setattr(converter.shutil, "which", lambda _: None)
    def unavailable():
        raise OSError("not installed")
    monkeypatch.setitem(sys.modules, "pypandoc", SimpleNamespace(get_pandoc_path=unavailable))
    with pytest.raises(ValueError, match="Pandoc is unavailable"):
        converter.pandoc_binary()


def test_existing_output_is_never_replaced(manuscript, tmp_path, monkeypatch):
    folder = tmp_path / "existing"
    folder.mkdir()
    sentinel = folder / "keep.txt"
    sentinel.write_bytes(b"original")
    monkeypatch.setattr(exporter, "pandoc_binary", lambda _: manuscript)
    with pytest.raises(FileExistsError):
        exporter.export(manuscript, "ieee-software", folder)
    assert sentinel.read_bytes() == b"original"
    assert list(folder.iterdir()) == [sentinel]


def _mock_native_export(monkeypatch, *, mutate_source=False):
    monkeypatch.setattr(exporter, "pandoc_binary", lambda _: Path(sys.executable))
    _, original = templates.official_template("ieee-software")
    def convert(source, output, **kwargs):
        digest = exporter.digest_file(source)
        output.write_bytes(b"native word fixture")
        if mutate_source:
            source.write_text(ARTICLE + "\nChanged\n", encoding="utf-8")
        return {"input_sha256": digest, "output_sha256": exporter.digest_file(output),
                "official_template": original}
    def render(source, output):
        output.write_bytes(b"%PDF-native fixture")
        return {"source_docx_sha256": exporter.digest_file(source), "visible": False, "read_only": True}
    monkeypatch.setattr(exporter, "convert_docx", convert)
    monkeypatch.setattr(exporter, "word_pdf", render)
    def report(*args, **kwargs):
        receipts = kwargs["conversion_receipts"]
        assert receipts["pdf"]["source_docx_sha256"] == exporter.digest_file(args[1])
        return {"content_format_passed": True, "submission_ready": False,
                "page_qa": {"visual_review_status": "pending"}}
    monkeypatch.setattr(exporter, "compliance_report", report)


def test_export_binds_real_bytes_and_never_infers_approval(manuscript, tmp_path, monkeypatch):
    _mock_native_export(monkeypatch)
    before = manuscript.read_bytes()
    output = tmp_path / "new"
    result = exporter.export(manuscript, "ieee-software", output)
    assert result["content_format_passed"] and not result["submission_ready"]
    assert (output / "manuscript.md").read_bytes() == before == manuscript.read_bytes()
    provenance = json.loads((output / "provenance.json").read_text(encoding="utf-8"))
    assert provenance["model_calls_made_by_tool"] == 0
    assert provenance["author_approval_inferred"] is False
    manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
    assert all(exporter.digest_file(output / row["path"]) == row["sha256"] for row in manifest["files"])


def test_changed_source_fails_but_retains_diagnostic_output(manuscript, tmp_path, monkeypatch):
    _mock_native_export(monkeypatch, mutate_source=True)
    output = tmp_path / "new"
    with pytest.raises(ValueError, match="changed during direct export"):
        exporter.export(manuscript, "ieee-software", output)
    assert (output / "manuscript.md").read_text(encoding="utf-8") == ARTICLE
    assert (output / "manuscript-ieee-software.docx").is_file()
    assert not (output / "provenance.json").exists()


@pytest.mark.parametrize("venue", ["ieee-software", "ieee-security-privacy"])
def test_retained_publisher_styles_columns_and_citations(tmp_path, venue):
    """Exercise actual template mappings, without Pandoc or Microsoft Word."""
    from docx import Document
    from docx.oxml.ns import qn
    template, record = templates.official_template(venue)
    assert exporter.digest_file(template) == record["sha256"]
    doc = Document(template)
    for child in list(doc._element.body):
        if child.tag != qn("w:sectPr"):
            doc._element.body.remove(child)
    for text, style in [("A concrete maintenance decision", "Heading 1"),
                        ("Abstract", "Heading 2"), ("A bounded static example.", "Normal"),
                        ("Keywords: one; two; three", "Normal"), ("Introduction", "Heading 2"),
                        ("Source supports this statement. [1]", "Normal"),
                        ("References", "Heading 2"), ("[1] Authors, A source, 2025.", "Normal")]:
        doc.add_paragraph(text, style)
    reference_heading = next(p for p in doc.paragraphs if p.text == "References")
    caption = doc.add_paragraph("Table 1. Retained seven-column content")
    reference_heading._p.addprevious(caption._p)
    table = doc.add_table(rows=2, cols=7)
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.text = f"r{row_index}c{column_index}"
    reference_heading._p.addprevious(table._tbl)
    expected_cells = [[cell.text for cell in row.cells] for row in table.rows]
    path = tmp_path / (venue + ".docx")
    doc.save(path)
    proof = templates.apply_official_styles(path, template)
    structure = checks._native_structure(path, venue == "ieee-software")
    assert not structure["issues"]
    assert structure["reference_count"] == 1
    assert proof["citation_typography"]["callout_groups_formatted"] == 1
    checked = Document(path)
    assert checked.paragraphs[0].style.name == "Title_document"
    assert any(s._sectPr.find(qn("w:cols")).get(qn("w:num")) == "2" for s in checked.sections)
    assert len(checked.tables) == 1 and len(checked.tables[0].columns) == 7
    assert [[cell.text for cell in row.cells] for row in checked.tables[0].rows] == expected_cells
    assert checked.tables[0].rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
    assert next(p for p in checked.paragraphs if p.text.startswith("Table 1.")).style.name == "TableCaption"


def test_skill_copy_preflight_works_without_old_checkout(tmp_path):
    isolated = tmp_path / "standalone-skill"
    (isolated / "scripts").mkdir(parents=True)
    for path in (ROOT / "scripts").glob("*.py"):
        if path.name in {"export_manuscript.py", "_author_profile.py"} or path.name.startswith("_publication_"):
            shutil.copyfile(path, isolated / "scripts" / path.name)
    shutil.copytree(ROOT / "assets" / "templates", isolated / "assets" / "templates")
    manuscript = tmp_path / "input.md"
    manuscript.write_text(ARTICLE, encoding="utf-8")
    env = {key: value for key, value in os.environ.items() if key not in {"PYTHONPATH", "PAPER_FACTORY_ROOT"}}
    result = subprocess.run([sys.executable, str(isolated / "scripts" / "export_manuscript.py"),
                             "preflight", "--manuscript", str(manuscript), "--venue", "ieee-software"],
                            cwd=tmp_path, env=env, capture_output=True, text=True, encoding="utf-8", timeout=30)
    assert result.returncode == 0, result.stderr
    report = json.loads(result.stdout)
    assert report["prefilter_passed"] and report["model_calls_made_by_tool"] == 0
    assert Path(report["official_template"]["path"]).is_relative_to(isolated)
    assert not (isolated / "desktop").exists() and not (isolated / "src").exists()
