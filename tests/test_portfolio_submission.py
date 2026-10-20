"""Submission workflow checks on explicitly synthetic source studies only."""

import base64
import copy
from datetime import datetime, timezone
import hashlib
import io
import json
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

from paper_factory import ipc
from paper_factory.autonomous.portfolio_submission import PORTALS, _submission_manuscript, _submission_biography
from paper_factory.workflow import WorkflowError
from test_repository_paper_workflow import study, draft, review
from test_repository_portfolio import portfolio, analyze_portfolio, retain_portfolio_response
from test_repository_portfolio_author_export import author_draft, PROFILE, COUNTS


PNG = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+j/aMAAAAASUVORK5CYII="


def submission(service, rid):
    return service.status(rid)["repository_paper"]["portfolio"]["submission"]


@pytest.fixture
def ready(author_draft, monkeypatch):
    service, rid, venue_id, _ = author_draft
    service.set_author_metadata(rid, PROFILE)
    service.export_portfolio_author(rid, venue_id)

    def native_fixture(source, output, **kwargs):
        if output.suffix == ".docx":
            document = Document()
            document.add_paragraph("Explicitly synthetic supplementary submission fixture; no scholarly claim is established.")
            document.add_paragraph(source.read_text(encoding="utf-8"))
            document.add_table(rows=2, cols=2)
            document.save(output)
        else:
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            writer.write(output)
        return {"input_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "output_sha256": hashlib.sha256(output.read_bytes()).hexdigest(), "engine": "synthetic-submission-fixture"}

    monkeypatch.setattr("paper_factory.conversion.convert", native_fixture)
    def compliance_fixture(markdown, docx, venue, **kwargs):
        whole_words = 4650 if venue == "ieee-security-privacy" else 3000
        return {"content_format_passed": True, "blockers": [], "pending": ["Explicitly synthetic author declarations remain pending."],
            "items": [{"id": "native_word", "stage": "native", "status": "passed", "reason": "Synthetic native fixture"},
                      {"id": "author_identity", "stage": "submission", "status": "pending", "reason": "Synthetic human fixture"}],
            "binding": {"markdown_sha256": hashlib.sha256(markdown.read_bytes()).hexdigest(),
                "docx_sha256": hashlib.sha256(docx.read_bytes()).hexdigest(),
                "pdf_sha256": hashlib.sha256(kwargs["pdf"].read_bytes()).hexdigest()},
            "metrics": {"native_word": {"whole_words": whole_words, "pages": 1, "abstract_words": 20, "reference_count": 2},
                        "effective_words": whole_words + 250, "word_count_formula": "Synthetic native whole words plus visual allowance"},
            "page_qa": {"pages": [], "page_count": 1}, "scope": "Explicit synthetic fixture only, no native Word measurement is established."}
    monkeypatch.setattr("paper_factory.journal_compliance.compliance_report", compliance_fixture)
    return service, rid, venue_id


def prepare_value(service, rid, venue_id):
    candidate = next(item for item in submission(service, rid)["candidates"] if item["venue_id"] == venue_id)
    return {**candidate["expected_binding"], "author_biography": "Synthetic Author is a researcher at Synthetic Institute. This explicit fixture does not establish a human identity or affiliation.",
        "biography_details": {key: "Explicitly synthetic " + key for key in ("role", "location", "research_interests", "highest_degree", "degree_institution", "professional_membership")},
        "keywords": ["synthetic fixture", "source analysis", "software design"],
        "photo": {"filename": "fixture.png", "media_type": "image/png", "base64": PNG},
        "funding_statement": "This synthetic fixture has no actual funding assertion.",
        "conflicts_statement": "This synthetic fixture does not assert real conflicts or their absence.",
        "cover_letter_note": "This is an explicitly synthetic test package, not an actual journal submission."}


def confirm_value(service, rid, *, word_count=3000):
    prepared = submission(service, rid)["preparation"]
    return {"preparation_sha256": prepared["sha256"], "expected_docx_sha256": prepared["docx_sha256"],
        "expected_pdf_sha256": prepared["pdf_sha256"],
        "qa": {"word_count": word_count, "docx_pages": 1, "pdf_pages": 1, "table_count": 1, "figure_count": 0},
        "declarations": {key: True for key in ("identity_confirmed", "all_authors_approved", "originality_confirmed", "rights_confirmed", "ai_disclosure_reviewed", "guidelines_reviewed", "all_pages_reviewed")}}


def record_value(service, rid, status="submitted", *, manuscript_id="SYNTHETIC-PORTAL-ID"):
    return {"preparation_sha256": submission(service, rid)["preparation"]["sha256"], "status": status,
        "manuscript_id": manuscript_id, "occurred_at": datetime.now(timezone.utc).isoformat(),
        "confirmed": True, "note": "Explicitly synthetic portal record; no real publication action occurred."}


def test_prepare_generates_new_files_without_approval_and_preserves_all_scientific_bytes(ready):
    service, rid, venue_id = ready
    before = service.status(rid)
    original = {key: service.artifact_path(rid, key).read_bytes() for key in before["artifacts"]}
    value = prepare_value(service, rid, venue_id)
    value["author_biography"] += " <script>alert(1)</script> {{source:invented}} ![unsafe](https://example.invalid)"
    current = service.prepare_portfolio_submission(rid, venue_id, value)
    state = current["repository_paper"]["portfolio"]["submission"]
    prepared = state["preparation"]
    assert state["status"] == "awaiting_file_check" and state["portal_record"] is None and state["history"] == []
    assert prepared["current"] and prepared["generation_status"] == "completed"
    assert prepared["confirmation"] is prepared["package_artifact_id"] is prepared["package_sha256"] is None
    source = service.artifact_path(rid, prepared["source_artifact_id"]).read_bytes()
    output = service.artifact_path(rid, prepared["md_artifact_id"]).read_bytes()
    assert prepared["provided_author_biography"] == value["author_biography"]
    assert all(fact in prepared["author_biography"] for fact in value["biography_details"].values())
    assert output == _submission_manuscript(source, prepared["author_biography"], value["keywords"], **prepared["supplements"])[0]
    assert value["funding_statement"].replace(".", r"\.").encode() in output
    assert value["conflicts_statement"].replace(".", r"\.").encode() in output
    assert b"Keywords: synthetic fixture; source analysis; software design" in output
    assert b"## Author Biography" in output and b"\\<script\\>" in output
    assert b"\\{\\{source:invented\\}\\}" in output and b"\\!\\[unsafe\\]" in output
    assert all(service.artifact_path(rid, key).read_bytes() == raw for key, raw in original.items())
    assert {key: current[key] for key in COUNTS} == {key: before[key] for key in COUNTS}
    assert {key for key in current["artifacts"] if key.startswith("model-evidence-")} == {key for key in before["artifacts"] if key.startswith("model-evidence-")}
    assert all(not item["submission_ready"] for item in current["repository_paper"]["portfolio"]["venues"])


def test_preparation_rejects_stale_original_binding_and_missing_software_photo(ready):
    service, rid, venue_id = ready
    original = service.status(rid)["artifacts"]
    value = prepare_value(service, rid, venue_id)
    value["expected_review_sha256"] = "0" * 64
    with pytest.raises(WorkflowError) as error:
        service.prepare_portfolio_submission(rid, venue_id, value)
    assert error.value.code == "PORTFOLIO_SUBMISSION_STALE"
    value = prepare_value(service, rid, venue_id)
    value["photo"] = None
    with pytest.raises(ValueError, match="photograph"):
        service.prepare_portfolio_submission(rid, venue_id, value)
    assert service.status(rid)["artifacts"] == original


@pytest.mark.parametrize("check", ["identity_confirmed", "all_authors_approved", "originality_confirmed", "rights_confirmed", "ai_disclosure_reviewed", "guidelines_reviewed", "all_pages_reviewed"])
def test_file_confirmation_requires_each_explicit_human_check(ready, check):
    service, rid, venue_id = ready
    service.prepare_portfolio_submission(rid, venue_id, prepare_value(service, rid, venue_id))
    before = service.status(rid)["artifacts"]
    value = confirm_value(service, rid)
    value["declarations"][check] = False
    with pytest.raises(ValueError, match="affirmative human"):
        service.confirm_portfolio_submission(rid, value)
    assert service.status(rid)["artifacts"] == before
    assert submission(service, rid)["status"] == "awaiting_file_check"


def test_confirmation_checks_new_hash_visuals_and_word_budget_before_package_writes(ready):
    service, rid, venue_id = ready
    service.prepare_portfolio_submission(rid, venue_id, prepare_value(service, rid, venue_id))
    before = service.status(rid)["artifacts"]
    for damage, expected_error in (("hash", "PORTFOLIO_SUBMISSION_STALE"), ("visuals", "PORTFOLIO_SUBMISSION_VISUALS"), ("length", "PORTFOLIO_SUBMISSION_QA")):
        value = confirm_value(service, rid)
        if damage == "hash": value["expected_docx_sha256"] = "0" * 64
        elif damage == "visuals": value["qa"]["table_count"] = 0
        else: value["qa"]["word_count"] = 4200
        with pytest.raises(WorkflowError) as error:
            service.confirm_portfolio_submission(rid, value)
        assert error.value.code == expected_error
        assert service.status(rid)["artifacts"] == before


def test_successful_confirmation_packages_exact_new_files_and_a_hash_manifest(ready):
    service, rid, venue_id = ready
    before = service.status(rid)
    service.prepare_portfolio_submission(rid, venue_id, prepare_value(service, rid, venue_id))
    state_before = submission(service, rid)
    current = service.confirm_portfolio_submission(rid, confirm_value(service, rid))
    state = current["repository_paper"]["portfolio"]["submission"]
    prepared = state["preparation"]
    assert state["status"] == "prepared" and state["portal_record"] is None and state["history"] == []
    assert prepared["sha256"] == state_before["preparation"]["sha256"]
    assert prepared["confirmation"]["qa"]["table_count"] == 1
    assert prepared["confirmation"]["metrics"]["effective_words"] == 3250
    raw = service.artifact_path(rid, prepared["package_artifact_id"]).read_bytes()
    assert hashlib.sha256(raw).hexdigest() == prepared["package_sha256"]
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        assert archive.read("paper.docx") == service.artifact_path(rid, prepared["docx_artifact_id"]).read_bytes()
        assert archive.read("paper.pdf") == service.artifact_path(rid, prepared["pdf_artifact_id"]).read_bytes()
        final_compliance = json.loads(archive.read("journal-compliance.json"))
        assert final_compliance["submission_ready"] and final_compliance["human_confirmation"]["declarations"]["identity_confirmed"]
        assert final_compliance["binding"] == prepared["compliance"]["binding"]
        assert archive.read("paper.md") == service.artifact_path(rid, prepared["md_artifact_id"]).read_bytes()
        assert archive.read("author-photo.png") == base64.b64decode(PNG)
        for item in manifest["files"]:
            content = archive.read(item["name"])
            assert len(content) == item["size"] and hashlib.sha256(content).hexdigest() == item["sha256"]
        assert "no concurrent submission" in archive.read("cover-letter.txt").decode()
    assert {key: current[key] for key in COUNTS} == {key: before[key] for key in COUNTS}
    with pytest.raises(WorkflowError, match="once"):
        service.confirm_portfolio_submission(rid, confirm_value(service, rid))


def test_portal_records_require_actual_receipt_and_block_concurrent_venue_or_repreparation(ready):
    service, rid, venue_id = ready
    value = prepare_value(service, rid, venue_id)
    service.prepare_portfolio_submission(rid, venue_id, value)
    with pytest.raises(WorkflowError):
        service.record_portfolio_submission(rid, record_value(service, rid))
    service.confirm_portfolio_submission(rid, confirm_value(service, rid))
    record = record_value(service, rid)
    for field in ("manuscript_id", "confirmed"):
        bad = copy.deepcopy(record)
        bad[field] = "" if field == "manuscript_id" else False
        with pytest.raises(ValueError):
            service.record_portfolio_submission(rid, bad)
    submitted = service.record_portfolio_submission(rid, record)
    state = submitted["repository_paper"]["portfolio"]["submission"]
    assert state["status"] == "submitted" and not state["portal_record"]["externally_verified"]
    assert state["portal_record"]["portal_url"] == PORTALS[venue_id]
    assert all(not candidate["prepare_available"] for candidate in state["candidates"])
    with pytest.raises(WorkflowError) as error:
        service.prepare_portfolio_submission(rid, venue_id, value)
    assert error.value.code == "PORTFOLIO_SUBMISSION_BLOCKED"
    with pytest.raises(WorkflowError):
        service.record_portfolio_submission(rid, record)
    bad = record_value(service, rid, "withdrawn", manuscript_id="SYNTHETIC-WRONG-ID")
    with pytest.raises(WorkflowError) as error:
        service.record_portfolio_submission(rid, bad)
    assert error.value.code == "PORTFOLIO_SUBMISSION_ID"
    released = service.record_portfolio_submission(rid, record_value(service, rid, "withdrawn"))
    assert released["repository_paper"]["portfolio"]["submission"]["status"] == "withdrawn"
    assert len(submission(service, rid)["history"]) == 2
    with pytest.raises(WorkflowError):
        service.record_portfolio_submission(rid, record_value(service, rid))
    reprepare = service.prepare_portfolio_submission(rid, venue_id, prepare_value(service, rid, venue_id))
    assert reprepare["repository_paper"]["portfolio"]["submission"]["status"] == "awaiting_file_check"
    assert len(submission(service, rid)["history"]) == 2


def test_changed_author_makes_preparation_stale_but_keeps_submitted_package_and_event(ready):
    service, rid, venue_id = ready
    service.prepare_portfolio_submission(rid, venue_id, prepare_value(service, rid, venue_id))
    service.confirm_portfolio_submission(rid, confirm_value(service, rid))
    original = submission(service, rid)
    changed = service.set_author_metadata(rid, {**PROFILE, "display_name": "Changed Synthetic Author"})
    stale = changed["repository_paper"]["portfolio"]["submission"]
    assert stale["status"] == "stale" and not stale["preparation"]["current"]
    assert stale["preparation"]["sha256"] == original["preparation"]["sha256"]
    with pytest.raises(WorkflowError):
        service.record_portfolio_submission(rid, record_value(service, rid))
    service.set_author_metadata(rid, PROFILE)
    service.record_portfolio_submission(rid, record_value(service, rid))
    event = submission(service, rid)["portal_record"]
    changed = service.set_author_metadata(rid, {**PROFILE, "display_name": "Again Changed Synthetic Author"})
    state = changed["repository_paper"]["portfolio"]["submission"]
    assert state["status"] == "submitted" and not state["preparation"]["current"]
    assert state["portal_record"] == event
    assert state["preparation"]["package_sha256"] == original["preparation"]["package_sha256"]


def test_generation_failure_and_pending_funding_keep_unconfirmed_outputs(ready, monkeypatch):
    service, rid, venue_id = ready
    convert = __import__("paper_factory.conversion", fromlist=["convert"]).convert

    def fail_pdf(source, output, **kwargs):
        if output.suffix == ".pdf":
            output.write_bytes(b"Explicitly synthetic interrupted PDF bytes.")
            raise ValueError("Synthetic PDF conversion failure")
        return convert(source, output)

    monkeypatch.setattr("paper_factory.conversion.convert", fail_pdf)
    service.prepare_portfolio_submission(rid, venue_id, prepare_value(service, rid, venue_id))
    assert submission(service, rid)["status"] == "generation_failed"
    assert submission(service, rid)["preparation"]["confirmation"] is None
    with pytest.raises(WorkflowError):
        service.confirm_portfolio_submission(rid, confirm_value(service, rid))
    monkeypatch.setattr("paper_factory.conversion.convert", convert)
    value = prepare_value(service, rid, venue_id)
    value["funding_statement"] = "저자 확인 대기"
    service.prepare_portfolio_submission(rid, venue_id, value)
    with pytest.raises(ValueError, match="funding"):
        service.confirm_portfolio_submission(rid, confirm_value(service, rid))
    assert submission(service, rid)["status"] == "awaiting_file_check"


def test_strict_ipc_routes_and_untrusted_photo_are_rejected_before_mutation(ready):
    service, rid, venue_id = ready
    assert ipc.METHODS["workflow.preparePortfolioSubmission"] == {"researchId", "venueId", "value"}
    assert ipc.METHODS["workflow.confirmPortfolioSubmission"] == {"researchId", "value"}
    assert ipc.METHODS["workflow.recordPortfolioSubmission"] == {"researchId", "value"}
    before = service.status(rid)["artifacts"]
    for damage in ("bytes", "mime", "path", "large"):
        value = prepare_value(service, rid, venue_id)
        if damage == "bytes": value["photo"]["base64"] = base64.b64encode(b"not a real image").decode()
        elif damage == "mime": value["photo"]["media_type"] = "image/jpeg"
        elif damage == "path": value["photo"]["filename"] = "../photo.png"
        else: value["photo"]["base64"] = base64.b64encode(b"x" * (512 * 1024 + 1)).decode()
        with pytest.raises(ValueError):
            service.prepare_portfolio_submission(rid, venue_id, value)
        assert service.status(rid)["artifacts"] == before


def admit_second_venue(service, rid):
    venue_id = "ieee-security-privacy"
    value = draft(service, rid, prepare_analysis=False)
    value["title"] += " for the second synthetic venue"
    key, _ = retain_portfolio_response(service, rid, value, "venue-manuscript", venue_id)
    service.submit_portfolio_manuscript(rid, {"venue_id": venue_id, "inference_artifact_id": key, "manuscript": value})
    assessment = review(True)
    key, _ = retain_portfolio_response(service, rid, assessment, "venue-manuscript-review", venue_id)
    service.submit_portfolio_review(rid, {"venue_id": venue_id, "inference_artifact_id": key, "review": assessment})
    service.export_portfolio_author(rid, venue_id)
    return venue_id


def test_sp_minimum_uses_complete_new_word_count_without_double_counting_biography(ready):
    service, rid, _ = ready
    venue_id = admit_second_venue(service, rid)
    value = prepare_value(service, rid, venue_id)
    value["photo"] = None
    service.prepare_portfolio_submission(rid, venue_id, value)
    with pytest.raises(WorkflowError) as error:
        service.confirm_portfolio_submission(rid, confirm_value(service, rid, word_count=4649))
    assert error.value.code == "PORTFOLIO_SUBMISSION_QA"
    service.confirm_portfolio_submission(rid, confirm_value(service, rid, word_count=4650))
    metrics = submission(service, rid)["preparation"]["confirmation"]["metrics"]
    assert metrics["effective_words"] == 4900 and metrics["biography_words"] > 0


def test_existing_keywords_and_biography_are_replaced_once_without_changing_science():
    original = b"# Synthetic\n\n## Abstract\n\nSynthetic abstract.\nKeywords: former; metadata; only\n\n## Analysis\n\nLiteral synthetic scientific body.\n\n## References\n\n[1] Synthetic reference.\n\n## Author Biography\n\nPrevious supplied author metadata.\n"
    output, proof = _submission_manuscript(original, "Explicitly synthetic replacement author biography only.", ["new one", "new two", "new three"])
    assert output.count(b"Keywords:") == output.count(b"## Author Biography") == 1
    assert b"Literal synthetic scientific body." in output and b"[1] Synthetic reference." in output
    assert b"former; metadata" not in output and b"Previous supplied" not in output
    assert proof["original_sha256"] == hashlib.sha256(original).hexdigest()


@pytest.mark.parametrize("existing_ack", [False, True])
def test_actual_declaration_supplements_preserve_science_and_existing_ai_acknowledgment(existing_ack):
    source = ("# Synthetic\n\n## Abstract\n\nSynthetic abstract.\n\n## Evidence\n\nLiteral scientific claim.\n\n"
              + ("## Acknowledgment\n\nActual retained AI disclosure.\n\n" if existing_ack else "")
              + "## References\n\n[1] Actual retained reference.\n").encode()
    old, old_proof = _submission_manuscript(source, "Actual supplied biography text.", ["one", "two", "three"])
    explicit_old, explicit_proof = _submission_manuscript(source, "Actual supplied biography text.", ["one", "two", "three"], funding_statement=None, conflicts_statement=None)
    assert (old, old_proof) == (explicit_old, explicit_proof)
    new, proof = _submission_manuscript(source, "Actual supplied biography text.", ["one", "two", "three"], funding_statement="Actual support from X.", conflicts_statement="None declared by the author.")
    assert b"Literal scientific claim." in new and b"[1] Actual retained reference." in new
    assert new.count(b"## Acknowledgment") == 1
    assert b"Funding: Actual support from X\\." in new and b"Conflicts of interest: None declared by the author\\." in new
    assert new.index(b"Funding:") < new.index(b"## References")
    assert (b"Actual retained AI disclosure." in new) == existing_ack
    assert proof["original_sha256"] == hashlib.sha256(source).hexdigest()


def test_actual_biography_facts_are_rendered_as_one_paragraph_without_inference():
    supplied = {key: "Actual supplied " + key for key in ("role", "location", "research_interests", "highest_degree", "degree_institution", "professional_membership")}
    result = _submission_biography("Actual author prose.\n\nAdditional supplied sentence.", supplied)
    assert "\n" not in result and all(value in result for value in supplied.values())
    assert "approved" not in result and "Actual author prose." in result


def test_unpassed_native_format_is_retained_and_blocks_final_package(ready, monkeypatch):
    service, rid, venue = ready
    original_report = __import__("paper_factory.journal_compliance", fromlist=["compliance_report"]).compliance_report
    def failed_report(*args, **kwargs):
        report = original_report(*args, **kwargs)
        return {**report, "content_format_passed": False, "blockers": ["Explicitly synthetic native abstract overflow."]}
    monkeypatch.setattr("paper_factory.journal_compliance.compliance_report", failed_report)
    before = service.status(rid)
    service.prepare_portfolio_submission(rid, venue, prepare_value(service, rid, venue))
    state = submission(service, rid)
    assert state["status"] == "generation_failed" and not state["preparation"]["compliance"]["content_format_passed"]
    assert any("abstract overflow" in message for message in state["preparation"]["errors"])
    with pytest.raises(WorkflowError):
        service.confirm_portfolio_submission(rid, confirm_value(service, rid))
    assert {key: service.status(rid)[key] for key in COUNTS} == {key: before[key] for key in COUNTS}


def test_other_ready_venue_requires_explicit_release_of_current_selection(ready):
    service, rid, venue_id = ready
    other = admit_second_venue(service, rid)
    service.prepare_portfolio_submission(rid, venue_id, prepare_value(service, rid, venue_id))
    candidate = next(item for item in submission(service, rid)["candidates"] if item["venue_id"] == other)
    assert not candidate["prepare_available"] and any("현재 준비" in text for text in candidate["missing"])
    with pytest.raises(WorkflowError) as error:
        service.prepare_portfolio_submission(rid, other, prepare_value(service, rid, other))
    assert error.value.code == "PORTFOLIO_SUBMISSION_BLOCKED"
    released = service.record_portfolio_submission(rid, record_value(service, rid, "withdrawn", manuscript_id=""))
    state = released["repository_paper"]["portfolio"]["submission"]
    assert state["status"] == "withdrawn" and "no portal submission" in state["portal_record"]["scope"]
    service.prepare_portfolio_submission(rid, other, prepare_value(service, rid, other))
    assert submission(service, rid)["selected_venue_id"] == other
    assert submission(service, rid)["status"] == "awaiting_file_check" and len(submission(service, rid)["history"]) == 1


def test_pending_scientific_revision_invalidates_generated_submission_files(ready):
    service, rid, venue_id = ready
    service.prepare_portfolio_submission(rid, venue_id, prepare_value(service, rid, venue_id))
    original = submission(service, rid)["preparation"]
    service.submit_portfolio_revision_feedback(rid, venue_id,
        "This explicitly synthetic feedback requests a supported interpretation revision without supplying approval or new evidence.")
    state = submission(service, rid)
    assert state["status"] == "stale" and not state["preparation"]["current"]
    assert state["preparation"]["sha256"] == original["sha256"]
    assert any("원고 보완" in text for text in state["candidates"][0]["missing"])
    with pytest.raises(WorkflowError):
        service.confirm_portfolio_submission(rid, confirm_value(service, rid))


def finish_quality(service, rid, *, recommendation="consider_submission"):
    """Consume explicit synthetic receipts through the real bounded native flow."""
    from test_manuscript_improvement import assessment, improvement_plan, reserve, retain, revision

    while True:
        state = service.status(rid)["repository_paper"]["manuscript_improvement"]
        if state["next_task"] is None:
            break
        task = reserve(service, rid)
        if task["purpose"] == "review":
            revised = len(state["versions"][task["venue_id"]]) > 1
            value = assessment(task, 8 if revised else 6)
            if revised:
                value["recommendation"] = recommendation
            retain(service, rid, task, value)
        elif task["purpose"] == "plan":
            retain(service, rid, task, improvement_plan())
        else:
            retain(service, rid, task, revision(service, rid, task))
    return service.finish_manuscript_improvement(rid)


def test_quality_start_stales_old_preparation_without_changing_old_files_or_counts(ready):
    service, rid, venue_id = ready
    admit_second_venue(service, rid)
    old_value = prepare_value(service, rid, venue_id)
    service.prepare_portfolio_submission(rid, venue_id, old_value)
    before = service.status(rid)
    old_preparation = submission(service, rid)["preparation"]
    old_bytes = {key: service.artifact_path(rid, key).read_bytes() for key in before["artifacts"]}

    started = service.start_manuscript_improvement(rid)
    state = started["repository_paper"]["portfolio"]["submission"]
    assert state["status"] == "stale" and not state["preparation"]["current"]
    assert state["preparation"]["sha256"] == old_preparation["sha256"]
    assert all(not item["prepare_available"] for item in state["candidates"])
    assert all(service.artifact_path(rid, key).read_bytes() == raw for key, raw in old_bytes.items())
    assert {key: started[key] for key in COUNTS} == {key: before[key] for key in COUNTS}
    with pytest.raises(WorkflowError) as error:
        service.prepare_portfolio_submission(rid, venue_id, old_value)
    assert error.value.code == "PORTFOLIO_SUBMISSION_BLOCKED"


def test_quality_submission_uses_best_official_author_output_and_never_infers_approval(ready):
    service, rid, venue_id = ready
    admit_second_venue(service, rid)
    before = service.status(rid)
    old_draft = next(item for item in before["repository_paper"]["portfolio"]["venues"] if item["venue_id"] == venue_id)
    service.start_manuscript_improvement(rid)
    finished = finish_quality(service, rid)
    quality = finished["repository_paper"]["manuscript_improvement"]
    candidate = next(item for item in submission(service, rid)["candidates"] if item["venue_id"] == venue_id)
    exported = quality["exports"][venue_id]
    selected = exported["selected_version"]
    assert candidate["prepare_available"] and not quality["target_achieved"]
    assert candidate["expected_binding"]["expected_draft_sha256"] == selected["draft_sha256"]
    assert candidate["expected_binding"]["expected_draft_sha256"] != old_draft["draft_sha256"]
    assert candidate["expected_binding"]["expected_review_sha256"] == finished["artifacts"][exported["review_artifact_id"]]["sha256"]
    assert candidate["expected_binding"]["expected_docx_sha256"] == exported["outputs"]["docx"]["sha256"]
    assert {key: finished[key] for key in COUNTS} == {key: before[key] for key in COUNTS}

    prepared = service.prepare_portfolio_submission(rid, venue_id, prepare_value(service, rid, venue_id))
    state = prepared["repository_paper"]["portfolio"]["submission"]
    assert state["status"] == "awaiting_file_check" and state["preparation"]["confirmation"] is None
    assert state["preparation"]["package_artifact_id"] is None and state["portal_record"] is None
    assert service.artifact_path(rid, state["preparation"]["source_artifact_id"]).read_bytes() == service.artifact_path(rid, exported["input_artifact_id"]).read_bytes()
    assert not service.confirm_portfolio_submission(rid, confirm_value(service, rid))["repository_paper"]["portfolio"]["submission"]["preparation"]["externally_verified"]


def test_nonrecommended_quality_blocks_legacy_pass_and_changed_author_blocks_quality_export(ready):
    service, rid, venue_id = ready
    admit_second_venue(service, rid)
    legacy = prepare_value(service, rid, venue_id)
    service.start_manuscript_improvement(rid)
    finish_quality(service, rid, recommendation="not_recommended")
    candidate = next(item for item in submission(service, rid)["candidates"] if item["venue_id"] == venue_id)
    assert not candidate["prepare_available"] and any("투고 검토를 권하지" in text for text in candidate["missing"])
    with pytest.raises(WorkflowError) as error:
        service.prepare_portfolio_submission(rid, venue_id, legacy)
    assert error.value.code == "PORTFOLIO_SUBMISSION_BLOCKED"
    service.set_author_metadata(rid, {**PROFILE, "display_name": "Changed Synthetic Author"})
    candidate = next(item for item in submission(service, rid)["candidates"] if item["venue_id"] == venue_id)
    assert not candidate["prepare_available"] and any("저자 정보를 저장" in text for text in candidate["missing"])


def test_native_converter_helper_bytes_are_frozen_with_the_generated_manuscript(ready, monkeypatch):
    service, rid, venue_id = ready
    convert = __import__("paper_factory.conversion", fromlist=["convert"]).convert

    def convert_with_helpers(source, output, **kwargs):
        value = convert(source, output)
        output.with_name("submission-table-layout.lua").write_text("-- Explicitly synthetic retained converter helper\n", encoding="utf-8")
        if output.suffix == ".pdf":
            output.with_suffix(".typ").write_text("// Explicitly synthetic retained typesetting input\n", encoding="utf-8")
            output.with_name("submission-pandoc-metadata.json").write_text('{"synthetic": true}\n', encoding="utf-8")
        return value

    monkeypatch.setattr("paper_factory.conversion.convert", convert_with_helpers)
    current = service.prepare_portfolio_submission(rid, venue_id, prepare_value(service, rid, venue_id))
    keys = [key for key in current["artifacts"] if key.startswith("repository-portfolio-submission-generated-")]
    assert len(keys) == 3
    service.artifact_path(rid, keys[0]).write_bytes(b"Changed synthetic helper bytes")
    with pytest.raises(WorkflowError) as error:
        service.status(rid)
    assert error.value.code == "ARTIFACT_CHANGED"
