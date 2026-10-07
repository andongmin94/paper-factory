"""Synthetic orchestration contracts; positive review fixtures prove no research value."""

import base64
import copy
from concurrent.futures import Future, ThreadPoolExecutor
import hashlib
import json
import os
from pathlib import Path
import re
import threading
import time
from types import SimpleNamespace
import zipfile

import pytest

from paper_factory.autonomous import literature, science
from paper_factory.autonomous.models import FrozenArtifact, ResearchPlan
from paper_factory.workflow import WorkflowError, WorkflowService, _scientific_inputs
from paper_factory.workflow_models import Workflow
from paper_factory.workspace import Workspace, digest_file, write_json


def protocol():
    return {
        "feasible": True, "reason": "A production callable supports controlled synthetic contract checks.",
        "title": "Controlled transformation contract study for native workflow validation",
        "question": "How does the selected production transformation preserve protected fixture values?",
        "research_gap": "Synthetic orchestration fixture; it does not establish a gap in actual literature.",
        "expected_contribution": "Synthetic positive proposal for workflow tests, with no claim of academic contribution.",
        "comparison_rationale": "The comparator exercises paired fixture plumbing and is not a research baseline.",
        "sampling_rationale": "The fixed annotations exercise grid validation rather than representative sampling.",
        "runtime": "quickjs", "source_files": ["transform.js"], "production_entrypoint": "transform.js:transform",
        "dependencies": [], "conditions": ["production", "ablation"],
        "metrics": [{"name": "error", "unit": "events", "description": "Independently detected transformation contract violations."}],
        "comparator": "An explicit guard removal ablation retains the other transformation rules.",
        "independent_oracle": "Fixture annotations define expected values independently of the comparator.",
        "sampling_unit": "An independently seeded input with protected fixture annotations.",
        "units_per_seed": 3, "seeds": [11, 37], "parameters": {"fixture_size": 12},
        "procedure": ["Generate seeded inputs and independent annotations.", "Invoke unchanged source and an explicit comparator.",
                      "Measure paired errors and inspect positive and negative controls."],
        "analysis_method": "Descriptive statistics with differences paired by seed and fixture unit.",
        "limitations": ["Synthetic inputs do not represent natural populations.", "Instrumentation affects measured execution time."],
        "literature_queries": ["independent software testing oracle"],
        "research_claim": {"mode": "finite_enumeration", "claim": "Synthetic controlled fixture checks establish no real academic contribution.",
                           "scope": "Only synthetic orchestration inputs; no population or universal correctness claim.",
                           "importance": "This simulated declaration tests product contracts and establishes no scholarly importance.",
                           "validation_plan": "Retained synthetic fixture rows exercise deterministic paired-result binding."},
    }


def observations():
    data = b'{"synthetic":true,"annotations":[0,1,2]}'
    return {"observations": [
        {"unit_id": f"case-{index}", "seed": seed, "condition": condition, "metric": "error",
         "value": 0 if condition == "production" else index + 1}
        for seed in (11, 37) for index in range(3) for condition in ("production", "ablation")
    ], "controls": [
        {"name": "positive control", "passed": True, "details": "Input agrees with independent expected annotations."},
        {"name": "negative control", "passed": True, "details": "An intentional deletion is detected by the fixture oracle."},
    ], "fixtures": [{"label": "synthetic-inputs.json", "encoding": "base64",
                      "content": base64.b64encode(data).decode(), "sha256": hashlib.sha256(data).hexdigest()}]}


def manuscript():
    paragraph = (
        "This synthetic manuscript fixture checks the transfer of a frozen software protocol into an evidence linked document. "
        "Its observations are simulated inputs for orchestration validation and cannot establish production correctness. "
        "The comparison separates a baseline condition from a declared ablation while preserving paired sampling units. "
        "The trusted analyzer determines descriptive statistics from retained records rather than accepting model supplied measurements. "
        "Source hashes and independent reproduction protect the relationship between these records and displayed numerical references. "
        "The fixture exercises document integrity checks and offers no evidence of scientific novelty or model inference. "
    )
    sections = []
    for heading in science.REQUIRED_SECTIONS:
        text = paragraph * 2
        if heading == "Abstract":
            text += "The paired mean difference is {{result:error.paired_2_minus_1.mean}} events."
        elif heading == "Results":
            text += ("The production mean is {{result:error.condition_1.mean}} events and the ablation mean is "
                     "{{result:error.condition_2.mean}} events. The paired mean difference is "
                     "{{result:error.paired_2_minus_1.mean}} events.")
        elif heading == "Experimental Setup":
            text += "The protocol uses {{parameter:units_per_seed}} units for each seed."
            text += " Recorded instrumentation: {{parameter:setting.execution_instrumentation}}."
        elif heading == "Related Work":
            text += "The synthetic passage describes independently defined expected outcomes {{citation:fixture-oracle}}."
        sections.append({"heading": heading, "text": text})
    return {"title": protocol()["title"], "sections": sections}


REVIEW = {"accepted": True, "issues": [], "checks": ["Synthetic native host review fixture"]}
SYNTHETIC_PASSAGE = ("Synthetic passage: independent software test oracles define expected outcomes separately from the implementation. "
                     "Controlled fixtures compare behavior with annotations and cannot establish population failure rates.")
SYNTHETIC_ABSTRACT = "Synthetic abstract describes fixture checks without establishing actual prior-work methods, findings, or any scientific importance."
SYNTHETIC_PREFIX = "Abstract\n" + SYNTHETIC_ABSTRACT + "\n\nIntroduction\n"
SYNTHETIC_SUFFIX = "\nReferences\nSynthetic bibliography is contextual metadata and is not a research method or finding."
SYNTHETIC_BODY_START = len(SYNTHETIC_PREFIX)


def publication_readiness(*, manuscript=False):
    return {**{name: {"passed": True, "reason": "Synthetic positive fixture only; no real journal readiness is attested."}
               for name in ("novelty", "significance", "validation")},
            "claim": protocol()["research_claim"]["claim"], "scope": protocol()["research_claim"]["scope"],
            "evidence_mode": "finite_enumeration",
            "evidence_basis": "Synthetic deterministic records bind this simulated assessment to fixture results.",
            "closest_work": [{"source_id": "fixture-oracle", "excerpt_index": 0, "quote": SYNTHETIC_PASSAGE[:100],
                              "known_result": "Synthetic independent-oracle passage is a test artifact rather than actual prior work.",
                              "difference": "This synthetic comparison exercises binding without asserting any scientific novelty."}],
            "analysis_keys": ["error.paired_2_minus_1.mean"] if manuscript else [], "fixture_labels": [],
            "proof_section": None, "proof_quote": None}


STUDY_CRITERIA = ("question", "contribution", "literature", "comparison", "sampling", "feasibility")
STUDY_REVIEW = {"accepted": True, "issues": [], **{
    name: {"passed": True, "reason": "Synthetic positive decision fixture for orchestration; not an academic assessment."}
    for name in STUDY_CRITERIA}, "selected_sources": [{
        "source_id": "fixture-oracle", "excerpt_index": 0,
        "relevance": "Synthetic passage selection exercises retained literature binding, not research relevance."}],
        "publication_readiness": publication_readiness()}
MANUSCRIPT_REVIEW = {**REVIEW, **{
    name: {"passed": True, "reason": "Synthetic positive manuscript decision fixture; no scholarly adequacy is attested."}
    for name in ("contribution", "literature", "interpretation", "presentation")}, "remediation": None,
    "publication_readiness": publication_readiness(manuscript=True)}
BUNDLE = {"runtime": "quickjs", "entrypoint": "experiment.mjs", "files": [
    {"path": "experiment.mjs", "content": "export default function run() { throw new Error('Unexecuted synthetic fixture'); }\n"}],
    "explanation": "Controlled experiment fixture for deterministic workflow orchestration checks."}


def collect(queries, root, *, limit, cancel):
    assert queries and limit in {3, 6}
    text = SYNTHETIC_PASSAGE
    raw = Path(root) / "literature" / "fixture-source.json"
    metadata = raw.with_name("fixture-metadata.json")
    decoded = raw.with_suffix(".txt")
    write_json(raw, {"fixture": True, "reading_text": text})
    write_json(metadata, {"title": "Synthetic oracle passage"})
    full_text = SYNTHETIC_PREFIX + text + SYNTHETIC_SUFFIX
    decoded.write_bytes(full_text.encode("utf-8"))
    source = {"id": "fixture-oracle", "scope": "full_text", "title": "Synthetic oracle passage", "authors": ["Fixture Author"],
              "simulation": True, "raw_path": "literature/fixture-source.json", "sha256": digest_file(raw),
              "metadata_path": "literature/fixture-metadata.json", "metadata_sha256": digest_file(metadata),
              "text_path": "literature/fixture-source.txt", "text_sha256": digest_file(decoded), "excerpts": [text],
              "excerpt_ranges": [{"start": SYNTHETIC_BODY_START, "end": SYNTHETIC_BODY_START + len(text), "page_start": 1, "page_end": 1}],
              "body_range": literature.full_text_body_range(full_text)}
    return {"sources": [source], "cancelled": False, "simulation": True,
            "searches": [{"query": query, "provider": "Crossref", "status": "succeeded", "attempted": True,
                          "resolved_ids": [source["id"]]} for query in queries]}


def collect_arxiv_preprint(queries, root, *, limit, cancel, arxiv_id="1311.3903v1"):
    """Synthetic collector output: identity/byte plumbing, not actual literature."""
    evidence = collect(queries, root, limit=limit, cancel=cancel)
    source = evidence["sources"][0]
    source_id = "source-" + hashlib.sha256(("arxiv:" + arxiv_id).encode()).hexdigest()[:20]
    metadata_bytes = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<feed xmlns="http://www.w3.org/2005/Atom"><entry>'
        f'<id>http://arxiv.org/abs/{arxiv_id}</id>'
        '<title>Synthetic exact-version preprint retention fixture</title>'
        '<author><name>Fixture Author</name></author>'
        '<published>2013-11-15T00:00:00Z</published>'
        '<updated>2014-02-15T00:00:00Z</updated>'
        '<summary>Synthetic identity fixture with no scholarly assessment.</summary>'
        '</entry></feed>\n').encode("utf-8")
    pdf_bytes = b"%PDF-1.4\n% Synthetic retention fixture, not an actual publication: " + arxiv_id.encode() + b"\n"
    text_bytes = (SYNTHETIC_PREFIX + SYNTHETIC_PASSAGE + SYNTHETIC_SUFFIX).encode("utf-8")
    files = {}
    for kind, suffix, content in (("metadata", "xml", metadata_bytes), ("pdf", "pdf", pdf_bytes),
                                 ("text", "txt", text_bytes), ("search", "xml", metadata_bytes)):
        digest = hashlib.sha256(content).hexdigest()
        relative = f"literature/{source_id}-{kind}-{digest[:16]}.{suffix}"
        path = root / relative
        path.write_bytes(content)
        files[kind] = (relative, digest)
    source.update(id=source_id, provider="arXiv", publication_type="preprint", arxiv_id=arxiv_id,
        url="https://arxiv.org/pdf/" + arxiv_id, title="Synthetic exact-version preprint retention fixture",
        year=2013, published="2013-11-15T00:00:00Z", updated="2014-02-15T00:00:00Z",
        metadata_path=files["metadata"][0], metadata_sha256=files["metadata"][1],
        raw_path=files["pdf"][0], sha256=files["pdf"][1],
        text_path=files["text"][0], text_sha256=files["text"][1], text_chars=len(text_bytes.decode()),
        reading_scope="Synthetic exact-version full-text passages; no real scholarly value is assessed.")
    evidence["searches"] = [{"query": query, "provider": "arXiv", "lookup": "arxiv_id",
        "requested_arxiv_id": arxiv_id, "status": "succeeded", "attempted": True,
        "raw_path": files["search"][0], "sha256": files["search"][1], "resolved_ids": [source_id]}
        for query in queries]
    return evidence


class FixtureRunner:
    def __init__(self):
        self.outputs = observations()
        self.calls = 0
        self.scientific_inputs = None
        self.entered = threading.Event()
        self.release = None
        self.production_calls = True
        self.coverage_truncated = False
        self.exit_status = "succeeded"
        self.cleanup_confirmed = True
        self.cleanable = True

    def status(self):
        return {"ready": True, "simulation": True, "backend": "quickjs-wasm", "runtimes": ["quickjs"], "dependencies": [], "declared_limits": {"guest_os_processes": 0}}

    def stop(self, handle):
        assert handle["kind"] == "fixture"
        if self.release:
            self.release.set()
        return self.cleanable

    def run(self, source_dir, bundle_dir, output_dir, *, runtime, entrypoint, production_entrypoint, scientific_inputs, timeout_seconds, cancel, on_handle):
        assert runtime == "quickjs" and timeout_seconds == 300
        assert production_entrypoint == "transform.js:transform"
        assert Path(source_dir, "transform.js").is_file() and Path(bundle_dir, entrypoint).is_file()
        self.calls += 1
        self.scientific_inputs = copy.deepcopy(scientific_inputs)
        on_handle({"kind": "fixture", "pid": 123, "simulation": True})
        self.entered.set()
        if self.release:
            while not self.release.wait(0.01):
                if cancel():
                    return {"status": "cancelled", "cleanup_confirmed": True, "simulation": True}
        write_json(Path(output_dir) / "observations.json", copy.deepcopy(self.outputs))
        return {"status": self.exit_status, "simulation": True, "exit_code": 0,
                "coverage_mechanism": "synthetic", "coverage_truncated": self.coverage_truncated,
                "cleanup_confirmed": self.cleanup_confirmed,
                "active_handle": {"kind": "fixture", "pid": 123} if not self.cleanup_confirmed else {},
                "production_calls": [{"path": "transform.js", "function": "transform", "calls": 6}] if self.production_calls else []}


def test_controller_requires_explicit_owned_home_and_runner(tmp_path):
    with pytest.raises(TypeError):
        WorkflowService(runner=FixtureRunner())
    with pytest.raises(TypeError):
        WorkflowService(tmp_path / "home")
    with pytest.raises(ValueError, match="explicit QuickJS"):
        WorkflowService(tmp_path / "home", runner=None)
    assert not (tmp_path / "home").exists()


@pytest.fixture
def setup(tmp_path):
    source = tmp_path / "input"
    source.mkdir()
    (source / "transform.js").write_text("export function transform(values) { return values.slice(); }\n", encoding="utf-8")
    runner = FixtureRunner()
    service = WorkflowService(tmp_path / "home", runner=runner, collector=collect)
    created = service.create(str(source), "Validate a bounded controlled production transformation study.")
    yield service, runner, created["id"]
    runner.cleanable = True
    if runner.release:
        runner.release.set()
    try:
        service.close()
    except WorkflowError:
        pass


def approve_study(service, research_id):
    """Simulate an accepted model decision; the fixture is not meaningful research."""
    service.submit_proposal(research_id, protocol())
    service.collect_literature(research_id)
    return service.submit_study_review(research_id, copy.deepcopy(STUDY_REVIEW))


def prepare(service, research_id):
    approve_study(service, research_id)
    service.submit_code(research_id, copy.deepcopy(BUNDLE), REVIEW)


def evidence_selection_receipt():
    prompt = 'Select retained explanatory fixtures. 원래 입력과 실제 생산 응답을 확인하세요.'
    text = json.dumps({"fixture_labels": ["synthetic-inputs.json"],
                       "reason": "Retained input annotations explain the original frozen fixture result."},
                      ensure_ascii=False)
    return {"id": "12345678-1234-1234-1234-123456789abc", "phase": "evidence-selection",
            "at": "2026-10-07T00:00:00Z", "model": "gpt-6-astra", "profileId": "owned-test-profile",
            "prompt": prompt, "promptSha256": hashlib.sha256(prompt.encode()).hexdigest(),
            "text": text, "textSha256": hashlib.sha256(text.encode()).hexdigest(), "outcome": "completed"}


def test_evidence_selection_receipt_is_original_hash_bound_and_append_only(setup):
    from paper_factory import ipc

    service, runner, research_id = setup
    initial = service.status(research_id, include_materials=False)
    completed = evidence_selection_receipt()
    started = {key: value for key, value in completed.items() if key not in {"text", "textSha256"}}
    started["outcome"] = "started"
    retained = {}
    for value in (started, completed):
        request = ipc.request(json.dumps({"id": "selection-receipt", "method": "workflow.recordInference",
                                         "params": {"researchId": research_id, "receipt": value}}).encode())
        result = service.record_inference(request["params"]["researchId"], request["params"]["receipt"])
        original = service.artifact_path(research_id, result["artifactId"])
        assert json.loads(original.read_bytes()) == value
        assert digest_file(original) == result["sha256"]
        journal = service.artifact_path(research_id, f'model-journal-{value["id"]}-{value["outcome"]}')
        assert json.loads(journal.read_bytes()) == {
            "event": "model-inference", "id": value["id"], "at": value["at"],
            "phase": "evidence-selection", "outcome": value["outcome"], "receipt_sha256": result["sha256"],
        }
        retained[original], retained[journal] = original.read_bytes(), journal.read_bytes()
        assert service.record_inference(research_id, value) == result
    for replacement in ({"profileId": "another-owned-profile"}, {"phase": "code-review"},
                        {"text": "Changed selection", "textSha256": hashlib.sha256(b"Changed selection").hexdigest()}):
        with pytest.raises(WorkflowError) as conflict:
            service.record_inference(research_id, {**completed, **replacement})
        assert conflict.value.code == "INFERENCE_EVIDENCE_CONFLICT"
    assert all(path.read_bytes() == original for path, original in retained.items())
    final = service.status(research_id, include_materials=False)
    assert (final["stage"], final["status"], final["execution_attempt"]) == ("created", "ready", 0)
    assert all(final["artifacts"][key] == value for key, value in initial["artifacts"].items())
    assert len(final["artifacts"]) == len(initial["artifacts"]) + 4 and runner.calls == 0


@pytest.mark.parametrize("defect", ["prompt_hash", "text_hash", "missing_text_hash", "missing_text",
                                    "empty_profile", "long_profile", "invalid_phase", "missing_timezone"])
def test_evidence_selection_receipt_rejects_invalid_binding_without_public_input_leak(setup, defect):
    from paper_factory import ipc

    service, runner, research_id = setup
    before = service.status(research_id, include_materials=False)
    value = evidence_selection_receipt()
    private_marker = "private-selection-profile-marker"
    if defect == "prompt_hash":
        value["prompt"] += private_marker
    elif defect == "text_hash":
        value["text"] += private_marker
    elif defect == "missing_text_hash":
        value.pop("textSha256")
    elif defect == "missing_text":
        value.pop("text")
    elif defect == "empty_profile":
        value["profileId"] = ""
    elif defect == "long_profile":
        value["profileId"] = private_marker * 8
    elif defect == "invalid_phase":
        value["phase"] = private_marker
    else:
        value["at"] = "2026-10-07T00:00:00"
    with pytest.raises(ValueError) as invalid:
        service.record_inference(research_id, value)
    public = ipc.public_error(invalid.value)
    assert public["code"] == "INVALID_ARGUMENT" and private_marker not in public["message"]
    after = service.status(research_id, include_materials=False)
    assert after["artifacts"] == before["artifacts"] and after["execution_attempt"] == 0 and runner.calls == 0


def test_evidence_selection_journal_tampering_is_detected_by_native_artifact_verification(setup):
    service, runner, research_id = setup
    value = evidence_selection_receipt()
    saved = service.record_inference(research_id, value)
    receipt_path = service.artifact_path(research_id, saved["artifactId"])
    original = receipt_path.read_bytes()
    journal_id = f'model-journal-{value["id"]}-{value["outcome"]}'
    journal_path = service.artifact_path(research_id, journal_id)
    journal = json.loads(journal_path.read_bytes())
    journal["receipt_sha256"] = "0" * 64
    write_json(journal_path, journal)
    with pytest.raises(WorkflowError) as changed:
        service.artifact_path(research_id, journal_id)
    assert changed.value.code == "ARTIFACT_CHANGED" and receipt_path.read_bytes() == original and runner.calls == 0


def test_scientific_inputs_bind_declared_source_and_exact_imported_document_bytes(setup):
    service, runner, research_id = setup
    raw = '\ufeff한글 원문\r\nSecond line.\r\n'.encode('utf-8')
    imported = service.add_evidence(research_id, [{"name": "source.txt", "contentBase64": base64.b64encode(raw).decode()}])
    document = imported["supporting_documents"][0]
    ws = service._workspace(research_id)
    source_raw = ws.path('source/transform.js').read_bytes()
    prepare(service, research_id)
    service.start_experiment(research_id)
    runner.entered.wait(2)
    assert runner.calls == 1
    assert runner.scientific_inputs == {
        'source/transform.js': {'name': 'transform.js', 'text': source_raw.decode('utf-8'),
                                'sha256': hashlib.sha256(source_raw).hexdigest()},
        document['id']: {'name': 'source.txt', 'text': raw.decode('utf-8'), 'sha256': hashlib.sha256(raw).hexdigest()},
    }


@pytest.mark.parametrize('area', ['source', 'supporting'])
def test_scientific_inputs_reject_changed_bytes_before_runner_dispatch(setup, area):
    service, runner, research_id = setup
    imported = service.add_evidence(research_id, [{"name": "source.txt", "contentBase64": base64.b64encode(b'Original text.\n').decode()}])
    ws = service._workspace(research_id)
    record = ws.get('workflow', research_id, Workflow)
    document = imported['supporting_documents'][0]
    path = ws.path('source/transform.js') if area == 'source' else ws.path(record.artifacts[document['id']].path)
    path.chmod(0o600)
    path.write_bytes(b'Changed input.\n')
    with pytest.raises(WorkflowError) as error:
        _scientific_inputs(ws, record, ResearchPlan.model_validate(protocol()))
    assert error.value.code == 'ARTIFACT_CHANGED'
    assert runner.calls == 0


def test_scientific_inputs_cannot_read_undeclared_source_paths(setup):
    service, runner, research_id = setup
    ws = service._workspace(research_id)
    record = ws.get('workflow', research_id, Workflow)
    altered = protocol() | {'source_files': ['../private.txt']}
    with pytest.raises(WorkflowError) as error:
        _scientific_inputs(ws, record, ResearchPlan.model_validate(altered))
    assert error.value.code == 'MATERIAL_NOT_DECLARED'
    assert runner.calls == 0


def test_planning_context_discloses_projection_and_only_retained_compiler_evidence(setup):
    service, _, research_id = setup
    context = service.status(research_id)["source_context"]
    marker = "Controller runtime capabilities and declared limits (not repository instructions): "
    capabilities = json.loads(context.split(marker, 1)[1])
    assert capabilities["json_projection_only"] is True
    assert "JSON.stringify" in capabilities["result_serialization"]
    assert "Map entries" in capabilities["unsupported_return_encodings"]
    compilation = capabilities["typescript_compilation_evidence"]
    assert "original/compiled SHA256" in compilation
    assert "per-file transformation_options" in compilation
    assert "shared options" in compilation
    assert all(field in compilation for field in ("transformer", "name", "version", "options"))
    assert set(capabilities["unavailable_evidence"]) == {
        "emitted JavaScript bytes", "separate pre-call syntax/builtin-probe receipt",
    }


def finished(service, research_id):
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        result = service.status(research_id)
        if result["status"] != "running":
            return result
        time.sleep(0.01)
    raise AssertionError("Synthetic workflow worker did not complete")


def test_native_submissions_analyze_actual_runner_artifacts(setup):
    service, runner, research_id = setup
    assert set(service.status(research_id)["schemas"]) == {"plan", "code", "review", "manuscript", "study_review", "manuscript_review"}
    prepare(service, research_id)
    started = service.start_experiment(research_id)
    assert started["status"] == "running" and "active_handle" not in started
    result = finished(service, research_id)
    assert result["stage"] == "analyzed" and result["status"] == "ready"
    assert result["analysis"]["results"]["error.condition_1.mean"]["value"] == 0
    assert runner.calls == 1
    result = service.submit_manuscript(research_id, manuscript(), MANUSCRIPT_REVIEW)
    assert result["stage"] == "manuscript"
    review = json.loads(service.artifact_path(research_id, "manuscript-review").read_text())
    assert review["origin"] == "native_host_submission"
    assert review["analysis_sha256"] == result["artifacts"]["analysis"]["sha256"]


def test_proposal_requires_literature_and_suitability_before_protocol_freezes(setup):
    service, runner, research_id = setup
    proposed = service.submit_proposal(research_id, protocol())
    assert proposed["stage"] == "proposed" and proposed["proposal_attempt"] == 1
    assert "proposal" in proposed["artifacts"] and "plan" not in proposed["artifacts"]
    assert "plan" not in proposed and runner.calls == 0
    with pytest.raises(WorkflowError):
        service.submit_study_review(research_id, STUDY_REVIEW)
    with pytest.raises(WorkflowError):
        service.submit_code(research_id, BUNDLE, REVIEW)
    service.collect_literature(research_id)
    accepted = service.submit_study_review(research_id, STUDY_REVIEW)
    assert accepted["stage"] == "planned" and accepted["status"] == "ready"
    assert {"proposal", "proposal-1", "plan", "study-review", "study-review-1", "selected-literature"} <= set(accepted["artifacts"])
    assert service.artifact_path(research_id, "proposal").read_bytes() == service.artifact_path(research_id, "plan").read_bytes()
    approval = json.loads(service.artifact_path(research_id, "study-review").read_bytes())
    for field, key in (("proposal_sha256", "proposal"), ("protocol_sha256", "plan"),
                       ("literature_sha256", "literature"), ("selected_literature_sha256", "selected-literature")):
        assert approval[field] == accepted["artifacts"][key]["sha256"]
    assert runner.calls == 0


def test_proposed_study_review_receives_original_goal_even_when_proposal_changes_required_design(setup):
    service, runner, research_id = setup
    goal = ('필수: 고정 corpus="original.json"; 조건 production/full_matrix_lcs/linear_space_lcs; '
            '지표 partial_attainability/context_bundle_fraction; seeds 17/29를 유지하세요.\n'
            '실제 사용자 효과를 대신 주장하지 마세요.')
    created = service.create(str(service._workspace(research_id).path("source")), goal)
    proposal = protocol()  # Deliberately substitutes conditions, metrics and seeds.
    service.submit_proposal(created["id"], proposal)
    state = service.collect_literature(created["id"])
    assert state["stage"] == "proposed" and state["goal"] == goal
    original, _ = json.JSONDecoder().raw_decode(state["instructions"].split("Original requested research goal:\n", 1)[1])
    assert original == goal and original != proposal["question"]
    assert state["proposal"]["conditions"] == ["production", "ablation"]
    assert [item["name"] for item in state["proposal"]["metrics"]] == ["error"]
    assert state["proposal"]["seeds"] == [11, 37]
    assert "silently relaxing the original goal" in state["instructions"]
    assert "Independently derive every proposed positive-control" in state["instructions"]
    assert "plan" not in state["artifacts"] and runner.calls == 0


@pytest.mark.parametrize("criterion", STUDY_CRITERIA)
def test_study_acceptance_cannot_override_a_failed_quality_criterion(setup, criterion):
    service, runner, research_id = setup
    service.submit_proposal(research_id, protocol())
    service.collect_literature(research_id)
    contradictory = copy.deepcopy(STUDY_REVIEW)
    contradictory[criterion]["passed"] = False
    with pytest.raises(ValueError, match="acceptance"):
        service.submit_study_review(research_id, contradictory)
    state = service.status(research_id)
    assert state["stage"] == "proposed" and "plan" not in state["artifacts"]
    assert "study-review" not in state["artifacts"] and runner.calls == 0


def rejected_study_review():
    review = copy.deepcopy(STUDY_REVIEW)
    review["accepted"] = False
    review["contribution"]["passed"] = False
    review["contribution"]["reason"] = "This is a trivial contract check with no justified contribution beyond existing behavior."
    review["issues"] = ["A routine regression fixture does not establish a worthwhile academic contribution."]
    review["selected_sources"] = []
    return review


def test_weak_study_remains_unapproved_when_resumed_for_bounded_evidence(setup):
    service, runner, research_id = setup
    service.submit_proposal(research_id, protocol())
    service.collect_literature(research_id)
    rejected = service.submit_study_review(research_id, rejected_study_review())
    assert rejected["status"] == "blocked" and rejected["code"] == "STUDY_REJECTED"
    assert rejected["stage"] == "proposed" and rejected["resume_kind"] == "preparation"
    assert "plan" not in rejected["artifacts"]
    assert not {"bundle", "observations", "analysis", "manuscript", "export-pdf"} & set(rejected["artifacts"])
    for operation in (lambda: service.submit_code(research_id, BUNDLE, REVIEW),
                      lambda: service.start_experiment(research_id), lambda: service.export(research_id)):
        with pytest.raises(WorkflowError):
            operation()
    assert runner.calls == 0
    receipt = json.loads(service.artifact_path(research_id, "study-review-1").read_bytes())
    assert receipt["review"] == rejected_study_review()
    resumed = service.resume(research_id)
    assert resumed["stage"] == "proposed" and resumed["study_review"]["accepted"] is False
    assert "plan" not in resumed["artifacts"] and runner.calls == 0
    cancelled = service.cancel(research_id)
    assert cancelled["resume_kind"] == "preparation"
    assert service.resume(research_id)["study_review"]["accepted"] is False
    assert json.loads(service.artifact_path(research_id, "study-review-1").read_bytes()) == receipt
    assert runner.calls == 0


def test_rejected_proposals_can_be_improved_at_most_three_times_without_erasing_history(setup):
    service, runner, research_id = setup
    preserved = {}
    for attempt in range(1, 4):
        proposal = protocol()
        proposal["expected_contribution"] += f" Synthetic revision label {attempt}."
        state = service.submit_proposal(research_id, proposal)
        assert state["proposal_attempt"] == attempt and state["status"] == "ready"
        service.collect_literature(research_id)
        service.submit_study_review(research_id, rejected_study_review())
        for key in (f"proposal-{attempt}", f"study-review-{attempt}"):
            preserved[key] = service.artifact_path(research_id, key).read_bytes()
        assert all(service.artifact_path(research_id, key).read_bytes() == data for key, data in preserved.items())
    with pytest.raises(WorkflowError) as blocked:
        service.submit_proposal(research_id, protocol())
    assert blocked.value.code == "PROPOSAL_LIMIT"
    state = service.status(research_id)
    assert state["proposal_attempt"] == 3 and state["code"] == "STUDY_REJECTED"
    assert all(service.artifact_path(research_id, key).read_bytes() == data for key, data in preserved.items())
    assert "plan" not in state["artifacts"] and runner.calls == 0


def test_rejected_normalized_proposal_cannot_repeat_or_duplicate_instrumentation(setup):
    service, runner, research_id = setup
    service.submit_proposal(research_id, protocol())
    normalized = json.loads(service.artifact_path(research_id, "proposal").read_bytes())
    instrumentation = "Production-call instrumentation affects execution overhead; measurements cannot establish uninstrumented production performance."
    assert sum(text.count(instrumentation) for text in normalized["limitations"]) == 1
    service.collect_literature(research_id)
    service.submit_study_review(research_id, rejected_study_review())
    frozen = {key: service.artifact_path(research_id, key).read_bytes()
              for key in ("proposal-1", "study-review-1", "literature")}
    for unchanged in (copy.deepcopy(normalized), protocol()):
        with pytest.raises(WorkflowError) as rejected:
            service.submit_proposal(research_id, unchanged)
        assert rejected.value.code == "PROPOSAL_UNCHANGED"
        assert service.status(research_id)["proposal_attempt"] == 1
        assert all(service.artifact_path(research_id, key).read_bytes() == content for key, content in frozen.items())
    revised = copy.deepcopy(normalized)
    revised["expected_contribution"] += " A revised synthetic design must distinguish annotated migration cases."
    proposed = service.submit_proposal(research_id, revised)
    assert proposed["proposal_attempt"] == 2 and proposed["stage"] == "proposed"
    assert sum(text.count(instrumentation) for text in proposed["proposal"]["limitations"]) == 1
    assert all(service.artifact_path(research_id, key).read_bytes() == content
               for key, content in frozen.items() if key != "literature")
    assert runner.calls == 0


def test_revised_proposal_gets_fresh_literature_budget_and_retries_old_failed_queries(setup):
    service, runner, research_id = setup
    proposal = protocol()
    proposal["literature_queries"] = ["previously failed query", "irrelevant initial topic"]
    service.submit_proposal(research_id, proposal)
    calls = []

    def revised_search(queries, root, *, limit, cancel):
        calls.append((list(queries), limit))
        evidence = collect(queries, root, limit=limit, cancel=cancel)
        source = evidence["sources"][0]
        if len(calls) == 1:
            evidence["sources"] = [{**source, "id": f"irrelevant-{index}",
                                    "title": f"Unrelated synthetic reading {index}"} for index in range(6)]
            evidence["searches"][0].update(status="failed", resolved_ids=[], error="HTTPStatusError", http_status=429)
        else:
            evidence["sources"].append({**source, "id": "second-fixture-oracle"})
        return evidence

    service.collector = revised_search
    initial = service.collect_literature(research_id)
    assert len(initial["literature"]["sources"]) == 6
    rejected = rejected_study_review()
    rejected["literature"] = {"passed": False, "reason": "All six retained readings are unrelated to the original research question."}
    service.submit_study_review(research_id, rejected)
    old_path = service.artifact_path(research_id, "literature")
    old_bytes = old_path.read_bytes()
    old_digest = digest_file(old_path)
    prior = service.status(research_id)
    retained = {key: service.artifact_path(research_id, key).read_bytes()
                for key in prior["artifacts"] if key not in {"proposal", "study-review", "literature"}}

    revised = copy.deepcopy(proposal)
    revised["literature_queries"] = ["previously failed query", "refined relevant method query"]
    proposed = service.submit_proposal(research_id, revised)
    history_key = "literature-history-" + old_digest
    assert proposed["proposal_attempt"] == 2 and proposed["stage"] == "proposed"
    assert "literature" not in proposed["artifacts"] and "literature" not in proposed
    assert service.artifact_path(research_id, history_key) == old_path
    assert proposed["artifacts"][history_key]["sha256"] == old_digest
    assert proposed["study_review"] is None and "plan" not in proposed["artifacts"]
    with pytest.raises(WorkflowError) as missing:
        service.submit_study_review(research_id, STUDY_REVIEW)
    assert missing.value.code == "ARTIFACT_MISSING"

    completed = service.collect_literature(research_id)
    assert calls == [(proposal["literature_queries"], 6), (revised["literature_queries"], 6)]
    assert len(completed["literature"]["sources"]) == 2
    assert [search["query"] for search in completed["literature"]["searches"]] == revised["literature_queries"]
    assert all(search["status"] == "succeeded" for search in completed["literature"]["searches"])
    assert service.artifact_path(research_id, "literature") != old_path and old_path.read_bytes() == old_bytes
    assert all(service.artifact_path(research_id, key).read_bytes() == content for key, content in retained.items())
    fresh_review = copy.deepcopy(STUDY_REVIEW)
    fresh_review["selected_sources"].append({**fresh_review["selected_sources"][0], "source_id": "second-fixture-oracle"})
    approved = service.submit_study_review(research_id, fresh_review)
    assert approved["stage"] == "planned" and runner.calls == 0
    assert len(approved["literature"]["sources"]) == 2
    assert old_path.read_bytes() == old_bytes


@pytest.mark.parametrize("defect", ["unknown", "metadata", "out_of_bounds"])
def test_suitability_selection_must_bind_a_retrieved_readable_source_excerpt(setup, defect):
    service, runner, research_id = setup
    service.submit_proposal(research_id, protocol())
    if defect == "metadata":
        def metadata_collector(*args, **kwargs):
            evidence = collect(*args, **kwargs)
            evidence["sources"][0].update(scope="metadata_only", excerpts=[])
            return evidence
        service.collector = metadata_collector
        with pytest.raises(WorkflowError, match="Metadata"):
            service.collect_literature(research_id)
    else:
        service.collect_literature(research_id)
    review = copy.deepcopy(STUDY_REVIEW)
    if defect == "unknown":
        review["selected_sources"][0]["source_id"] = "not-retrieved"
    elif defect == "out_of_bounds":
        review["selected_sources"][0]["excerpt_index"] = 1
    with pytest.raises(WorkflowError) as rejected:
        service.submit_study_review(research_id, review)
    assert rejected.value.code == "LITERATURE_SELECTION_INVALID"
    assert "plan" not in service.status(research_id)["artifacts"] and runner.calls == 0


def test_only_selected_literature_is_available_to_writer_and_citation_validation(setup):
    service, runner, research_id = setup
    def two_sources(*args, **kwargs):
        evidence = collect(*args, **kwargs)
        first, second = evidence["sources"][0]["excerpts"][0].split(". ", 1)
        evidence["sources"][0]["excerpts"] = [first + ".", second]
        evidence["sources"][0]["excerpt_ranges"] = [{"start": SYNTHETIC_BODY_START, "end": SYNTHETIC_BODY_START + len(first) + 1},
                                                    {"start": SYNTHETIC_BODY_START + len(first) + 2, "end": SYNTHETIC_BODY_START + len(SYNTHETIC_PASSAGE)}]
        evidence["sources"].append({**evidence["sources"][0], "id": "unselected-source",
                                   "title": "Unrelated source that was retrieved but not selected"})
        return evidence
    service.collector = two_sources
    service.submit_proposal(research_id, protocol())
    service.collect_literature(research_id)
    raw_path = service.artifact_path(research_id, "literature")
    raw_bytes = raw_path.read_bytes()
    review = copy.deepcopy(STUDY_REVIEW)
    review["selected_sources"][0]["excerpt_index"] = 1
    work = review["publication_readiness"]["closest_work"][0]
    work.update(excerpt_index=1, quote=json.loads(raw_bytes)["sources"][0]["excerpts"][1][:100])
    service.submit_study_review(research_id, review)
    service.submit_code(research_id, BUNDLE, REVIEW)
    state = service.status(research_id)
    assert [source["id"] for source in state["literature"]["sources"]] == ["fixture-oracle"]
    raw = json.loads(raw_bytes)
    assert raw_path.read_bytes() == raw_bytes
    assert {source["id"] for source in raw["sources"]} == {"fixture-oracle", "unselected-source"}
    selected_source = state["literature"]["sources"][0]
    assert selected_source["excerpts"] == [raw["sources"][0]["excerpts"][1]]
    assert selected_source["selected_excerpt_index"] == 1
    service.start_experiment(research_id)
    analyzed = finished(service, research_id)
    assert "unselected-source" not in analyzed["instructions"]
    assert raw["sources"][0]["excerpts"][0] not in analyzed["instructions"]
    bad = manuscript()
    related = next(section for section in bad["sections"] if section["heading"] == "Related Work")
    related["text"] = related["text"].replace("{{citation:fixture-oracle}}", "{{citation:unselected-source}}")
    manuscript_review = copy.deepcopy(MANUSCRIPT_REVIEW)
    manuscript_review["publication_readiness"]["closest_work"] = review["publication_readiness"]["closest_work"]
    with pytest.raises(ValueError, match="cite its inspected closest work"):
        service.submit_manuscript(research_id, bad, manuscript_review)
    assert service.status(research_id)["stage"] == "analyzed" and runner.calls == 1


@pytest.mark.parametrize("criterion", ["contribution", "literature", "interpretation", "presentation"])
def test_manuscript_acceptance_requires_all_quality_criteria(setup, criterion):
    service, runner, research_id = setup
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert finished(service, research_id)["stage"] == "analyzed"
    review = copy.deepcopy(MANUSCRIPT_REVIEW)
    review[criterion]["passed"] = False
    with pytest.raises(ValueError, match="acceptance"):
        service.submit_manuscript(research_id, manuscript(), review)
    assert "manuscript" not in service.status(research_id)["artifacts"] and runner.calls == 1


def test_rejected_manuscript_retains_decision_and_draft_without_export_or_new_experiment(setup):
    service, runner, research_id = setup
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert finished(service, research_id)["stage"] == "analyzed"
    rejected_review = copy.deepcopy(MANUSCRIPT_REVIEW)
    rejected_review["accepted"] = False
    rejected_review["contribution"]["passed"] = False
    rejected_review["contribution"]["reason"] = "The draft turns a trivial contract check into a paper without useful new knowledge."
    rejected_review["issues"] = ["The evidence does not justify the manuscript's asserted research contribution."]
    rejected_review["remediation"] = {"strategy": "revise_manuscript", "reason": "The retained synthetic observations support only a bounded contract statement.",
                                      "actions": [{"criterion": "contribution", "action": "Remove the unsupported novelty claim and state the synthetic contract scope."}],
                                      "evidence_gaps": []}
    rejected = service.submit_manuscript(research_id, manuscript(), rejected_review)
    assert rejected["stage"] == "analyzed" and rejected["status"] == "blocked"
    assert rejected["code"] == "MANUSCRIPT_REJECTED" and rejected["manuscript_review"] == rejected_review
    assert "manuscript" not in rejected["artifacts"] and "export-pdf" not in rejected["artifacts"]
    frozen = {key: service.artifact_path(research_id, key).read_bytes()
              for key in ("draft-1", "manuscript-review-1", "observations", "analysis")}
    with pytest.raises(WorkflowError):
        service.export(research_id)
    accepted = service.submit_manuscript(research_id, manuscript(), MANUSCRIPT_REVIEW)
    assert accepted["stage"] == "manuscript" and accepted["status"] == "ready" and accepted["code"] is None
    assert accepted["draft_attempt"] == 2 and accepted["manuscript_review"]["accepted"] is True
    assert all(service.artifact_path(research_id, key).read_bytes() == content for key, content in frozen.items())
    assert runner.calls == 1


def redesign_review(strategy="redesign_study"):
    review = copy.deepcopy(MANUSCRIPT_REVIEW)
    review.update(accepted=False, issues=["Synthetic evidence lacks the required independent comparison."])
    review["contribution"] = {"passed": False, "reason": "Synthetic evidence cannot support the declared scientific contribution."}
    review["remediation"] = {"strategy": strategy,
                             "reason": "A new preregistered comparison is required to address the missing independent evidence.",
                             "actions": [{"criterion": "contribution", "action": "Design an independently justified comparison without changing prior observed results."}],
                             "evidence_gaps": [] if strategy == "revise_manuscript" else ["The frozen study has no independent validation of its asserted general application."]}
    return review


def reject_for_redesign(service, research_id):
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert finished(service, research_id)["stage"] == "analyzed"
    return service.submit_manuscript(research_id, manuscript(), redesign_review())


def committed_workflows(service):
    return [record for path in service.root.iterdir() if path.is_dir() and (path / "records.sqlite3").is_file()
            for record in Workspace(path).list("workflow", Workflow)]


def test_study_redesign_reuses_exact_source_provenance_and_supporting_bytes_without_dispatch(setup, monkeypatch):
    from paper_factory import project as project_module
    from paper_factory.models import Project
    service, runner, research_id = setup
    raw = b"Original imported corpus with CRLF.\r\n"
    imported_docs = service.add_evidence(research_id, [{"name": "corpus.txt", "contentBase64": base64.b64encode(raw).decode()}])
    parent_ws = service._workspace(research_id)
    original_project = parent_ws.latest("project", Project)
    original_project.source_commit = "a" * 40
    original_project.source = "https://github.com/example/synthetic"
    parent_ws.save("project", original_project)
    write_json(parent_ws.path("project.json"), original_project)
    parent = parent_ws.get("workflow", research_id, Workflow)
    collection = parent_ws.path("source-collection.json")
    write_json(collection, {"commit": original_project.source_commit, "scope": "Synthetic provenance fixture; no network retrieval"})
    parent.artifacts["source-collection"] = FrozenArtifact(path="source-collection.json", sha256=digest_file(collection), size=collection.stat().st_size)
    parent_ws.save("workflow", parent)
    rejected = reject_for_redesign(service, research_id)
    before = {key: service.artifact_path(research_id, key).read_bytes() for key in rejected["artifacts"]}
    monkeypatch.setattr(project_module, "ingest", lambda *args: pytest.fail("Redesign downloaded or re-imported moving source"))
    child = service.redesign_study(research_id)
    assert child["id"] != research_id and child["parent_research_id"] == research_id
    assert child["root_research_id"] == research_id and child["redesign_attempt"] == 1
    assert (child["stage"], child["status"], child["execution_attempt"], child["proposal_attempt"], child["draft_attempt"]) == ("created", "ready", 0, 0, 0)
    assert child["study_review"] is None and child["manuscript_review"] is None
    assert not any(key in child["artifacts"] for key in ("plan", "bundle", "observations", "analysis", "literature", "execution"))
    child_ws = service._workspace(child["id"])
    assert child_ws.latest("project", Project) == original_project
    assert child_ws.path("project.json").read_bytes() == parent_ws.path("project.json").read_bytes()
    assert child_ws.path("source-collection.json").read_bytes() == collection.read_bytes()
    assert child_ws.path("source/transform.js").read_bytes() == parent_ws.path("source/transform.js").read_bytes()
    assert child["supporting_documents"] == imported_docs["supporting_documents"]
    document = child["supporting_documents"][0]
    assert service.artifact_path(child["id"], document["id"]).read_bytes() == raw
    assert all(service.artifact_path(research_id, key).read_bytes() == data for key, data in before.items())
    assert child["prior_study"]["analysis"] == rejected["analysis"]
    assert child["prior_study"]["review"] == rejected["manuscript_review"]
    assert child["prior_study"]["execution"]["status"] == "succeeded"
    assert "exploratory history" in child["instructions"] and "Do not repeat the same design" in child["instructions"]
    full_history = json.dumps(child["prior_study"], ensure_ascii=False)
    assert full_history in child["instructions"] and full_history not in child["planning_instructions"]
    assert "exploratory history" in child["planning_instructions"]
    assert service.status(research_id)["followup_research_id"] == child["id"]
    assert runner.calls == 1
    with pytest.raises(WorkflowError):
        service.start_experiment(child["id"])
    assert service.redesign_study(research_id)["id"] == child["id"]
    assert len(committed_workflows(service)) == 2 and runner.calls == 1


def test_study_redesign_lost_reply_reads_active_child_without_duplicate_dispatch(setup):
    service, runner, research_id = setup
    reject_for_redesign(service, research_id)
    child = service.redesign_study(research_id)
    prepare(service, child["id"])
    runner.release = threading.Event(); runner.entered.clear()
    try:
        started = service.start_experiment(child["id"])
        assert started["status"] == "running" and runner.entered.wait(2)
        retained = service.redesign_study(research_id)
        assert retained["id"] == child["id"] and retained["status"] == "running"
        assert retained["execution_attempt"] == 1 and runner.calls == 2
        assert len(committed_workflows(service)) == 2
    finally:
        runner.release.set()
        assert finished(service, child["id"])["stage"] == "analyzed"


def test_study_redesign_recovers_same_intent_after_partial_copy_and_restart(setup, monkeypatch):
    import paper_factory.workflow as workflow_module
    service, runner, research_id = setup
    reject_for_redesign(service, research_id)
    copy_file = workflow_module._copy_retained
    counter = 0
    def interrupted_copy(*args):
        nonlocal counter
        counter += 1
        if counter == 3:
            raise OSError("Synthetic import interruption after a durable parent intent")
        return copy_file(*args)
    with monkeypatch.context() as scope:
        scope.setattr(workflow_module, "_copy_retained", interrupted_copy)
        with pytest.raises(OSError, match="Synthetic import interruption"):
            service.redesign_study(research_id)
    pending = service.status(research_id)
    intent_id = json.loads(service.artifact_path(research_id, "redesign-intent").read_bytes())["child_id"]
    assert pending["followup_research_id"] is None and pending["redesign_pending"] is True
    assert pending["improvement_available"] is False and len(committed_workflows(service)) == 1 and runner.calls == 1
    with pytest.raises(WorkflowError) as refused:
        service.improve_writing(research_id)
    assert refused.value.code == "WRITING_IMPROVEMENT_NOT_ALLOWED"
    service.close()
    reopened = WorkflowService(service.home, runner=runner, collector=collect)
    try:
        child = reopened.redesign_study(research_id)
        assert child["id"] == intent_id and child["execution_attempt"] == 0 and runner.calls == 1
        assert reopened.status(research_id)["followup_research_id"] == intent_id
        assert reopened.status(research_id)["redesign_pending"] is False
        assert reopened.redesign_study(research_id)["id"] == intent_id
        assert len(committed_workflows(reopened)) == 2
    finally:
        reopened.close()


@pytest.mark.parametrize("failure", ["control", "active", "review-binding"])
def test_pending_redesign_marker_requires_safe_current_science_and_bound_review(setup, monkeypatch, failure):
    import paper_factory.workflow as workflow_module
    service, runner, research_id = setup
    reject_for_redesign(service, research_id)
    with monkeypatch.context() as scope:
        scope.setattr(workflow_module, "_copy_retained", lambda *args: (_ for _ in ()).throw(OSError("Synthetic pending clone")))
        with pytest.raises(OSError, match="Synthetic pending clone"):
            service.redesign_study(research_id)
    ws = service._workspace(research_id)
    record = ws.get("workflow", research_id, Workflow)
    if failure == "control":
        record.terminal_control_failure = True
    elif failure == "active":
        record.active_handle = {"kind": "fixture", "pid": 123}
    else:
        path = service.artifact_path(research_id, "manuscript-review")
        receipt = json.loads(path.read_bytes()); receipt["analysis_sha256"] = "0" * 64
        write_json(path, receipt)
        record.artifacts["manuscript-review"] = record.artifacts["manuscript-review-1"] = FrozenArtifact(
            path=record.artifacts["manuscript-review"].path, sha256=digest_file(path), size=path.stat().st_size)
    ws.save("workflow", record)
    state = service.status(research_id, include_materials=False)
    assert state["followup_research_id"] is None and state["redesign_pending"] is False
    assert state["improvement_available"] is False
    with pytest.raises((ValueError, WorkflowError)):
        service.redesign_study(research_id)
    assert runner.calls == 1 and len(committed_workflows(service)) == 1


def test_study_redesign_is_bounded_and_keeps_all_prior_negative_evidence(setup):
    service, runner, research_id = setup
    current = research_id
    chain = []
    for index in range(3):
        rejected = reject_for_redesign(service, current)
        chain.append(current)
        assert rejected["redesign_attempt"] == index and rejected["execution_attempt"] == 1
        if index < 2:
            current = service.redesign_study(current)["id"]
    assert runner.calls == 3 and len(committed_workflows(service)) == 3
    with pytest.raises(WorkflowError) as limited:
        service.redesign_study(current)
    assert limited.value.code == "STUDY_REDESIGN_LIMIT"
    final = service.status(current)
    assert final["root_research_id"] == research_id and final["prior_study"]["earlier_study"]["parent_id"] == research_id
    for prior_id in chain[:-1]:
        retained = "prior-study-" + prior_id + "-observations"
        assert service.artifact_path(current, retained).read_bytes() == service.artifact_path(prior_id, "observations").read_bytes()
    assert runner.calls == 3


def test_study_redesign_archive_keeps_prior_raw_distinct_from_new_observations(setup):
    service, runner, research_id = setup
    reject_for_redesign(service, research_id)
    child = service.redesign_study(research_id)
    runner.outputs = observations(); runner.outputs["observations"][0]["value"] = 9
    prepare(service, child["id"]); service.start_experiment(child["id"])
    assert finished(service, child["id"])["stage"] == "analyzed"
    service.submit_manuscript(child["id"], manuscript(), MANUSCRIPT_REVIEW)
    ws = service._workspace(child["id"])
    record = ws.get("workflow", child["id"], Workflow)
    root = ws.path("research/exports/synthetic-lineage-archive")
    root.mkdir(parents=True)
    # These native-format placeholders test archive membership only, never export quality.
    for key, name in (("export-md", "paper.md"), ("export-pdf", "paper.pdf"), ("export-docx", "paper.docx"),
                      ("export-tex", "paper.tex"), ("conversion", "conversion-receipts.json")):
        path = root / name; path.write_bytes(("Synthetic archive fixture only: " + key).encode())
        record.artifacts[key] = FrozenArtifact(path=path.relative_to(ws.root).as_posix(), sha256=digest_file(path), size=path.stat().st_size)
    ws.save("workflow", record)
    service._bundle(ws, record, root)
    with zipfile.ZipFile(ws.path(record.artifacts["reproducibility"].path)) as archive:
        current = archive.read("observations.json")
        prior_artifact = record.artifacts["prior-study-" + research_id + "-observations"]
        prior_name = "prior-studies/" + ws.path(prior_artifact.path).relative_to(ws.path("research/prior-studies")).as_posix()
        prior = archive.read(prior_name)
        assert current == service.artifact_path(child["id"], "observations").read_bytes()
        assert prior == service.artifact_path(research_id, "observations").read_bytes() and current != prior
        assert json.loads(archive.read("redesign-origin.json"))["parent_id"] == research_id
        assert "not current observations" in archive.read("README.md").decode()
        inventory = json.loads(archive.read("inventory.json"))
        for name in ("observations.json", prior_name, "redesign-origin.json"):
            raw = archive.read(name)
            assert inventory[name] == {"size": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}
        assert archive.testzip() is None and len(archive.namelist()) == len(set(archive.namelist()))
    assert runner.calls == 2


@pytest.mark.parametrize("field", ["protocol", "analysis", "execution", "review", "earlier_study", "source_commit", "snapshot_digest", "supporting_documents", "scope"])
def test_study_redesign_summary_cannot_forge_or_hide_retained_evidence(setup, field):
    service, runner, research_id = setup
    reject_for_redesign(service, research_id)
    child = service.redesign_study(research_id)
    ws = service._workspace(child["id"])
    record = ws.get("workflow", child["id"], Workflow)
    path = service.artifact_path(child["id"], "prior-study")
    summary = json.loads(path.read_bytes())
    summary[field] = {"fabricated": "Synthetic altered prior evidence, not an actual observation"}
    write_json(path, summary)
    changed = FrozenArtifact(path=record.artifacts["prior-study"].path, sha256=digest_file(path), size=path.stat().st_size)
    record.artifacts["prior-study"] = record.artifacts["prior-study-" + research_id + "-summary"] = changed
    ws.save("workflow", record)
    with pytest.raises(WorkflowError) as rejected:
        service.status(child["id"])
    assert rejected.value.code == "ARTIFACT_CHANGED" and runner.calls == 1


@pytest.mark.parametrize("strategy", ["revise_manuscript", "infeasible", None])
def test_study_redesign_requires_current_explicit_new_evidence_remediation(setup, strategy):
    service, runner, research_id = setup
    rejected = reject_for_redesign(service, research_id)
    ws = service._workspace(research_id)
    record = ws.get("workflow", research_id, Workflow)
    artifact = record.artifacts["manuscript-review"]
    receipt = json.loads(ws.path(artifact.path).read_bytes())
    receipt["review"]["remediation"] = None if strategy is None else redesign_review(strategy)["remediation"]
    write_json(ws.path(artifact.path), receipt)
    updated = FrozenArtifact(path=artifact.path, sha256=digest_file(ws.path(artifact.path)), size=ws.path(artifact.path).stat().st_size)
    record.artifacts["manuscript-review"] = record.artifacts["manuscript-review-1"] = updated
    ws.save("workflow", record)
    assert service.status(research_id)["manuscript_review"]["remediation"] == receipt["review"]["remediation"]
    with pytest.raises(WorkflowError) as rejected:
        service.redesign_study(research_id)
    assert rejected.value.code == "STUDY_REDESIGN_NOT_ALLOWED" and runner.calls == 1
    assert len(committed_workflows(service)) == 1


@pytest.mark.parametrize("failure", ["source", "observations", "review-binding", "control", "cleanup", "execution", "cancelled", "lineage"])
def test_study_redesign_refuses_tampered_failed_ambiguous_or_cancelled_science(setup, failure):
    service, runner, research_id = setup
    reject_for_redesign(service, research_id)
    ws = service._workspace(research_id)
    record = ws.get("workflow", research_id, Workflow)
    if failure in {"source", "observations"}:
        path = ws.path("source/transform.js") if failure == "source" else service.artifact_path(research_id, "observations")
        path.chmod(0o600)
        path.write_bytes(b"Changed retained scientific bytes.")
    elif failure == "review-binding":
        path = service.artifact_path(research_id, "manuscript-review")
        receipt = json.loads(path.read_bytes()); receipt["analysis_sha256"] = "0" * 64
        write_json(path, receipt)
        record.artifacts["manuscript-review"] = record.artifacts["manuscript-review-1"] = FrozenArtifact(path=record.artifacts["manuscript-review"].path, sha256=digest_file(path), size=path.stat().st_size)
    elif failure == "control":
        record.terminal_control_failure = True
    elif failure == "cleanup":
        record.active_handle = {"kind": "fixture", "pid": 123}
        record.code = "CLEANUP_UNCONFIRMED"
    elif failure == "execution":
        record.execution_attempt = 2
    elif failure == "cancelled":
        record.cancellation_requested = True
    elif failure == "lineage":
        child = service.redesign_study(research_id)
        current_ws = service._workspace(child["id"])
        current = current_ws.get("workflow", child["id"], Workflow)
        altered = current.model_dump(); altered.update(parent_research_id=None, root_research_id=None, redesign_attempt=0)
        current_ws.save("workflow", Workflow.model_validate(altered))
        with pytest.raises(WorkflowError, match="discard retained redesign lineage"):
            service.status(child["id"])
        assert runner.calls == 1
        return
    ws.save("workflow", record)
    with pytest.raises((ValueError, WorkflowError)):
        service.redesign_study(research_id)
    assert len(committed_workflows(service)) == 1 and runner.calls == 1


def test_new_rejected_manuscript_requires_actionable_remediation_before_retention(setup):
    service, runner, research_id = setup
    prepare(service, research_id); service.start_experiment(research_id)
    assert finished(service, research_id)["stage"] == "analyzed"
    review = redesign_review(); review["remediation"] = None
    with pytest.raises(WorkflowError) as rejected:
        service.submit_manuscript(research_id, manuscript(), review)
    assert rejected.value.code == "REVIEW_REMEDIATION_REQUIRED"
    assert service.status(research_id)["draft_attempt"] == 0 and runner.calls == 1


def test_explicit_writing_improvement_reopens_old_rejection_without_touching_frozen_science(setup):
    service, runner, research_id = setup
    reject_for_redesign(service, research_id)
    ws = service._workspace(research_id)
    record = ws.get("workflow", research_id, Workflow)
    path = service.artifact_path(research_id, "manuscript-review")
    receipt = json.loads(path.read_bytes()); receipt["review"]["remediation"] = None
    write_json(path, receipt)
    record.artifacts["manuscript-review"] = record.artifacts["manuscript-review-1"] = FrozenArtifact(
        path=record.artifacts["manuscript-review"].path, sha256=digest_file(path), size=path.stat().st_size)
    ws.save("workflow", record)
    held = service.status(research_id)
    assert held["improvement_available"] is True and held["manuscript_review"]["remediation"] is None
    before = ws.get("workflow", research_id, Workflow).model_dump(mode="json")
    frozen = {key: service.artifact_path(research_id, key).read_bytes() for key in held["artifacts"]}
    reopened = service.improve_writing(research_id)
    assert (reopened["stage"], reopened["status"], reopened["code"], reopened["execution_attempt"], reopened["draft_attempt"]) == ("analyzed", "ready", None, 1, 1)
    assert reopened["resume_kind"] == "authoring" and reopened["improvement_available"] is False
    assert reopened["manuscript_review"]["remediation"] is None
    identifier = next(key for key in reopened["artifacts"] if key.startswith("writing-improvement-"))
    evidence = json.loads(service.artifact_path(research_id, identifier).read_bytes())
    assert evidence["event"] == "writing-improvement" and evidence["previous_workflow"] == before
    assert evidence["execution_attempt"] == 1
    for key in ("plan", "execution", "observations", "analysis", "manuscript-review"):
        assert evidence[key + "_sha256"] == before["artifacts"][key]["sha256"]
    assert identifier in {item["name"] for item in reopened["material_manifest"]["evidence"]}
    assert all(service.artifact_path(research_id, key).read_bytes() == raw for key, raw in frozen.items())
    assert runner.calls == 1
    with pytest.raises(WorkflowError) as refused:
        service.improve_writing(research_id)
    assert refused.value.code == "WRITING_IMPROVEMENT_NOT_ALLOWED"
    with pytest.raises(WorkflowError) as refused:
        service.start_experiment(research_id)
    assert refused.value.code == "EXPERIMENT_ALREADY_DISPATCHED" and runner.calls == 1
    accepted = service.submit_manuscript(research_id, manuscript(), MANUSCRIPT_REVIEW)
    assert accepted["stage"] == "manuscript" and accepted["draft_attempt"] == 2 and runner.calls == 1


@pytest.mark.parametrize("failure", ["infeasible", "active", "control", "failed", "cancelled", "execution", "source", "binding", "unsuccessful", "cleanup"])
def test_writing_improvement_availability_refuses_unsafe_or_unsupported_studies(setup, failure):
    service, runner, research_id = setup
    reject_for_redesign(service, research_id)
    ws = service._workspace(research_id)
    record = ws.get("workflow", research_id, Workflow)
    if failure in {"infeasible", "binding"}:
        path = service.artifact_path(research_id, "manuscript-review")
        receipt = json.loads(path.read_bytes())
        if failure == "infeasible":
            receipt["review"]["remediation"] = redesign_review("infeasible")["remediation"]
        else:
            receipt["analysis_sha256"] = "0" * 64
        write_json(path, receipt)
        record.artifacts["manuscript-review"] = record.artifacts["manuscript-review-1"] = FrozenArtifact(
            path=record.artifacts["manuscript-review"].path, sha256=digest_file(path), size=path.stat().st_size)
    elif failure == "active":
        record.active_handle = {"kind": "fixture", "pid": 123}
    elif failure == "control":
        record.terminal_control_failure = True
    elif failure in {"failed", "cancelled"}:
        record.status = failure
    elif failure == "execution":
        record.execution_attempt = 2
    elif failure == "source":
        source = ws.path("source/transform.js"); source.chmod(0o600); source.write_bytes(b"Modified production source.")
    elif failure in {"unsuccessful", "cleanup"}:
        path = service.artifact_path(research_id, "execution")
        receipt = json.loads(path.read_bytes())
        receipt["status" if failure == "unsuccessful" else "cleanup_confirmed"] = "failed" if failure == "unsuccessful" else False
        write_json(path, receipt)
        record.artifacts["execution"] = FrozenArtifact(path=record.artifacts["execution"].path, sha256=digest_file(path), size=path.stat().st_size)
    ws.save("workflow", record)
    assert service.status(research_id, include_materials=False)["improvement_available"] is False
    with pytest.raises(WorkflowError) as refused:
        service.improve_writing(research_id)
    assert refused.value.code == "WRITING_IMPROVEMENT_NOT_ALLOWED" and runner.calls == 1
    assert not any(key.startswith("writing-improvement-") for key in service.status(research_id, include_materials=False)["artifacts"])


def test_writing_improvement_is_unavailable_after_committed_followup_and_redesign_budget(setup):
    service, runner, research_id = setup
    reject_for_redesign(service, research_id)
    assert service.status(research_id)["improvement_available"] is True
    first = service.redesign_study(research_id)
    assert service.status(research_id)["improvement_available"] is False
    with pytest.raises(WorkflowError):
        service.improve_writing(research_id)
    reject_for_redesign(service, first["id"])
    second = service.redesign_study(first["id"])
    reject_for_redesign(service, second["id"])
    assert service.status(second["id"])["improvement_available"] is False
    with pytest.raises(WorkflowError):
        service.improve_writing(second["id"])
    assert runner.calls == 3


@pytest.mark.parametrize("target", ["proposal", "study-review", "literature", "selected-literature"])
def test_study_binding_cannot_be_bypassed_by_rehashing_a_changed_artifact(setup, target):
    service, runner, research_id = setup
    prepare(service, research_id)
    ws = Workspace(service.root / research_id)
    record = ws.get("workflow", research_id, Workflow)
    artifact = record.artifacts[target]
    path = ws.path(artifact.path)
    content = json.loads(path.read_bytes())
    if target == "proposal":
        content["expected_contribution"] += " Altered after research review."
    elif target == "study-review":
        content["literature_sha256"] = "0" * 64
    else:
        content["sources"][0]["title"] += " Altered after research review."
    write_json(path, content)
    changed = FrozenArtifact(path=artifact.path, sha256=digest_file(path), size=path.stat().st_size)
    for key, existing in list(record.artifacts.items()):
        if existing.path == artifact.path:
            record.artifacts[key] = changed
    ws.save("workflow", record)
    with pytest.raises(WorkflowError) as rejected:
        service.start_experiment(research_id)
    assert rejected.value.code == "ARTIFACT_CHANGED" and runner.calls == 0


def test_successful_evidence_cannot_be_rerun_or_replanned(setup):
    service, runner, research_id = setup
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert finished(service, research_id)["stage"] == "analyzed"
    for operation in (lambda: service.start_experiment(research_id), lambda: service.submit_code(research_id, BUNDLE, REVIEW)):
        with pytest.raises(WorkflowError) as rejected:
            operation()
        assert rejected.value.code == "EXPERIMENT_ALREADY_DISPATCHED"
    with pytest.raises(WorkflowError, match="stage"):
        approve_study(service, research_id)
    assert runner.calls == 1


@pytest.mark.parametrize("field", ["accepted", "issues"])
def test_native_review_rejection_prevents_dispatch(setup, field):
    service, runner, research_id = setup
    approve_study(service, research_id)
    review = copy.deepcopy(REVIEW)
    review[field] = False if field == "accepted" else ["Independent oracle defect"]
    with pytest.raises(WorkflowError, match="defect|accept"):
        service.submit_code(research_id, BUNDLE, review)
    assert runner.calls == 0 and service.status(research_id)["stage"] == "planned"


@pytest.mark.parametrize("target", ["proposal", "plan", "study-review", "selected-literature", "literature", "bundle", "source", "generated", "extra"])
def test_frozen_artifacts_cannot_change_before_execution(setup, target):
    service, runner, research_id = setup
    prepare(service, research_id)
    if target == "source":
        path = service.home / "workflows" / research_id / "source" / "transform.js"
    elif target in {"generated", "extra"}:
        path = service.artifact_path(research_id, "bundle").parent / "experiment.mjs"
        if target == "extra":
            path = path.with_name("injected.js")
    else:
        path = service.artifact_path(research_id, target)
    if path.exists():
        path.chmod(0o644)
    path.write_text("changed", encoding="utf-8")
    with pytest.raises(ValueError, match="changed|modified"):
        service.start_experiment(research_id)
    assert runner.calls == 0


@pytest.mark.parametrize("failed_exit", [False, True])
def test_failed_scientific_control_is_terminal_even_after_nonzero_exit(setup, failed_exit):
    service, runner, research_id = setup
    runner.outputs["controls"][1]["passed"] = False
    if failed_exit:
        runner.exit_status = "failed"
    prepare(service, research_id)
    service.start_experiment(research_id)
    result = finished(service, research_id)
    assert result["code"] == "CONTROL_FAILED" and result["terminal_control_failure"]
    assert "observations" in result["artifacts"]
    for operation in (lambda: service.start_experiment(research_id), lambda: service.submit_code(research_id, BUNDLE, REVIEW)):
        with pytest.raises(WorkflowError) as rejected:
            operation()
        assert rejected.value.code == "EXPERIMENT_ALREADY_DISPATCHED"
    assert runner.calls == 1


@pytest.mark.parametrize("defect", ["production_calls", "coverage_truncated", "rows", "boolean", "fixture"])
def test_runner_evidence_must_satisfy_production_and_observation_contracts(setup, defect):
    service, runner, research_id = setup
    if defect == "production_calls":
        runner.production_calls = False
    elif defect == "coverage_truncated":
        runner.coverage_truncated = True
    elif defect == "rows":
        runner.outputs["observations"].pop()
    elif defect == "boolean":
        runner.outputs["observations"][0]["value"] = True
    else:
        runner.outputs["fixtures"][0]["sha256"] = "0" * 64
    prepare(service, research_id)
    service.start_experiment(research_id)
    result = finished(service, research_id)
    assert result["status"] == "blocked" and "analysis" not in result
    assert "observations" in result["artifacts"]


def test_code_validation_repairs_remain_available_before_first_dispatch(setup):
    service, runner, research_id = setup
    approve_study(service, research_id)
    forged = copy.deepcopy(BUNDLE)
    forged["observations"] = observations()
    with pytest.raises(ValueError):
        service.submit_code(research_id, forged, REVIEW)
    forged = copy.deepcopy(BUNDLE)
    forged["files"][0]["path"] = forged["entrypoint"] = "../escape.mjs"
    with pytest.raises(ValueError, match="safe relative"):
        service.submit_code(research_id, forged, REVIEW)
    result = service.status(research_id)
    assert result["stage"] == "planned" and result["execution_attempt"] == runner.calls == 0
    accepted = service.submit_code(research_id, copy.deepcopy(BUNDLE), REVIEW)
    first_bundle = service.artifact_path(research_id, "bundle")
    first_bytes = first_bundle.read_bytes()
    assert accepted["code_attempt"] == 1 and accepted["execution_attempt"] == runner.calls == 0
    forged = copy.deepcopy(BUNDLE)
    forged["files"].append({"path": "bundle.json", "content": "{}"})
    with pytest.raises(ValueError, match="reserved controller metadata"):
        service.submit_code(research_id, forged, REVIEW)
    assert service.status(research_id)["code_attempt"] == 1
    repaired = copy.deepcopy(BUNDLE)
    repaired["files"][0]["content"] += "// Synthetic pre-dispatch code revision.\n"
    accepted = service.submit_code(research_id, repaired, REVIEW)
    assert accepted["stage"] == "code_ready" and accepted["code_attempt"] == 2
    assert accepted["execution_attempt"] == runner.calls == 0
    assert service.artifact_path(research_id, "bundle") != first_bundle and first_bundle.read_bytes() == first_bytes
    service.start_experiment(research_id)
    assert finished(service, research_id)["stage"] == "analyzed" and runner.calls == 1


@pytest.mark.parametrize("entrypoint", ["experiment.py", "experiment.ts", "experiment.json"])
def test_code_bundle_cannot_execute_native_or_non_module_files(setup, entrypoint):
    service, runner, research_id = setup
    approve_study(service, research_id)
    forged = copy.deepcopy(BUNDLE)
    forged["files"][0]["path"] = forged["entrypoint"] = entrypoint
    with pytest.raises(ValueError):
        service.submit_code(research_id, forged, REVIEW)
    assert runner.calls == 0
    assert "bundle" not in service.status(research_id)["artifacts"]


def test_start_is_immediate_and_cancellation_stops_owned_execution(setup):
    service, runner, research_id = setup
    runner.release = threading.Event()
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert runner.entered.wait(2)
    with pytest.raises(WorkflowError) as rejected:
        service.start_experiment(research_id)
    assert rejected.value.code == "EXPERIMENT_ALREADY_DISPATCHED"
    cancelled = service.cancel(research_id)
    assert cancelled["cancellation_requested"]
    assert cancelled["status"] == "cancelled" and runner.calls == 1
    assert cancelled["cleanup_confirmed"] is True and cancelled["cleanup_pending"] is False
    assert service._jobs[research_id].done()
    assert not Workspace(service.root / research_id).get("workflow", research_id, Workflow).active_handle


def test_cancel_waits_for_worker_checkpoint_without_holding_workflow_locks(setup, monkeypatch):
    service, runner, research_id = setup
    entered, requested, release = threading.Event(), threading.Event(), threading.Event()

    def delayed_cleanup(*args, cancel, on_handle, **kwargs):
        runner.calls += 1
        on_handle({"kind": "fixture", "pid": 123, "simulation": True})
        entered.set()
        deadline = time.monotonic() + 5
        while not cancel() and time.monotonic() < deadline:
            requested.wait(0.01)
        assert cancel()
        requested.set()
        assert release.wait(5)
        return {"status": "cancelled", "cleanup_confirmed": True, "simulation": True}

    monkeypatch.setattr(runner, "run", delayed_cleanup)
    monkeypatch.setattr(runner, "stop", lambda handle: False)
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert entered.wait(2)
    with ThreadPoolExecutor(max_workers=1) as pool:
        cancellation = pool.submit(service.cancel, research_id)
        try:
            assert requested.wait(2)
            assert not cancellation.done()
            pending = service.status(research_id)
            assert pending["status"] == "running" and pending["cleanup_pending"] is True
            assert pending["cancellation_requested"] is True
        finally:
            release.set()
        result = cancellation.result(timeout=2)
    assert result["status"] == "cancelled" and result["cleanup_confirmed"] is True
    assert result["cleanup_pending"] is False and runner.calls == 1


def test_cancel_deadline_preserves_pending_identity_and_retry_never_redispatches(setup, monkeypatch):
    from paper_factory import workflow as workflow_module
    service, runner, research_id = setup
    release = threading.Event()

    def delayed_cleanup(*args, on_handle, **kwargs):
        runner.calls += 1
        on_handle({"kind": "fixture", "pid": 123, "simulation": True})
        runner.entered.set()
        assert release.wait(5)
        return {"status": "cancelled", "cleanup_confirmed": True, "simulation": True}

    monkeypatch.setattr(runner, "run", delayed_cleanup)
    monkeypatch.setattr(runner, "stop", lambda handle: False)
    monkeypatch.setattr(workflow_module, "CANCEL_CLEANUP_SECONDS", 0.05)
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert runner.entered.wait(2)
    try:
        started = time.monotonic()
        with pytest.raises(WorkflowError) as rejected:
            service.cancel(research_id)
        assert rejected.value.code == "CLEANUP_UNCONFIRMED"
        assert time.monotonic() - started < 1
        pending = service.status(research_id)
        assert pending["code"] == "CLEANUP_UNCONFIRMED" and pending["cleanup_pending"] is True
        assert "cleanup_confirmed" not in pending
        record = Workspace(service.root / research_id).get("workflow", research_id, Workflow)
        assert record.active_handle["pid"] == 123 and record.execution_attempt == 1
    finally:
        release.set()
        service._jobs[research_id].result(timeout=2)
    result = service.cancel(research_id)
    assert result["status"] == "cancelled" and result["cleanup_pending"] is False
    assert result["cleanup_confirmed"] is True and result["execution_attempt"] == 1 and runner.calls == 1


def test_cancel_does_not_restart_the_deadline_for_post_worker_cleanup(setup, monkeypatch):
    from paper_factory import workflow as workflow_module
    service, runner, research_id = setup
    release, entered = threading.Event(), threading.Event()

    def unresolved_cleanup(*args, on_handle, **kwargs):
        runner.calls += 1
        handle = {"kind": "fixture", "pid": 123, "simulation": True}
        on_handle(handle)
        entered.set()
        assert release.wait(5)
        return {"status": "cancelled", "cleanup_confirmed": False, "active_handle": handle, "simulation": True}

    monkeypatch.setattr(runner, "run", unresolved_cleanup)
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert entered.wait(2)
    elapsed, stop_calls = [0], []

    def late_worker_exit(futures, timeout):
        assert timeout == workflow_module.CANCEL_CLEANUP_SECONDS
        release.set()
        for future in futures:
            future.result(timeout=2)
        elapsed[0] = workflow_module.CANCEL_CLEANUP_SECONDS - 0.5
        return set(futures), set()

    with monkeypatch.context() as clock_patch:
        clock_patch.setattr(workflow_module, "time", SimpleNamespace(monotonic=lambda: elapsed[0]))
        clock_patch.setattr(workflow_module, "wait", late_worker_exit)
        clock_patch.setattr(runner, "stop", lambda handle: stop_calls.append(handle) or True)
        try:
            with pytest.raises(WorkflowError) as rejected:
                service.cancel(research_id)
            assert rejected.value.code == "CLEANUP_UNCONFIRMED" and stop_calls == []
            assert service.status(research_id)["cleanup_pending"] is True
        finally:
            release.set()
    result = service.cancel(research_id)
    assert result["cleanup_confirmed"] is True and result["cleanup_pending"] is False and runner.calls == 1


@pytest.mark.parametrize("failed_control", [False, True])
def test_cancel_retry_confirms_cleanup_and_retains_success_or_failed_control_evidence(setup, failed_control):
    service, runner, research_id = setup
    runner.cleanup_confirmed = False
    runner.cleanable = False
    if failed_control:
        runner.outputs["controls"][1]["passed"] = False
    prepare(service, research_id)
    service.start_experiment(research_id)
    initial = finished(service, research_id)
    assert initial["code"] == "CLEANUP_UNCONFIRMED" and initial["cleanup_pending"] is True
    service._jobs[research_id].result(timeout=2)
    retained = {key: service.artifact_path(research_id, key).read_bytes() for key in initial["artifacts"]}
    with pytest.raises(WorkflowError) as rejected:
        service.cancel(research_id)
    assert rejected.value.code == "CLEANUP_UNCONFIRMED"
    pending = service.status(research_id)
    assert pending["code"] == "CLEANUP_UNCONFIRMED" and pending["cleanup_pending"] is True
    runner.cleanable = True
    result = service.cancel(research_id)
    assert result["cleanup_confirmed"] is True and result["cleanup_pending"] is False
    assert runner.calls == 1 and result["execution_attempt"] == 1
    assert all(service.artifact_path(research_id, key).read_bytes() == raw for key, raw in retained.items())
    cleanup = json.loads(service.artifact_path(research_id, "cleanup-1").read_bytes())
    assert cleanup["confirmed"] is True and cleanup["execution_sha256"] == initial["artifacts"]["execution"]["sha256"]
    if failed_control:
        assert result["status"] == "blocked" and result["code"] == "CONTROL_FAILED"
        assert result["terminal_control_failure"] is True and "analysis" not in result
    else:
        assert result["status"] == "cancelled" and result["stage"] == "analyzed"
        assert result["code"] == "CANCELLED" and "analysis" in result


def test_cancel_cannot_confirm_an_orphan_cleanup_failure_without_an_identity_or_receipt(setup):
    service, runner, research_id = setup
    ws = Workspace(service.root / research_id)
    record = ws.get("workflow", research_id, Workflow)
    record.status, record.code = "blocked", "CLEANUP_UNCONFIRMED"
    ws.save("workflow", record)
    for _ in range(2):
        with pytest.raises(WorkflowError) as rejected:
            service.cancel(research_id)
        assert rejected.value.code == "CLEANUP_UNCONFIRMED"
        pending = service.status(research_id)
        assert pending["status"] == "blocked" and pending["code"] == "CLEANUP_UNCONFIRMED"
        assert pending["cleanup_pending"] is True and "cleanup_confirmed" not in pending
    assert runner.calls == 0


def test_cancel_rejects_a_failed_future_until_the_retained_identity_is_actually_stopped(setup):
    service, runner, research_id = setup
    future = Future()
    future.set_exception(OSError("Synthetic worker checkpoint failure"))
    service._jobs[research_id] = future
    ws = Workspace(service.root / research_id)
    record = ws.get("workflow", research_id, Workflow)
    record.status, record.stage, record.execution_attempt = "running", "execute", 1
    ws.save("workflow", record)
    with pytest.raises(WorkflowError) as rejected:
        service.cancel(research_id)
    assert rejected.value.code == "CLEANUP_UNCONFIRMED"
    assert service.status(research_id)["cleanup_pending"] is True
    record = ws.get("workflow", research_id, Workflow)
    record.active_handle = {"kind": "fixture", "pid": 123, "simulation": True}
    ws.save("workflow", record)
    runner.cleanable = False
    with pytest.raises(WorkflowError) as rejected:
        service.cancel(research_id)
    assert rejected.value.code == "CLEANUP_UNCONFIRMED"
    runner.cleanable = True
    result = service.cancel(research_id)
    assert result["status"] == "cancelled" and result["stage"] == "code_ready"
    assert result["cleanup_confirmed"] is True and result["cleanup_pending"] is False
    assert "execution" not in result["artifacts"] and result["execution_attempt"] == 1 and runner.calls == 0
    assert service.cancel(research_id)["cleanup_confirmed"] is True


@pytest.mark.parametrize("previous_code", [None, "CLEANUP_UNCONFIRMED"])
def test_cancel_failure_without_identity_or_receipt_remains_pending_after_restart(setup, previous_code):
    service, runner, research_id = setup
    ws = Workspace(service.root / research_id)
    record = ws.get("workflow", research_id, Workflow)
    record.status, record.stage, record.code = "running", "execute", previous_code
    ws.save("workflow", record)
    recovered = WorkflowService(service.home, runner=runner, collector=collect)
    try:
        pending = recovered.status(research_id)
        assert pending["code"] == "CLEANUP_UNCONFIRMED" and pending["cleanup_pending"] is True
        with pytest.raises(WorkflowError) as rejected:
            recovered.cancel(research_id)
        assert rejected.value.code == "CLEANUP_UNCONFIRMED"
        assert recovered.status(research_id)["cleanup_pending"] is True and runner.calls == 0
    finally:
        with pytest.raises(WorkflowError) as rejected:
            recovered.close()
        assert rejected.value.code == "CLEANUP_UNCONFIRMED"


def test_cancel_waiting_for_a_worker_preserves_its_natural_completion(setup, monkeypatch):
    service, runner, research_id = setup
    entered, release = threading.Event(), threading.Event()

    def finish_naturally(identifier, lease):
        try:
            entered.set()
            assert release.wait(5)
            with service._operation(identifier) as (ws, record):
                record.status, record.stage, record.code = "completed", "exported", None
                service._save(ws, record)
        finally:
            lease.__exit__(None, None, None)

    monkeypatch.setattr(service, "_execute", finish_naturally)
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert entered.wait(2)
    with ThreadPoolExecutor(max_workers=1) as pool:
        cancellation = pool.submit(service.cancel, research_id)
        try:
            deadline = time.monotonic() + 2
            while not service.status(research_id)["cancellation_requested"] and time.monotonic() < deadline:
                time.sleep(0.01)
            assert service.status(research_id)["cancellation_requested"] is True
            assert not cancellation.done()
        finally:
            release.set()
        result = cancellation.result(timeout=2)
    assert result["status"] == "completed" and result["stage"] == "exported" and result["code"] is None
    assert result["cleanup_confirmed"] is True and result["cleanup_pending"] is False and runner.calls == 0


def test_cancel_preserves_a_completed_workflow_and_existing_artifacts(setup):
    service, runner, research_id = setup
    ws = retained_authoring_fixture(service, research_id)
    record = ws.get("workflow", research_id, Workflow)
    record.status, record.stage = "completed", "exported"
    ws.save("workflow", record)
    retained = {key: service.artifact_path(research_id, key).read_bytes() for key in record.artifacts}
    result = service.cancel(research_id)
    assert result["status"] == "completed" and result["stage"] == "exported"
    assert result["cancellation_requested"] is False
    assert result["cleanup_confirmed"] is True and result["cleanup_pending"] is False
    assert all(service.artifact_path(research_id, key).read_bytes() == raw for key, raw in retained.items())
    assert runner.calls == 0


def test_cancel_preserves_a_terminal_control_failure(setup):
    service, runner, research_id = setup
    runner.outputs["controls"][1]["passed"] = False
    prepare(service, research_id)
    service.start_experiment(research_id)
    initial = finished(service, research_id)
    assert initial["terminal_control_failure"] is True
    result = service.cancel(research_id)
    assert result["status"] == "blocked" and result["code"] == "CONTROL_FAILED"
    assert result["terminal_control_failure"] is True and result["artifacts"] == initial["artifacts"]
    assert result["cleanup_confirmed"] is True and result["cleanup_pending"] is False and runner.calls == 1


def retained_authoring_fixture(service, research_id):
    """Freeze an explicitly synthetic prior analysis without dispatching a runner."""
    prepare(service, research_id)
    ws = Workspace(service.root / research_id)
    record = ws.get("workflow", research_id, Workflow)
    raw = observations()
    plan = json.loads(service.artifact_path(research_id, "plan").read_bytes())
    execution = {"status": "succeeded", "cleanup_confirmed": True, "coverage_truncated": False,
                 "coverage_mechanism": "synthetic prior execution fixture",
                 "production_calls": [{"path": "transform.js", "function": "transform", "calls": 6}]}
    analysis = science._compute(raw, plan)
    for key, value in (("execution", execution), ("observations", raw)):
        path = ws.path("research/" + key + "-fixture.json")
        write_json(path, value)
        record.artifacts[key] = FrozenArtifact(path=path.relative_to(ws.root).as_posix(),
                                               sha256=digest_file(path), size=path.stat().st_size)
    analysis.update(raw_sha256=record.artifacts["observations"].sha256,
                    protocol_sha256=record.artifacts["plan"].sha256)
    path = ws.path("research/analysis-fixture.json")
    write_json(path, analysis)
    record.artifacts["analysis"] = FrozenArtifact(path=path.relative_to(ws.root).as_posix(),
                                                   sha256=digest_file(path), size=path.stat().st_size)
    record.execution_attempt, record.stage = 1, "analyzed"
    ws.save("workflow", record)
    return ws


@pytest.mark.parametrize("stage", ["created", "proposed", "planned", "code_ready"])
@pytest.mark.parametrize("previous_status", ["ready", "cancelled"])
def test_preparation_resume_preserves_verified_inputs_and_never_dispatches(setup, stage, previous_status, monkeypatch):
    service, runner, research_id = setup
    if stage == "proposed":
        service.submit_proposal(research_id, protocol())
    elif stage == "planned":
        approve_study(service, research_id)
    elif stage == "code_ready":
        prepare(service, research_id)
    state = service.cancel(research_id) if previous_status == "cancelled" else service.status(research_id)
    assert state["resume_kind"] == "preparation" and state["execution_attempt"] == 0
    ws = Workspace(service.root / research_id)
    prior = ws.get("workflow", research_id, Workflow).model_dump(mode="json")
    retained = {key: service.artifact_path(research_id, key).read_bytes() for key in state["artifacts"]}
    monkeypatch.setattr(runner, "run", lambda *args, **kwargs: pytest.fail("Preparation resume dispatched a measurement"))
    monkeypatch.setattr(runner, "status", lambda: pytest.fail("Preparation resume launched a worker probe"))
    result = service.resume(research_id)
    assert result["status"] == "ready" and result["stage"] == stage and result["resume_kind"] == "preparation"
    assert result["execution_attempt"] == 0 and result["cancellation_requested"] is False and runner.calls == 0
    keys = [key for key in result["artifacts"] if key.startswith("workflow-resume-")]
    assert len(keys) == 1
    resumed = json.loads(service.artifact_path(research_id, keys[0]).read_bytes())
    assert resumed["event"] == "workflow-resume" and resumed["kind"] == "preparation"
    assert resumed["previous_workflow"] == prior and resumed["execution_attempt"] == 0
    assert "execution_sha256" not in resumed
    assert all(service.artifact_path(research_id, key).read_bytes() == content for key, content in retained.items())


@pytest.mark.parametrize("defect,expected_code", [
    ("source_changed", None), ("plan_changed", "ARTIFACT_CHANGED"), ("bundle_missing", "ARTIFACT_MISSING"),
    ("review_rejected", "REVIEW_REJECTED"), ("execution_attempt", "INVALID_STATE"),
    ("execution_artifact", "INVALID_STATE"), ("execution_files", "EXPERIMENT_ALREADY_DISPATCHED"),
    ("pending_future", "CLEANUP_UNCONFIRMED"), ("terminal_control", "CONTROL_FAILED"),
    ("cleanup_unconfirmed", "CLEANUP_UNCONFIRMED"),
])
def test_preparation_resume_rejects_unknown_execution_or_changed_frozen_inputs(setup, defect, expected_code):
    service, runner, research_id = setup
    prepare(service, research_id)
    service.cancel(research_id)
    ws = Workspace(service.root / research_id)
    record = ws.get("workflow", research_id, Workflow)
    pending = None
    if defect == "source_changed":
        source = ws.path("source/transform.js")
        source.chmod(0o644)
        source.write_text("export function changed() { return 0; }", encoding="utf-8")
    elif defect == "plan_changed":
        service.artifact_path(research_id, "plan").write_text("{}", encoding="utf-8")
    elif defect == "bundle_missing":
        record.artifacts.pop("bundle")
    elif defect == "review_rejected":
        key = "code-review-1"
        path = service.artifact_path(research_id, key)
        review = json.loads(path.read_bytes())
        review["review"]["accepted"], review["review"]["issues"] = False, ["Synthetic reviewer rejection"]
        write_json(path, review)
        record.artifacts[key] = FrozenArtifact(path=path.relative_to(ws.root).as_posix(), sha256=digest_file(path), size=path.stat().st_size)
    elif defect == "execution_attempt":
        record.execution_attempt = 1
    elif defect == "execution_artifact":
        record.artifacts["execution-1"] = record.artifacts["context"]
    elif defect == "execution_files":
        path = ws.path("research/executions/attempt-1")
        path.mkdir(parents=True)
        write_json(path / "execution.json", {"fixture": "uncheckpointed execution bytes"})
    elif defect == "pending_future":
        pending = Future()
        service._jobs[research_id] = pending
    elif defect == "terminal_control":
        record.terminal_control_failure = True
    else:
        record.code = "CLEANUP_UNCONFIRMED"
    ws.save("workflow", record)
    before = record.model_dump(mode="json")
    try:
        with pytest.raises(ValueError) as rejected:
            service.resume(research_id)
        if expected_code:
            assert rejected.value.code == expected_code
        after = ws.get("workflow", research_id, Workflow).model_dump(mode="json")
        assert after == before and runner.calls == 0
        assert not any(key.startswith("workflow-resume-") for key in after["artifacts"])
    finally:
        if pending is not None:
            pending.set_result(None)
            service._jobs.pop(research_id)


def test_cancelled_preparation_requires_explicit_resume_before_first_and_only_dispatch(setup):
    service, runner, research_id = setup
    prepare(service, research_id)
    service.cancel(research_id)
    for submit in (lambda: service.start_experiment(research_id), lambda: service.submit_code(research_id, BUNDLE, REVIEW)):
        with pytest.raises(WorkflowError) as rejected:
            submit()
        assert rejected.value.code == "CANCELLED" and runner.calls == 0
    service.resume(research_id)
    service.start_experiment(research_id)
    result = finished(service, research_id)
    assert result["stage"] == "analyzed" and result["execution_attempt"] == 1 and runner.calls == 1
    service.cancel(research_id)
    assert service.resume(research_id)["resume_kind"] == "authoring"
    with pytest.raises(WorkflowError) as rejected:
        service.start_experiment(research_id)
    assert rejected.value.code == "EXPERIMENT_ALREADY_DISPATCHED" and runner.calls == 1


@pytest.mark.parametrize("executed", [False, True])
def test_resume_after_restart_preserves_dispatch_history_and_frozen_evidence(setup, executed, monkeypatch):
    service, runner, research_id = setup
    prepare(service, research_id)
    if executed:
        service.start_experiment(research_id)
        assert finished(service, research_id)["stage"] == "analyzed"
    prior = service.cancel(research_id)
    retained = {key: service.artifact_path(research_id, key).read_bytes() for key in prior["artifacts"]}
    service.close()
    recovered = WorkflowService(service.home, runner=runner, collector=collect)
    try:
        kind = "authoring" if executed else "preparation"
        assert recovered.status(research_id)["resume_kind"] == kind
        monkeypatch.setattr(runner, "run", lambda *args, **kwargs: pytest.fail("Resume redispatched an experiment"))
        result = recovered.resume(research_id)
        assert result["resume_kind"] == kind and result["execution_attempt"] == int(executed)
        assert all(recovered.artifact_path(research_id, key).read_bytes() == content for key, content in retained.items())
        if executed:
            with pytest.raises(WorkflowError) as rejected:
                recovered.start_experiment(research_id)
            assert rejected.value.code == "EXPERIMENT_ALREADY_DISPATCHED" and runner.calls == 1
    finally:
        recovered.close()


@pytest.mark.parametrize("stage", ["analyzed", "manuscript"])
@pytest.mark.parametrize("previous_status", ["ready", "cancelled"])
def test_authoring_resumes_without_dispatch_and_preserves_prior_state(setup, stage, previous_status, monkeypatch):
    service, runner, research_id = setup
    ws = retained_authoring_fixture(service, research_id)
    record = ws.get("workflow", research_id, Workflow)
    record.stage = stage
    ws.save("workflow", record)
    state = service.cancel(research_id) if previous_status == "cancelled" else service.status(research_id)
    assert state["resume_kind"] == "authoring"
    before = {key: service.artifact_path(research_id, key).read_bytes() for key in state["artifacts"]}
    prior = ws.get("workflow", research_id, Workflow).model_dump(mode="json")
    monkeypatch.setattr(service, "start_experiment", lambda *args: pytest.fail("Authoring resume dispatched an experiment"))
    monkeypatch.setattr(runner, "run", lambda *args, **kwargs: pytest.fail("Authoring resume called the runner"))
    resumed = service.resume(research_id)
    assert resumed["status"] == "ready" and resumed["stage"] == stage
    assert resumed["code"] is None and resumed["cancellation_requested"] is False
    assert resumed["execution_attempt"] == 1 and runner.calls == 0
    keys = [key for key in resumed["artifacts"] if key.startswith("workflow-resume-")]
    assert len(keys) == 1
    receipt = json.loads(service.artifact_path(research_id, keys[0]).read_bytes())
    assert receipt["previous_workflow"] == prior and receipt["previous_status"] == previous_status
    assert receipt["event"] == "workflow-resume" and receipt["kind"] == "authoring"
    assert receipt["previous_stage"] == stage and receipt["execution_attempt"] == 1
    assert receipt["execution_sha256"] == resumed["artifacts"]["execution"]["sha256"]
    assert receipt["at"].endswith("+00:00")
    for key, content in before.items():
        assert service.artifact_path(research_id, key).read_bytes() == content
    original_receipt = service.artifact_path(research_id, keys[0]).read_bytes()
    repeated = service.resume(research_id)
    assert len([key for key in repeated["artifacts"] if key.startswith("workflow-resume-")]) == 2
    service.cancel(research_id)
    again = service.resume(research_id)
    assert len([key for key in again["artifacts"] if key.startswith("workflow-resume-")]) == 3
    assert service.artifact_path(research_id, keys[0]).read_bytes() == original_receipt
    assert again["execution_attempt"] == 1 and runner.calls == 0


@pytest.mark.parametrize("defect,expected_code", [
    ("execution_stage", "INVALID_STATE"), ("terminal_control", "CONTROL_FAILED"),
    ("active_handle", "CLEANUP_UNCONFIRMED"), ("cleanup", "CLEANUP_UNCONFIRMED"),
    ("failed_execution", "EXPERIMENT_FAILED"), ("production_gate", "PRODUCTION_EXECUTION_UNVERIFIED"),
    ("failed_control", "CONTROL_FAILED"), ("missing_negative", "CONTROL_FAILED"),
    ("tampered_raw", "ARTIFACT_CHANGED"), ("analysis_link", "ARTIFACT_CHANGED"),
])
def test_authoring_resume_rejects_unsafe_or_changed_retained_evidence(setup, defect, expected_code):
    service, runner, research_id = setup
    ws = retained_authoring_fixture(service, research_id)
    service.cancel(research_id)
    record = ws.get("workflow", research_id, Workflow)
    if defect == "execution_stage":
        record.stage = "execute"
    elif defect == "terminal_control":
        record.terminal_control_failure = True
    elif defect == "active_handle":
        record.active_handle = {"kind": "fixture", "pid": 123}
    else:
        key = "observations" if defect in {"failed_control", "missing_negative", "tampered_raw"} else "analysis" if defect == "analysis_link" else "execution"
        path = service.artifact_path(research_id, key)
        value = json.loads(path.read_bytes())
        if defect == "cleanup":
            value["cleanup_confirmed"] = False
        elif defect == "failed_execution":
            value["status"] = "cancelled"
        elif defect == "production_gate":
            value["production_calls"] = []
        elif defect == "failed_control":
            value["controls"][1]["passed"] = False
        elif defect == "missing_negative":
            value["controls"] = value["controls"][:1]
        elif defect == "analysis_link":
            value["raw_sha256"] = "0" * 64
        else:
            value["observations"][0]["value"] = 99
        write_json(path, value)
        if defect != "tampered_raw":
            record.artifacts[key] = FrozenArtifact(path=path.relative_to(ws.root).as_posix(),
                                                   sha256=digest_file(path), size=path.stat().st_size)
    ws.save("workflow", record)
    with pytest.raises(WorkflowError) as rejected:
        service.resume(research_id)
    assert rejected.value.code == expected_code
    after = ws.get("workflow", research_id, Workflow)
    assert after.status == "cancelled" and after.execution_attempt == 1 and runner.calls == 0
    assert not any(key.startswith("workflow-resume-") for key in after.artifacts)
    if defect == "active_handle":
        # The fixture identity is not a real process; keep fixture teardown local.
        after.active_handle = {}
        ws.save("workflow", after)


def test_authoring_resume_accepts_hash_bound_confirmed_cleanup_receipt(setup):
    service, runner, research_id = setup
    ws = retained_authoring_fixture(service, research_id)
    service.cancel(research_id)
    record = ws.get("workflow", research_id, Workflow)
    execution_path = service.artifact_path(research_id, "execution")
    execution = json.loads(execution_path.read_bytes())
    execution["cleanup_confirmed"] = False
    write_json(execution_path, execution)
    record.artifacts["execution"] = FrozenArtifact(path=execution_path.relative_to(ws.root).as_posix(),
                                                   sha256=digest_file(execution_path), size=execution_path.stat().st_size)
    cleanup_path = ws.path("research/cleanup-fixture.json")
    write_json(cleanup_path, {"confirmed": True, "execution_sha256": record.artifacts["execution"].sha256})
    record.artifacts["cleanup-1"] = FrozenArtifact(path=cleanup_path.relative_to(ws.root).as_posix(),
                                                   sha256=digest_file(cleanup_path), size=cleanup_path.stat().st_size)
    ws.save("workflow", record)
    assert service.resume(research_id)["status"] == "ready" and runner.calls == 0


def test_close_cancels_workers_and_prevents_new_operations(setup):
    service, runner, research_id = setup
    runner.release = threading.Event()
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert runner.entered.wait(2)
    service.close()
    with pytest.raises(ValueError, match="closed"):
        service.start_experiment(research_id)
    ws = Workspace(service.home / "workflows" / research_id)
    record = ws.get("workflow", research_id, Workflow)
    assert record.status == "cancelled" and not record.active_handle


def test_close_timeout_keeps_worker_identity_and_allows_evidence_cancel_and_shutdown_retry(setup, monkeypatch):
    from test_standalone_ipc import receipt
    service, runner, research_id = setup
    entered, release = threading.Event(), threading.Event()

    def delayed_cleanup(*args, on_handle, **kwargs):
        runner.calls += 1
        on_handle({"kind": "fixture", "pid": 123, "simulation": True})
        entered.set()
        assert release.wait(5)
        return {"status": "cancelled", "cleanup_confirmed": True, "simulation": True}

    monkeypatch.setattr(runner, "run", delayed_cleanup)
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert entered.wait(2)
    try:
        started = time.monotonic()
        with pytest.raises(WorkflowError) as rejected:
            service.close(deadline=started + 0.05)
        assert rejected.value.code == "CLEANUP_UNCONFIRMED" and time.monotonic() - started < 1
        assert service._closing is True and service._closed is False
        pending = service.status(research_id)
        assert pending["cleanup_pending"] is True and pending["code"] == "CLEANUP_UNCONFIRMED"
        record = Workspace(service.root / research_id).get("workflow", research_id, Workflow)
        assert record.active_handle["pid"] == 123 and not service._jobs[research_id].done()
        with pytest.raises(WorkflowError) as rejected:
            service.start_experiment(research_id)
        assert rejected.value.code == "ENGINE_CLOSING"
        saved = service.record_inference(research_id, receipt("started"))
        receipt_path = service.artifact_path(research_id, saved["artifactId"])
        retained_receipt = receipt_path.read_bytes()
        with ThreadPoolExecutor(max_workers=1) as pool:
            cancellation = pool.submit(service.cancel, research_id)
            release.set()
            result = cancellation.result(timeout=3)
        assert result["cleanup_confirmed"] is True and result["cleanup_pending"] is False
        service.close()
        service.close()
        assert service._closed is True and runner.calls == 1 and receipt_path.read_bytes() == retained_receipt
    finally:
        release.set()


def test_close_rejects_failed_future_without_cleanup_proof_and_retries_exact_identity(setup):
    service, runner, research_id = setup
    future = Future()
    future.set_exception(OSError("Synthetic final checkpoint failed"))
    service._jobs[research_id] = future
    ws = Workspace(service.root / research_id)
    record = ws.get("workflow", research_id, Workflow)
    record.status, record.stage, record.execution_attempt = "running", "execute", 1
    ws.save("workflow", record)
    with pytest.raises(WorkflowError) as rejected:
        service.close()
    assert rejected.value.code == "CLEANUP_UNCONFIRMED"
    assert service._closed is False and service.status(research_id)["cleanup_pending"] is True
    record = ws.get("workflow", research_id, Workflow)
    record.active_handle = {"kind": "fixture", "pid": 123, "simulation": True}
    ws.save("workflow", record)
    runner.cleanable = False
    with pytest.raises(WorkflowError) as rejected:
        service.close()
    assert rejected.value.code == "CLEANUP_UNCONFIRMED"
    assert ws.get("workflow", research_id, Workflow).active_handle["pid"] == 123
    runner.cleanable = True
    service.close()
    service.close()
    assert service._closed is True and not ws.get("workflow", research_id, Workflow).active_handle and runner.calls == 0


def test_unconfirmed_cleanup_preserves_identity_and_blocks_reexecution(setup):
    service, runner, research_id = setup
    runner.cleanup_confirmed = False
    prepare(service, research_id)
    service.start_experiment(research_id)
    result = finished(service, research_id)
    assert result["code"] == "CLEANUP_UNCONFIRMED"
    ws = Workspace(service.home / "workflows" / research_id)
    assert ws.get("workflow", research_id, Workflow).active_handle["pid"] == 123
    with pytest.raises(WorkflowError) as rejected:
        service.start_experiment(research_id)
    assert rejected.value.code == "EXPERIMENT_ALREADY_DISPATCHED"


def test_native_export_reopens_files_and_recomputes_observations(setup, pandoc):
    service, runner, research_id = setup
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert finished(service, research_id)["stage"] == "analyzed"
    # An explicit image fixture tests companion export without automatic graphs.
    ws = Workspace(service.root / research_id)
    record = ws.get("workflow", research_id, Workflow)
    image = ws.path("research/analysis-fixture/figure-1.png")
    image.parent.mkdir(parents=True)
    image.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAADAAAAAgCAIAAADbtmxLAAAAOElEQVR4nO3OoQEAIAzAsP3/NPiZSBDNBZnz"
        "mXkd2ApJISkkhaSQFJJCUkgKSSEpJIWkkBSSQnIBd/rvALsGC8oAAAAASUVORK5CYII="))
    record.artifacts["analysis-figure-1.png"] = FrozenArtifact(
        path=image.relative_to(ws.root).as_posix(), sha256=digest_file(image), size=image.stat().st_size)
    ws.save("workflow", record)
    service.submit_manuscript(research_id, manuscript(), MANUSCRIPT_REVIEW)
    with pytest.MonkeyPatch.context() as monkeypatch:
        monkeypatch.setenv("PYPANDOC_PANDOC", pandoc)
        result = service.export(research_id)
    assert result["status"] == "completed" and result["stage"] == "exported"
    assert set(result["artifacts"]) >= {"export-md", "export-pdf", "export-docx", "export-tex", "validation", "reproducibility"}
    for artifact_id in ("export-md", "export-tex"):
        resolved = service.resolve_artifact(research_id, artifact_id)
        images = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", service.artifact_path(research_id, "export-md").read_text(encoding="utf-8"))
        assert {item["name"] for item in resolved["companions"]} == set(images) and images
        for item in resolved["companions"]:
            path = Path(item["path"])
            assert path.parent == Path(resolved["path"]).parent
            assert item["sha256"] == digest_file(path) and item["size"] == path.stat().st_size
    validation = json.loads(service.artifact_path(research_id, "validation").read_text())
    assert validation["passed"] is True
    assert validation["final_verification_in_archive"] is False
    journal_key = validation["verification_journal"]
    assert journal_key.startswith("verification-journal-") and journal_key in result["artifacts"]
    assert [row["phase"] for row in verification_rows(service.artifact_path(research_id, journal_key))] == VERIFICATION_PHASES
    with zipfile.ZipFile(service.artifact_path(research_id, "reproducibility")) as archive:
        assert archive.testzip() is None
        assert {"source/transform.js", "generated/experiment.mjs", "analysis/analysis.py", "observations.json"} <= set(archive.namelist())
        assert not {"verification.jsonl", "validation.json"} & set(archive.namelist())
        readme = archive.read("README.md").decode().casefold()
        assert "verification" in readme and "journal" in readme and "validation" in readme
        images = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", archive.read("paper.md").decode())
        assert images
        inventory = json.loads(archive.read("inventory.json"))
        for image in images:
            assert image in archive.namelist() and image in inventory
            assert "analysis/" + image in archive.namelist()
            assert archive.read(image) == archive.read("analysis/" + image)
    assert runner.calls == 1


def test_concise_manuscript_exports_real_native_documents_and_archive(setup, pandoc):
    from docx import Document
    from pypdf import PdfReader

    service, runner, research_id = setup
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert finished(service, research_id)["stage"] == "analyzed"
    prose = {
        "Abstract": (
            "This synthetic fixture verifies that concise evidence linked manuscripts survive native document conversion. "
            "The paired mean error difference is {{result:error.paired_2_minus_1.mean}} events. "
            "Measurements originate from an orchestration test runner, so the comparison validates artifact plumbing rather than scientific contribution or production correctness."),
        "Introduction": (
            "A document pipeline should preserve the content of an appropriately short manuscript without requiring repetition. "
            "This fixture addresses a conversion regression in which accepted prose could later fail a stricter document length check. "
            "The generated files provide concrete evidence of consistent handling across export formats."),
        "Related Work": (
            "The inspected synthetic passage describes expected outcomes defined independently of the implementation "
            "and warns that controlled annotations cannot establish population failure rates {{citation:fixture-oracle}}. "
            "It supplies a fixture citation for reference binding. "
            "It is not a published source and cannot establish an academic research gap."),
        "Research Questions": (
            "The fixture asks whether concise structured prose retains its declared results, required sections, and inspected reference "
            "when rendered into native documents. "
            "It also checks whether conversion receipts and archive membership preserve the same manuscript bytes. "
            "These are bounded engineering checks rather than claims about a population."),
        "Method": (
            "The test controller freezes source and protocol records before accepting simulated runner observations. "
            "Trusted analysis computes descriptive summaries from those records. "
            "The manuscript inserts numerical placeholders instead of copied quantities and selects an inspected fixture excerpt. "
            "Native converters then produce documents that the controller independently reopens."),
        "Experimental Setup": (
            "The protocol specifies {{parameter:units_per_seed}} units for each seed and retains the complete condition grid. "
            "Its instrumentation setting is {{parameter:setting.execution_instrumentation}}. "
            "The synthetic runner supplies explicit positive and intentional fault controls. "
            "Neither its replayed measurements nor bridge overhead justify claims about uninstrumented production performance."),
        "Results": (
            "The production mean error is {{result:error.condition_1.mean}} events, and the ablation mean error is "
            "{{result:error.condition_2.mean}} events. "
            "The paired mean difference is {{result:error.paired_2_minus_1.mean}} events. "
            "These values describe the retained synthetic matrix. "
            "Their agreement across prose and trusted tables tests numerical reference preservation."),
        "Discussion": (
            "A concise document can explain the fixture purpose and report its actual summaries without repeating generic text. "
            "Successful conversion means that native output preserves the supplied structure and evidence references. "
            "It does not promote a routine testing fixture into academically useful research or establish novelty."),
        "Threats to Validity": (
            "The test runner supplies simulated outputs and invocation records rather than observing an application workload. "
            "The annotations and protocol are deliberately narrow. "
            "Consequently this fixture cannot demonstrate independent sampling, realistic failure prevalence, or general software correctness. "
            "Converter success also cannot substitute for scientific manuscript assessment."),
        "Limitations": (
            "The manuscript exercises native conversion and archive integrity using a small controlled fixture. "
            "It contains no runtime benchmark, human study, or external research evaluation. "
            "Its positive model review is a test double. "
            "The retained source and observations support orchestration checks and must not be described as publication evidence."),
        "Conclusion": (
            "The concise fixture links declared summaries to retained observations while preserving its scientific limitations. "
            "Its native documents and archive test the expected export path without adding prose solely to satisfy a larger length threshold. "
            "Research quality remains a separate question that this regression fixture does not answer."),
    }
    draft = {"title": "Concise synthetic manuscript for native conversion regression",
             "sections": [{"heading": heading, "text": prose[heading]} for heading in science.REQUIRED_SECTIONS]}
    count = sum(len(re.findall(r"\b[\w'-]+\b", science.PLACEHOLDER.sub("evidence", section["text"])))
                for section in draft["sections"])
    assert 300 <= count <= 700 and len(draft["sections"]) == 11
    service.submit_manuscript(research_id, draft, MANUSCRIPT_REVIEW)
    with pytest.MonkeyPatch.context() as monkeypatch:
        monkeypatch.setenv("PYPANDOC_PANDOC", pandoc)
        exported = service.export(research_id)
    assert exported["stage"] == "exported" and exported["status"] == "completed"
    markdown = service.artifact_path(research_id, "export-md").read_text(encoding="utf-8")
    assert draft["title"] in markdown and "{{result:" not in markdown
    word = Document(service.artifact_path(research_id, "export-docx"))
    word_text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    assert 300 <= len(re.findall(r"\b[\w'-]+\b", word_text)) < 1000
    assert draft["title"] in word_text and all(heading in word_text for heading in science.REQUIRED_SECTIONS)
    pdf = PdfReader(service.artifact_path(research_id, "export-pdf"))
    assert pdf.pages and any(page.extract_text() for page in pdf.pages)
    tex = service.artifact_path(research_id, "export-tex").read_text(encoding="utf-8")
    assert "\\begin{document}" in tex and "\\end{document}" in tex
    with zipfile.ZipFile(service.artifact_path(research_id, "reproducibility")) as archive:
        assert archive.testzip() is None
        assert {"paper.md", "paper.pdf", "paper.docx", "paper.tex"} <= set(archive.namelist())
        assert archive.read("paper.md") == service.artifact_path(research_id, "export-md").read_bytes()
    assert runner.calls == 1


VERIFICATION_PHASES = ["artifact-source-validation", "trusted-analysis", "manuscript-rendering",
                       "pdf", "docx", "tex-conversion", "zip", "complete"]


def verification_rows(path):
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]


def synthetic_verification_export(setup, monkeypatch):
    """Exercise trusted verification using fake native documents, not native executables."""
    from paper_factory import conversion
    import docx
    import pypdf

    service, runner, research_id = setup
    ws = retained_authoring_fixture(service, research_id)
    service.submit_manuscript(research_id, manuscript(), MANUSCRIPT_REVIEW)

    def no_analysis_rendering(*args, **kwargs):
        pytest.fail("Export verification must recompute values without analysis files or figures")

    def fixture_conversion(markdown, output, **kwargs):
        content = (b"\\begin{document}\nSynthetic verification fixture.\n\\end{document}\n"
                   if output.suffix == ".tex" else b"Synthetic native document fixture: " + output.suffix.encode())
        output.write_bytes(content)
        return {"input_sha256": digest_file(markdown), "output_sha256": digest_file(output)}

    monkeypatch.setattr(science, "analyze", no_analysis_rendering)
    monkeypatch.setattr(conversion, "convert", fixture_conversion)
    monkeypatch.setattr(pypdf, "PdfReader", lambda *args, **kwargs: SimpleNamespace(
        is_encrypted=False, pages=[SimpleNamespace(extract_text=lambda: "Synthetic readable PDF fixture.")]))
    def fixture_word_document(path, *args, **kwargs):
        markdown = Path(path).with_suffix(".md").read_text(encoding="utf-8")
        return SimpleNamespace(paragraphs=[SimpleNamespace(text=markdown)])

    monkeypatch.setattr(docx, "Document", fixture_word_document)
    return service, runner, research_id, ws


def test_verification_recomputes_equal_values_without_rendering_analysis_and_fsyncs_each_phase(setup, monkeypatch):
    from paper_factory import workflow as workflow_module

    service, runner, research_id, ws = synthetic_verification_export(setup, monkeypatch)
    raw = json.loads(service.artifact_path(research_id, "observations").read_bytes())
    protocol_value = json.loads(service.artifact_path(research_id, "plan").read_bytes())
    retained = json.loads(service.artifact_path(research_id, "analysis").read_bytes())
    compute = science._compute
    expected = {key: value for key, value in retained.items() if key not in {"raw_sha256", "protocol_sha256"}}
    inputs = []

    def verified_compute(observations_value, frozen_protocol):
        inputs.append((copy.deepcopy(observations_value), copy.deepcopy(frozen_protocol)))
        result = compute(observations_value, frozen_protocol)
        assert result == expected
        return result

    monkeypatch.setattr(science, "_compute", verified_compute)
    journal_streams, durable_phases = [], []
    original_open, original_fsync, original_freeze = Path.open, os.fsync, workflow_module._freeze

    def tracked_open(path, *args, **kwargs):
        stream = original_open(path, *args, **kwargs)
        mode = args[0] if args else kwargs.get("mode", "r")
        if path.name == "verification.jsonl" and mode == "x":
            journal_streams.append(stream)
        return stream

    def durable_sync(descriptor):
        original_fsync(descriptor)
        for stream in journal_streams:
            if not stream.closed and stream.fileno() == descriptor:
                rows = verification_rows(Path(stream.name))
                assert rows, "The journal phase must be flushed before fsync"
                durable_phases.append(rows[-1]["phase"])

    def freeze_closed_journal(workspace, record, key, path):
        if key.startswith("verification-journal-"):
            assert journal_streams and all(stream.closed for stream in journal_streams)
        return original_freeze(workspace, record, key, path)

    monkeypatch.setattr(Path, "open", tracked_open)
    monkeypatch.setattr(workflow_module.os, "fsync", durable_sync)
    monkeypatch.setattr(workflow_module, "_freeze", freeze_closed_journal)
    result = service.export(research_id)
    validation = json.loads(service.artifact_path(research_id, "validation").read_bytes())
    journal = service.artifact_path(research_id, validation["verification_journal"])
    rows = verification_rows(journal)
    assert inputs == [(raw, protocol_value)]
    assert [row["phase"] for row in rows] == durable_phases == VERIFICATION_PHASES
    assert all(set(row) == {"at", "research_id", "phase"} and row["research_id"] == research_id
               and row["at"].endswith("+00:00") for row in rows)
    assert len(journal_streams) == 1 and journal_streams[0].closed
    assert journal.parent.name.startswith("export-attempt-")
    assert validation["passed"] is True and validation["final_verification_in_archive"] is False
    assert result["stage"] == "exported" and result["status"] == "completed"
    assert result["execution_attempt"] == 1 and runner.calls == 0


@pytest.mark.parametrize("defect", ["too_short", "missing_title", "missing_heading"])
def test_word_verification_requires_content_title_and_each_canonical_heading(setup, monkeypatch, defect):
    import docx

    service, runner, research_id, ws = synthetic_verification_export(setup, monkeypatch)
    canonical = json.loads(service.artifact_path(research_id, "canonical").read_bytes())
    paragraphs = [] if defect == "missing_title" else [canonical["title"]]
    for section in canonical["sections"]:
        if not (defect == "missing_heading" and section["heading"] == "Threats to Validity"):
            paragraphs.append(section["heading"])
        if defect != "too_short":
            paragraphs.append(section["text"])
    word_count = len(re.findall(r"\b[\w'-]+\b", "\n".join(paragraphs)))
    if defect == "too_short":
        assert word_count < 300
    else:
        assert word_count >= 300
    monkeypatch.setattr(docx, "Document", lambda *args, **kwargs: SimpleNamespace(
        paragraphs=[SimpleNamespace(text=text) for text in paragraphs]))

    with pytest.raises(ValueError, match="Native Word manuscript is incomplete"):
        service.export(research_id)
    journals = list(ws.path("research/exports").glob("*/verification.jsonl"))
    assert len(journals) == 1
    assert [row["phase"] for row in verification_rows(journals[0])] == VERIFICATION_PHASES[:5]
    state = service.status(research_id)
    assert state["stage"] == "manuscript" and state["execution_attempt"] == 1
    assert runner.calls == 0


@pytest.mark.parametrize("field", ["results", "parameters", "controls", "protocol_digest", "observation_digest", "summaries", "paired_deltas"])
def test_verification_rejects_changed_analysis_even_with_a_matching_artifact_hash(setup, monkeypatch, field):
    service, runner, research_id, ws = synthetic_verification_export(setup, monkeypatch)
    path = service.artifact_path(research_id, "analysis")
    analysis = json.loads(path.read_bytes())
    if field == "results":
        analysis[field]["error.condition_1.mean"]["value"] = 99
    elif field == "parameters":
        analysis[field]["units_per_seed"] = 99
    elif field == "controls":
        analysis[field][0]["details"] = "Synthetic altered interpretation of the retained control."
    elif field in {"protocol_digest", "observation_digest"}:
        analysis[field] = "0" * 64
    else:
        analysis[field][0]["mean"] = 99
    write_json(path, analysis)
    record = ws.get("workflow", research_id, Workflow)
    record.artifacts["analysis"] = FrozenArtifact(path=path.relative_to(ws.root).as_posix(),
                                                 sha256=digest_file(path), size=path.stat().st_size)
    ws.save("workflow", record)
    with pytest.raises(ValueError, match="Analysis differs from raw observation recomputation: " + field):
        service.export(research_id)
    journals = list(ws.path("research/exports").glob("*/verification.jsonl"))
    assert len(journals) == 1
    assert [row["phase"] for row in verification_rows(journals[0])] == VERIFICATION_PHASES[:2]
    after = ws.get("workflow", research_id, Workflow)
    assert not any(key.startswith("verification-journal-") for key in after.artifacts)
    assert after.stage == "manuscript" and after.execution_attempt == 1 and runner.calls == 0


@pytest.mark.parametrize("phase", VERIFICATION_PHASES[:-1])
def test_verification_failure_preserves_durable_journal_and_prior_attempt_bytes(setup, monkeypatch, phase):
    from paper_factory import conversion, workflow as workflow_module
    import docx
    import pypdf

    service, runner, research_id, ws = synthetic_verification_export(setup, monkeypatch)
    verify = service._verify
    journals, durable_phases = [], []
    original_fsync = os.fsync

    def interrupted(*args, **kwargs):
        raise RuntimeError("Synthetic verification interruption at " + phase)

    def interrupted_verification(workspace, record, journal):
        journals.append(journal)
        targets = {
            "artifact-source-validation": (workflow_module, "_verify_artifacts"),
            "trusted-analysis": (science, "_compute"),
            "manuscript-rendering": (science, "validate_and_render"),
            "pdf": (pypdf, "PdfReader"),
            "docx": (docx, "Document"),
            "tex-conversion": (conversion, "verify_receipts"),
            "zip": (zipfile, "ZipFile"),
        }
        target, name = targets[phase]
        with pytest.MonkeyPatch.context() as patch:
            patch.setattr(target, name, interrupted)
            return verify(workspace, record, journal)

    def durable_sync(descriptor):
        original_fsync(descriptor)
        for journal in journals:
            if journal.exists():
                opened = os.fstat(descriptor)
                retained = journal.stat()
                if (opened.st_dev, opened.st_ino) == (retained.st_dev, retained.st_ino):
                    rows = verification_rows(journal)
                    assert rows, "The interrupted phase must be flushed before fsync"
                    durable_phases.append(rows[-1]["phase"])

    monkeypatch.setattr(service, "_verify", interrupted_verification)
    monkeypatch.setattr(workflow_module.os, "fsync", durable_sync)
    expected_phases = VERIFICATION_PHASES[:VERIFICATION_PHASES.index(phase) + 1]
    with pytest.raises(RuntimeError, match="Synthetic verification interruption at " + phase):
        service.export(research_id)
    assert [row["phase"] for row in verification_rows(journals[0])] == durable_phases == expected_phases
    first = {path: path.read_bytes() for path in journals[0].parent.rglob("*") if path.is_file()}
    with pytest.raises(RuntimeError, match="Synthetic verification interruption at " + phase):
        service.export(research_id)
    assert journals[0] != journals[1] and journals[0].parent != journals[1].parent
    assert all(path.read_bytes() == content for path, content in first.items())
    assert [row["phase"] for row in verification_rows(journals[1])] == expected_phases
    assert durable_phases == expected_phases * 2
    after = ws.get("workflow", research_id, Workflow)
    assert not any(key.startswith("verification-journal-") for key in after.artifacts)
    assert "validation" not in after.artifacts
    assert after.stage == "manuscript" and after.execution_attempt == 1 and runner.calls == 0


def test_already_exported_verification_adds_only_a_new_journal_and_preserves_completed_outputs(setup, monkeypatch):
    service, runner, research_id, ws = synthetic_verification_export(setup, monkeypatch)
    completed = service.export(research_id)
    original = {key: service.artifact_path(research_id, key).read_bytes() for key in completed["artifacts"]}
    original_metadata = copy.deepcopy(completed["artifacts"])
    validation = json.loads(original["validation"])
    first_journal = validation["verification_journal"]
    repeated = service.export(research_id)
    added = set(repeated["artifacts"]) - set(completed["artifacts"])
    assert len(added) == 1
    new_journal = added.pop()
    assert new_journal.startswith("verification-journal-") and new_journal != first_journal
    path = service.artifact_path(research_id, new_journal)
    assert path.parent.name.startswith("verification-attempt-")
    assert [row["phase"] for row in verification_rows(path)] == VERIFICATION_PHASES
    assert all(service.artifact_path(research_id, key).read_bytes() == content for key, content in original.items())
    assert all(repeated["artifacts"][key] == value for key, value in original_metadata.items())
    assert json.loads(service.artifact_path(research_id, "validation").read_bytes())["verification_journal"] == first_journal
    assert repeated["status"] == "completed" and repeated["stage"] == "exported"
    assert repeated["execution_attempt"] == 1 and runner.calls == 0
    record = ws.get("workflow", research_id, Workflow)
    with pytest.raises(FileExistsError):
        service._verify(ws, record, service.artifact_path(research_id, first_journal))
    assert all(service.artifact_path(research_id, key).read_bytes() == content for key, content in original.items())
    assert set(record.artifacts) == set(repeated["artifacts"])


def test_completed_export_revision_preserves_bytes_and_approval_then_requires_a_new_draft(setup, monkeypatch):
    service, runner, research_id, ws = synthetic_verification_export(setup, monkeypatch)
    completed = service.export(research_id)
    prior = ws.get("workflow", research_id, Workflow).model_dump(mode="json")
    original = {key: (ws.path(item["path"]), ws.path(item["path"]).read_bytes())
                for key, item in prior["artifacts"].items()}
    for name in ("start_experiment", "_execute", "_analyze"):
        monkeypatch.setattr(service, name, lambda *args, **kwargs: pytest.fail("Writing revision dispatched or analyzed science"))
    monkeypatch.setattr(runner, "run", lambda *args, **kwargs: pytest.fail("Writing revision ran the experiment"))
    with monkeypatch.context() as revision_only:
        revision_only.setattr(science, "validate_and_render", lambda *args, **kwargs: pytest.fail("Revision generated a draft"))
        revision_only.setattr(service, "_verify", lambda *args, **kwargs: pytest.fail("Revision re-rendered native exports"))
        result = service.revise_writing(research_id)
    assert (result["status"], result["stage"], result["code"], result["cancellation_requested"]) == ("ready", "analyzed", None, False)
    assert (result["execution_attempt"], result["code_attempt"], result["draft_attempt"], runner.calls) == (1, 1, 1, 0)
    revision = next(key for key in result["artifacts"] if key.startswith("authoring-revision-"))
    receipt = json.loads(service.artifact_path(research_id, revision).read_bytes())
    assert receipt["id"] == revision and receipt["event"] == "authoring-revision" and receipt["at"].endswith("+00:00")
    assert receipt["previous_workflow"] == prior
    assert receipt["previous_status"] == "completed" and receipt["previous_stage"] == "exported"
    assert receipt["execution_attempt"] == 1
    assert all(receipt[name + "_sha256"] == prior["artifacts"][key]["sha256"] for name, key in
               (("execution", "execution"), ("observations", "observations"), ("analysis", "analysis")))
    assert revision in {item["name"] for item in result["material_manifest"]["evidence"]}
    assert json.loads(service.read_material(research_id, "evidence", revision, limit=32000)["text"]) == receipt
    for key, item in prior["artifacts"].items():
        preserved_key = receipt["preserved_artifacts"].get(key, key)
        assert ws.get("workflow", research_id, Workflow).artifacts[preserved_key].model_dump(mode="json") == item
        assert original[key][0].read_bytes() == original[key][1]
    for key in receipt["preserved_artifacts"]:
        assert key not in result["artifacts"]
    with pytest.raises(WorkflowError) as rejected:
        service.export(research_id)
    assert rejected.value.code == "INVALID_STATE"
    with pytest.raises(WorkflowError) as rejected:
        service.submit_code(research_id, BUNDLE, REVIEW)
    assert rejected.value.code == "EXPERIMENT_ALREADY_DISPATCHED"

    revised_draft = manuscript()
    revised_draft["title"] += " revised"
    manuscript_state = service.submit_manuscript(research_id, revised_draft, MANUSCRIPT_REVIEW)
    assert manuscript_state["draft_attempt"] == 2
    fresh_review = json.loads(service.artifact_path(research_id, "manuscript-review").read_bytes())
    old_review = json.loads(original["manuscript-review"][1])
    assert fresh_review["manuscript_sha256"] != old_review["manuscript_sha256"]
    assert fresh_review["manuscript_sha256"] == manuscript_state["artifacts"]["manuscript"]["sha256"]
    second_export = service.export(research_id)
    assert second_export["execution_attempt"] == 1 and second_export["draft_attempt"] == 2 and runner.calls == 0
    with zipfile.ZipFile(service.artifact_path(research_id, "reproducibility")) as archive:
        names = archive.namelist()
        assert archive.testzip() is None and len(names) == len(set(names))
        assert not any(name.endswith(".zip") for name in names)
        assert json.loads(archive.read("authoring/" + revision + ".json")) == receipt
        inventory = json.loads(archive.read("inventory.json"))
        for key, preserved_key in receipt["preserved_artifacts"].items():
            if key == "reproducibility":
                continue
            path, content = original[key]
            name = "authoring/history/" + preserved_key + "/" + path.name
            assert archive.read(name) == content
            assert inventory[name] == {"sha256": prior["artifacts"][key]["sha256"], "size": len(content)}
        for key in prior["artifacts"]:
            if key.startswith(("draft-", "verification-journal-")):
                path, content = original[key]
                assert archive.read("authoring/retained/" + key + "/" + path.name) == content
        for name, item in inventory.items():
            content = archive.read(name)
            assert item == {"sha256": hashlib.sha256(content).hexdigest(), "size": len(content)}
    assert all(path.read_bytes() == content for path, content in original.values())

    second_prior = ws.get("workflow", research_id, Workflow).model_dump(mode="json")
    second_original = {key: (ws.path(item["path"]), ws.path(item["path"]).read_bytes())
                       for key, item in second_prior["artifacts"].items()}
    again = service.revise_writing(research_id)
    assert len([key for key in again["artifacts"] if key.startswith("authoring-revision-")]) == 2
    revised_draft["title"] += " again"
    service.submit_manuscript(research_id, revised_draft, MANUSCRIPT_REVIEW)
    service.export(research_id)
    with zipfile.ZipFile(service.artifact_path(research_id, "reproducibility")) as archive:
        assert archive.testzip() is None and not any(name.endswith(".zip") for name in archive.namelist())
        assert "authoring/" + revision + ".json" in archive.namelist()
    assert all(path.read_bytes() == content for path, content in second_original.values())
    assert all(path.read_bytes() == content for path, content in original.values())
    assert service.status(research_id)["execution_attempt"] == 1 and runner.calls == 0


@pytest.mark.parametrize("status,stage", [
    ("ready", "analyzed"), ("cancelled", "analyzed"), ("ready", "manuscript"),
    ("failed", "exported"), ("blocked", "exported"), ("running", "exported"), ("completed", "analyzed"),
])
def test_writing_revision_rejects_other_states_without_mutation(setup, monkeypatch, status, stage):
    service, runner, research_id, ws = synthetic_verification_export(setup, monkeypatch)
    service.export(research_id)
    record = ws.get("workflow", research_id, Workflow)
    record.status, record.stage = status, stage
    ws.save("workflow", record)
    prior = record.model_dump(mode="json")
    files = {path: path.read_bytes() for path in ws.path("research").rglob("*") if path.is_file()}
    with pytest.raises(WorkflowError) as rejected:
        service.revise_writing(research_id)
    assert rejected.value.code == "INVALID_STATE"
    assert ws.get("workflow", research_id, Workflow).model_dump(mode="json") == prior
    assert all(path.read_bytes() == content for path, content in files.items()) and runner.calls == 0


@pytest.mark.parametrize("defect,expected_code", [
    ("terminal_control", "CONTROL_FAILED"), ("active_handle", "CLEANUP_UNCONFIRMED"),
    ("cleanup", "CLEANUP_UNCONFIRMED"), ("failed_execution", "EXPERIMENT_FAILED"),
    ("production_gate", "PRODUCTION_EXECUTION_UNVERIFIED"), ("failed_control", "CONTROL_FAILED"),
    ("missing_negative", "CONTROL_FAILED"), ("analysis_link", "ARTIFACT_CHANGED"),
    ("tampered_raw", "ARTIFACT_CHANGED"), ("tampered_source", "ARTIFACT_CHANGED"),
    ("missing_validation", "ARTIFACT_MISSING"), ("failed_validation", "REVISION_UNVERIFIED"),
    ("incomplete_journal", "REVISION_UNVERIFIED"), ("review_binding", "ARTIFACT_CHANGED"),
    ("rejected_review", "REVIEW_REJECTED"), ("receipt_collision", "REVISION_EVIDENCE_CONFLICT"),
])
def test_writing_revision_rejects_unsafe_or_changed_completed_evidence_before_mutation(setup, monkeypatch, defect, expected_code):
    from paper_factory import workflow as workflow_module
    service, runner, research_id, ws = synthetic_verification_export(setup, monkeypatch)
    service.export(research_id)
    record = ws.get("workflow", research_id, Workflow)
    if defect == "terminal_control":
        record.terminal_control_failure = True
    elif defect == "active_handle":
        record.active_handle = {"kind": "synthetic", "pid": 123}
    elif defect == "missing_validation":
        del record.artifacts["validation"]
    elif defect == "tampered_source":
        source = ws.path("source/transform.js")
        source.chmod(0o644)
        source.write_text("Synthetic changed source, never executed.", encoding="utf-8")
    elif defect == "receipt_collision":
        monkeypatch.setattr(workflow_module, "uid", lambda prefix: prefix + "-abcdefabcdef")
        ws.path("research/authoring-revision-abcdefabcdef.json").write_bytes(b"Retained collision fixture.")
    else:
        key = ("observations" if defect in {"failed_control", "missing_negative", "tampered_raw"} else
               "analysis" if defect == "analysis_link" else "validation" if defect == "failed_validation" else
               "manuscript-review" if defect in {"review_binding", "rejected_review"} else "execution")
        if defect == "incomplete_journal":
            key = json.loads(service.artifact_path(research_id, "validation").read_bytes())["verification_journal"]
        path = service.artifact_path(research_id, key)
        if defect == "incomplete_journal":
            path.write_text("\n".join(path.read_text(encoding="utf-8").splitlines()[:-1]) + "\n", encoding="utf-8")
        else:
            value = json.loads(path.read_bytes())
            if defect == "cleanup":
                value["cleanup_confirmed"] = False
            elif defect == "failed_execution":
                value["status"] = "cancelled"
            elif defect == "production_gate":
                value["production_calls"] = []
            elif defect == "failed_control":
                value["controls"][1]["passed"] = False
            elif defect == "missing_negative":
                value["controls"] = value["controls"][:1]
            elif defect == "analysis_link":
                value["raw_sha256"] = "0" * 64
            elif defect == "failed_validation":
                value["passed"] = False
            elif defect == "review_binding":
                value["manuscript_sha256"] = "0" * 64
            elif defect == "rejected_review":
                value["review"]["accepted"] = False
                value["review"]["issues"] = ["Synthetic manuscript quality rejection."]
            else:
                value["observations"][0]["value"] = 99
            write_json(path, value)
        if defect != "tampered_raw":
            artifact_path = path.relative_to(ws.root).as_posix()
            changed = FrozenArtifact(path=artifact_path, sha256=digest_file(path), size=path.stat().st_size)
            for alias, artifact in list(record.artifacts.items()):
                if artifact.path == artifact_path:
                    record.artifacts[alias] = changed
    ws.save("workflow", record)
    prior = record.model_dump(mode="json")
    files = {path: path.read_bytes() for path in ws.path("research").rglob("*") if path.is_file()}
    with pytest.raises((WorkflowError, ValueError)) as rejected:
        service.revise_writing(research_id)
    if defect == "tampered_source":
        assert "Source snapshot was modified" in str(rejected.value)
    else:
        assert rejected.value.code == expected_code
    after = ws.get("workflow", research_id, Workflow)
    assert after.model_dump(mode="json") == prior and not any(key.startswith("authoring-revision-") for key in after.artifacts)
    assert all(path.read_bytes() == content for path, content in files.items()) and runner.calls == 0
    if defect == "active_handle":
        after.active_handle = {}
        ws.save("workflow", after)


@pytest.mark.parametrize("failure_phase", ["draft-render", "native-export"])
def test_failed_revised_authoring_preserves_all_previous_completed_bytes(setup, monkeypatch, failure_phase):
    from paper_factory import conversion
    service, runner, research_id, ws = synthetic_verification_export(setup, monkeypatch)
    service.export(research_id)
    prior = ws.get("workflow", research_id, Workflow).model_dump(mode="json")
    original = {key: (ws.path(item["path"]), ws.path(item["path"]).read_bytes())
                for key, item in prior["artifacts"].items()}
    revised = service.revise_writing(research_id)
    revision = next(key for key in revised["artifacts"] if key.startswith("authoring-revision-"))
    receipt = json.loads(service.artifact_path(research_id, revision).read_bytes())
    def interrupted(*args, **kwargs):
        raise RuntimeError("Synthetic revised authoring interruption")
    if failure_phase == "draft-render":
        monkeypatch.setattr(science, "validate_and_render", interrupted)
        with pytest.raises(RuntimeError, match="Synthetic revised"):
            service.submit_manuscript(research_id, manuscript(), MANUSCRIPT_REVIEW)
    else:
        service.submit_manuscript(research_id, manuscript(), MANUSCRIPT_REVIEW)
        monkeypatch.setattr(conversion, "convert", interrupted)
        with pytest.raises(RuntimeError, match="Synthetic revised"):
            service.export(research_id)
    current = ws.get("workflow", research_id, Workflow)
    for key, item in prior["artifacts"].items():
        preserved = receipt["preserved_artifacts"].get(key, key)
        assert current.artifacts[preserved].model_dump(mode="json") == item
        assert original[key][0].read_bytes() == original[key][1]
    assert current.execution_attempt == 1 and runner.calls == 0


def test_failed_export_attempts_preserve_partial_files_without_rerunning_science(setup, monkeypatch):
    from paper_factory import conversion

    service, runner, research_id = setup
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert finished(service, research_id)["stage"] == "analyzed"
    service.submit_manuscript(research_id, manuscript(), MANUSCRIPT_REVIEW)
    before = service.status(research_id)
    outputs = []

    def interrupted_conversion(markdown, output, **kwargs):
        outputs.append(output)
        output.write_bytes(b"synthetic interrupted native output " + str(len(outputs)).encode())
        raise TimeoutError("Synthetic export interruption after writing partial bytes")

    monkeypatch.setattr(conversion, "convert", interrupted_conversion)
    with pytest.raises(TimeoutError, match="Synthetic export"):
        service.export(research_id)
    first = {path: digest_file(path) for path in outputs[0].parent.iterdir() if path.is_file()}
    with pytest.raises(TimeoutError, match="Synthetic export"):
        service.export(research_id)
    assert outputs[0].parent != outputs[1].parent
    assert all(digest_file(path) == digest for path, digest in first.items())
    after = service.status(research_id)
    attempts = [key for key in after["artifacts"] if key.startswith("export-attempt-")]
    assert len(attempts) == 2
    for key in attempts:
        receipt = json.loads(service.artifact_path(research_id, key).read_bytes())
        assert receipt["event"] == "export-started"
        assert receipt["execution_attempt"] == 1
        assert receipt["manuscript_sha256"] == before["artifacts"]["manuscript"]["sha256"]
        assert receipt["execution_sha256"] == before["artifacts"]["execution"]["sha256"]
    assert after["artifacts"]["observations"] == before["artifacts"]["observations"]
    assert after["stage"] == "manuscript" and after["status"] == "ready"
    assert after["execution_attempt"] == runner.calls == 1


def test_artifact_resolution_never_accepts_host_paths(setup):
    service, runner, research_id = setup
    for invalid in ("../outside", "C:/Users", "research-../escape"):
        with pytest.raises(ValueError):
            service.status(invalid)
    with pytest.raises(WorkflowError, match="missing"):
        service.artifact_path(research_id, "../../private")


def test_restart_finishes_retained_success_without_favorable_rerun(setup):
    service, runner, research_id = setup
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert finished(service, research_id)["stage"] == "analyzed"
    service.close()
    ws = Workspace(service.home / "workflows" / research_id)
    record = ws.get("workflow", research_id, Workflow)
    record.status, record.stage = "running", "execute"
    # Simulate a hard interruption after receipt/observations checkpoint and
    # before committing analysis. Existing uncheckpointed analysis is retained.
    record.artifacts = {key: value for key, value in record.artifacts.items()
                        if key != "analysis" and not key.startswith("analysis-")}
    ws.save("workflow", record)
    recovered = WorkflowService(service.home, runner=runner, collector=collect)
    try:
        result = recovered.status(research_id)
        assert result["stage"] == "analyzed" and result["status"] == "ready"
        assert runner.calls == 1
        with pytest.raises(WorkflowError) as rejected:
            recovered.start_experiment(research_id)
        assert rejected.value.code == "EXPERIMENT_ALREADY_DISPATCHED"
    finally:
        recovered.close()


def test_restart_preserves_failed_controls_even_before_terminal_checkpoint(setup):
    service, runner, research_id = setup
    runner.outputs["controls"][1]["passed"] = False
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert finished(service, research_id)["code"] == "CONTROL_FAILED"
    service.close()
    ws = Workspace(service.home / "workflows" / research_id)
    record = ws.get("workflow", research_id, Workflow)
    record.status, record.stage, record.terminal_control_failure = "running", "execute", False
    ws.save("workflow", record)
    recovered = WorkflowService(service.home, runner=runner, collector=collect)
    try:
        result = recovered.status(research_id)
        assert result["code"] == "CONTROL_FAILED" and result["terminal_control_failure"]
        with pytest.raises(WorkflowError) as rejected:
            recovered.submit_code(research_id, BUNDLE, REVIEW)
        assert rejected.value.code == "EXPERIMENT_ALREADY_DISPATCHED"
        assert runner.calls == 1
    finally:
        recovered.close()


def test_another_service_does_not_reconcile_a_live_owned_worker(setup):
    service, runner, research_id = setup
    runner.release = threading.Event()
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert runner.entered.wait(2)
    other = WorkflowService(service.home, runner=runner, collector=collect)
    try:
        assert other.status(research_id)["status"] == "running"
        assert runner.calls == 1
    finally:
        other.close()
    service.cancel(research_id)
    assert finished(service, research_id)["status"] == "cancelled"


def test_failed_worker_diagnostics_are_not_public_host_paths(setup):
    service, runner, research_id = setup
    prepare(service, research_id)
    def fail(*args, **kwargs):
        raise OSError("private C:/Users/Private/credentials.json could not be accessed")
    runner.run = fail
    service.start_experiment(research_id)
    result = finished(service, research_id)
    assert result["status"] == "blocked"
    assert "C:/Users" not in result["message"] and "credentials.json" not in result["message"]
    assert "C:/Users" not in result["diagnostics"]["validation"]
    assert "credentials.json" not in result["diagnostics"]["validation"]


@pytest.mark.parametrize("state", ["blocked", "cancelled", "running", "analyzed"])
def test_dispatched_experiment_rejects_code_replacement_and_reexecution_without_mutation(setup, monkeypatch, state):
    service, runner, research_id = setup
    if state == "blocked":
        runner.outputs["observations"].pop()
    elif state in {"cancelled", "running"}:
        runner.release = threading.Event()
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert runner.entered.wait(2)
    if state == "cancelled":
        service.cancel(research_id)
    result = service.status(research_id) if state == "running" else finished(service, research_id)
    assert result["execution_attempt"] == runner.calls == 1
    if state == "analyzed":
        assert result["stage"] == "analyzed" and result["status"] == "ready"
    else:
        assert result["status"] == state
    if state == "blocked":
        diagnostic = result["diagnostics"]["validation"]
        assert "row" in diagnostic.lower()
        assert "C:/Users" not in diagnostic and str(service.home) not in diagnostic
        assert "do not replace its code or rerun this study" in result["instructions"]

    ws = Workspace(service.root / research_id)
    before = ws.get("workflow", research_id, Workflow).model_dump(mode="json")
    artifact_bytes = {key: service.artifact_path(research_id, key).read_bytes() for key in before["artifacts"]}
    research_files = {path.relative_to(ws.root).as_posix(): path.read_bytes()
                      for path in ws.path("research").rglob("*") if path.is_file()}

    def forbidden(*args, **kwargs):
        pytest.fail("A dispatched experiment must be rejected before state validation or runtime access")

    monkeypatch.setattr(service, "_require", forbidden)
    monkeypatch.setattr(runner, "status", forbidden)
    for operation in (lambda: service.submit_code(research_id, copy.deepcopy(BUNDLE), REVIEW),
                      lambda: service.start_experiment(research_id)):
        with pytest.raises(WorkflowError) as rejected:
            operation()
        assert rejected.value.code == "EXPERIMENT_ALREADY_DISPATCHED"
        assert ws.get("workflow", research_id, Workflow).model_dump(mode="json") == before
        assert service.status(research_id).get("diagnostics") == result.get("diagnostics")
        assert all(service.artifact_path(research_id, key).read_bytes() == content for key, content in artifact_bytes.items())
        assert {path.relative_to(ws.root).as_posix(): path.read_bytes()
                for path in ws.path("research").rglob("*") if path.is_file()} == research_files
        assert runner.calls == 1


def test_contract_instructions_cover_measurements_and_manuscript_requirements(setup):
    service, runner, research_id = setup
    created = service.status(research_id)
    assert "ResearchPlan" in created["instructions"]
    planned = approve_study(service, research_id)
    assert "returned observations envelope" in planned["instructions"]
    assert "intentional" in planned["instructions"] and "base64" in planned["instructions"]
    service.collect_literature(research_id)
    service.submit_code(research_id, BUNDLE, REVIEW)
    service.start_experiment(research_id)
    analyzed = finished(service, research_id)
    assert "padded" in analyzed["instructions"] and "twelve hundred" not in analyzed["instructions"]
    assert "{{result:key}}" in analyzed["instructions"] and "{{citation:id}}" in analyzed["instructions"]
    for heading in science.REQUIRED_SECTIONS:
        assert heading in analyzed["instructions"]


def test_incomplete_import_cannot_break_service_recovery(setup):
    service, runner, research_id = setup
    Workspace.create(service.root / "research-000000000001")
    other = WorkflowService(service.home, runner=runner, collector=collect)
    try:
        assert other.status(research_id)["stage"] == "created"
    finally:
        other.close()


def test_status_polling_can_exclude_materials(setup):
    service, runner, research_id = setup
    prepare(service, research_id)
    compact = service.status(research_id, include_materials=False)
    assert compact["id"] == research_id and compact["stage"] == "code_ready"
    assert set(compact["artifacts"]) >= {"context", "plan", "bundle", "literature"}
    assert not set(compact) & {"schemas", "source_context", "instructions", "plan", "literature", "analysis", "execution", "active_handle", "material_manifest"}
    assert set(service.status(research_id)) >= {"schemas", "source_context", "instructions", "plan", "literature"}


def test_retained_metadata_only_literature_never_becomes_usable_on_repeat(setup):
    service, runner, research_id = setup
    service.submit_proposal(research_id, protocol())
    def metadata_only(*args, **kwargs):
        result = collect(*args, **kwargs)
        result["sources"][0]["scope"] = "metadata_only"
        result["sources"][0]["excerpts"] = []
        return result
    service.collector = metadata_only
    for _ in range(2):
        with pytest.raises(WorkflowError, match="metadata|Metadata"):
            service.collect_literature(research_id)
    assert "literature" in service.status(research_id)["artifacts"]


def test_partial_literature_completes_missing_query_without_changing_retained_bytes(setup):
    service, runner, research_id = setup
    plan = protocol()
    plan["literature_queries"] = ["first frozen query", "second frozen query"]
    service.submit_proposal(research_id, plan)
    calls = []

    def partial_then_complete(queries, root, *, limit, cancel):
        calls.append((queries, limit))
        if len(calls) == 1:
            evidence = collect(queries[:1], root, limit=limit, cancel=cancel)
            evidence["sources"] = [dict(evidence["sources"][0], id=f"retained-source-{index}") for index in range(6)]
            return evidence
        assert queries == ["second frozen query"] and limit == 0
        path = Path(root) / "literature" / "second-query-search.json"
        write_json(path, {"status": "ok", "message": {"items": []}})
        return {"sources": [], "searches": [{"query": queries[0], "provider": "Crossref", "status": "succeeded",
                    "attempted": True, "resolved_ids": [], "raw_path": "literature/second-query-search.json",
                    "sha256": digest_file(path)}], "warnings": []}

    service.collector = partial_then_complete
    with pytest.raises(WorkflowError) as first:
        service.collect_literature(research_id)
    assert first.value.code == "LITERATURE_QUERIES_INCOMPLETE"
    original = service.artifact_path(research_id, "literature")
    original_bytes = original.read_bytes()
    original_digest = digest_file(original)
    retained = json.loads(original_bytes)
    result = service.collect_literature(research_id)
    current = service.artifact_path(research_id, "literature")
    assert current != original and original.read_bytes() == original_bytes
    assert result["literature"]["sources"] == retained["sources"]
    assert result["literature"]["searches"][:1] == retained["searches"]
    assert len(result["literature"]["sources"]) == 6 and len(result["literature"]["searches"]) == 2
    history = result["literature"]["history"][0]
    assert history["sha256"] == original_digest
    assert service.artifact_path(research_id, history["artifact_id"]) == original
    assert result["artifacts"][history["artifact_id"]]["sha256"] == original_digest
    search = result["literature"]["searches"][1]
    bindings = [key for key, artifact in result["artifacts"].items() if artifact["sha256"] == search["sha256"] and
                service.artifact_path(research_id, key).relative_to(service.root / research_id).as_posix() == "research/" + search["raw_path"]]
    assert len(bindings) == 1
    service.collect_literature(research_id)
    assert len(calls) == 2 and runner.calls == 0
    assert current.read_bytes() == service.artifact_path(research_id, "literature").read_bytes()


def test_metadata_only_partial_collection_does_not_exhaust_later_inspected_source_budget(setup):
    service, runner, research_id = setup
    proposal = protocol()
    proposal["literature_queries"] = ["metadata query", "readable query"]
    service.submit_proposal(research_id, proposal)
    calls = []
    def metadata_then_readable(queries, root, *, limit, cancel):
        calls.append((queries, limit))
        evidence = collect(queries[:1], root, limit=limit, cancel=cancel)
        if len(calls) == 1:
            evidence["sources"] = [{**evidence["sources"][0], "id": f"metadata-{index}",
                                    "scope": "metadata_only", "excerpts": []} for index in range(6)]
        return evidence
    service.collector = metadata_then_readable
    with pytest.raises(WorkflowError) as insufficient:
        service.collect_literature(research_id)
    assert insufficient.value.code == "LITERATURE_EVIDENCE_INSUFFICIENT"
    original = service.artifact_path(research_id, "literature")
    retained = original.read_bytes()
    completed = service.collect_literature(research_id)
    assert calls == [(["metadata query", "readable query"], 6), (["readable query"], 6)]
    sources = completed["literature"]["sources"]
    assert len(sources) == 6 and sources[0]["id"] == "fixture-oracle" and sources[0]["scope"] == "full_text"
    assert all(source["scope"] == "metadata_only" for source in sources[1:])
    assert original.read_bytes() == retained
    approved = service.submit_study_review(research_id, STUDY_REVIEW)
    assert approved["stage"] == "planned" and runner.calls == 0


@pytest.mark.parametrize("incomplete_flag", [None, "cancelled", "timed_out", "rate_limited"])
@pytest.mark.parametrize("old_scope,new_scope,retained_scope", [
    ("metadata_only", "abstract", "abstract"),
    ("metadata_only", "full_text", "full_text"),
    ("abstract", "full_text", "full_text"),
    ("abstract", "metadata_only", "abstract"),
    ("full_text", "abstract", "full_text"),
    ("full_text", "metadata_only", "full_text"),
    ("abstract", "abstract", "abstract"),
])
def test_same_source_recovery_upgrades_reading_without_downgrading_or_overwriting_history(
        setup, incomplete_flag, old_scope, new_scope, retained_scope):
    service, runner, research_id = setup
    proposal = protocol()
    proposal["literature_queries"] = ["first query", "second query"]
    service.submit_proposal(research_id, proposal)
    calls = []

    def recover_reading(queries, root, *, limit, cancel):
        calls.append((list(queries), limit))
        evidence = collect(queries if incomplete_flag or len(calls) > 1 else queries[:1], root, limit=6, cancel=cancel)
        source = evidence["sources"][0]
        scope = old_scope if len(calls) == 1 else new_scope
        source["scope"] = scope
        if scope == "metadata_only":
            source["excerpts"] = []
        raw_path = f"literature/recovered-source-{len(calls)}.json"
        write_json(Path(root) / raw_path, {"scope": scope, "title": source["title"], "excerpts": source["excerpts"]})
        source.update(raw_path=raw_path, sha256=digest_file(Path(root) / raw_path))
        if len(calls) == 1 and incomplete_flag:
            evidence[incomplete_flag] = True
        return evidence

    service.collector = recover_reading
    expected_error = "LITERATURE_EVIDENCE_INSUFFICIENT" if incomplete_flag or old_scope == "metadata_only" else "LITERATURE_QUERIES_INCOMPLETE"
    with pytest.raises(WorkflowError) as partial:
        service.collect_literature(research_id)
    assert partial.value.code == expected_error
    original = service.artifact_path(research_id, "literature")
    prior_bytes = original.read_bytes()
    original_source = json.loads(prior_bytes)["sources"][0]
    prior_artifact = service.root / research_id / "research" / original_source["raw_path"]
    prior_source_bytes = prior_artifact.read_bytes()
    if incomplete_flag:
        # Every query was attempted, but interrupted DOI reading cannot be approved.
        assert all(search["attempted"] for search in json.loads(prior_bytes)["searches"])
        assert len(json.loads(prior_bytes)["searches"]) == 2
        before = service.status(research_id)["artifacts"]
        with pytest.raises(WorkflowError) as premature:
            service.submit_study_review(research_id, STUDY_REVIEW)
        assert premature.value.code == "LITERATURE_EVIDENCE_INSUFFICIENT"
        if incomplete_flag in {"timed_out", "rate_limited"}:
            assert ("timed-out" if incomplete_flag == "timed_out" else "rate-limited") in str(premature.value)
        assert service.status(research_id)["artifacts"] == before
        assert "plan" not in before and "study-review" not in before

    completed = service.collect_literature(research_id)
    assert completed["literature"]["cancelled"] is False
    assert completed["literature"]["timed_out"] is False
    assert completed["literature"]["rate_limited"] is False
    sources = completed["literature"]["sources"]
    assert len(sources) == 1 and sources[0]["id"] == "fixture-oracle"
    assert sources[0]["scope"] == retained_scope and sources[0]["excerpts"]
    attempt = 2 if retained_scope != old_scope else 1
    assert sources[0]["raw_path"] == f"literature/recovered-source-{attempt}.json"
    assert calls == [(["first query", "second query"], 6),
                     (["first query", "second query"] if incomplete_flag else ["second query"],
                      6 if old_scope == "metadata_only" else 5)]
    history = completed["literature"]["history"]
    assert len(history) == 1 and history[0]["sha256"] == hashlib.sha256(prior_bytes).hexdigest()
    assert service.artifact_path(research_id, history[0]["artifact_id"]) == original
    assert original.read_bytes() == prior_bytes and prior_artifact.read_bytes() == prior_source_bytes
    assert completed["study_review"] is None and "plan" not in completed["artifacts"]
    invalid_selection = copy.deepcopy(STUDY_REVIEW)
    invalid_selection["selected_sources"][0]["excerpt_index"] = 99
    with pytest.raises(WorkflowError) as uninspected:
        service.submit_study_review(research_id, invalid_selection)
    assert uninspected.value.code == "LITERATURE_SELECTION_INVALID"
    assert "plan" not in service.status(research_id)["artifacts"]
    if retained_scope != "full_text":
        with pytest.raises(WorkflowError) as abstract_positioning:
            service.submit_study_review(research_id, STUDY_REVIEW)
        assert abstract_positioning.value.code == "PUBLICATION_EVIDENCE_INVALID" and runner.calls == 0
        return
    approved = service.submit_study_review(research_id, STUDY_REVIEW)
    assert approved["stage"] == "planned" and runner.calls == 0
    receipt = json.loads(service.artifact_path(research_id, "study-review").read_bytes())
    assert receipt["literature_sha256"] == approved["artifacts"]["literature"]["sha256"]
    selected = json.loads(service.artifact_path(research_id, "selected-literature").read_bytes())
    assert selected["sources"][0]["raw_path"] == sources[0]["raw_path"]


@pytest.mark.parametrize("defect,message", [
    ("short_excerpt", "insufficient inspected text"),
    ("mismatched_digest", "Literature differs from its retrieval digest"),
])
def test_same_source_upgrade_requires_substantive_reading_and_matching_artifact_before_commit(setup, defect, message):
    service, runner, research_id = setup
    proposal = protocol()
    proposal["literature_queries"] = ["first query", "second query"]
    service.submit_proposal(research_id, proposal)
    calls = []

    def unverified_upgrade(queries, root, *, limit, cancel):
        calls.append(list(queries))
        evidence = collect(queries[:1] if len(calls) == 1 else queries, root, limit=6, cancel=cancel)
        if len(calls) == 1:
            evidence["sources"][0]["scope"] = "abstract"
        if len(calls) == 2:
            source = evidence["sources"][0]
            source["scope"] = "full_text"
            if defect == "short_excerpt":
                source["excerpts"] = ["Too short to substantiate a claim."]
            else:
                source["sha256"] = "0" * 64
        return evidence

    service.collector = unverified_upgrade
    with pytest.raises(WorkflowError) as partial:
        service.collect_literature(research_id)
    assert partial.value.code == "LITERATURE_QUERIES_INCOMPLETE"
    original = service.artifact_path(research_id, "literature")
    prior_bytes = original.read_bytes()
    before = service.status(research_id)["artifacts"]
    with pytest.raises(ValueError, match=message):
        service.collect_literature(research_id)
    assert original.read_bytes() == prior_bytes
    assert service.artifact_path(research_id, "literature") == original
    assert service.status(research_id)["artifacts"] == before
    assert runner.calls == 0


def test_failed_attempt_completes_query_but_unattempted_error_remains_missing(setup):
    service, runner, research_id = setup
    plan = protocol()
    plan["literature_queries"] = ["first frozen query", "second frozen query"]
    service.submit_proposal(research_id, plan)
    calls = []

    def collector(queries, root, *, limit, cancel):
        calls.append(queries)
        if len(calls) == 1:
            evidence = collect(queries[:1], root, limit=limit, cancel=cancel)
            evidence["searches"].append({"query": queries[1], "status": "failed", "attempted": False, "error": "OSError"})
            return evidence
        assert queries == ["second frozen query"] and limit == 5
        return {"sources": [], "searches": [{"query": queries[0], "status": "failed", "attempted": True,
                                                "error": "HTTPStatusError"}], "warnings": ["Observed provider failure"]}

    service.collector = collector
    with pytest.raises(WorkflowError) as initial:
        service.collect_literature(research_id)
    assert initial.value.code == "LITERATURE_QUERIES_INCOMPLETE"
    prior = service.artifact_path(research_id, "literature").read_bytes()
    result = service.collect_literature(research_id)
    assert result["literature"]["searches"][:2] == json.loads(prior)["searches"]
    assert result["literature"]["searches"][-1]["error"] == "HTTPStatusError"
    assert result["literature"]["sources"] == json.loads(prior)["sources"]
    service.collect_literature(research_id)
    assert len(calls) == 2 and runner.calls == 0


def test_explicit_research_id_does_not_change_global_workspace_selection(setup):
    service, runner, research_id = setup
    pointer = service.home / "current.json"
    assert not pointer.exists()
    write_json(pointer, {"workspace": "existing-legacy-selection"})
    original = pointer.read_bytes()
    service.create(str(service.home.parent / "input"), "Validate another bounded controlled transformation study.")
    assert pointer.read_bytes() == original


def test_recovered_unconfirmed_worker_cannot_be_hidden_by_shutdown(setup):
    service, runner, research_id = setup
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert finished(service, research_id)["stage"] == "analyzed"
    service.close()
    ws = Workspace(service.home / "workflows" / research_id)
    record = ws.get("workflow", research_id, Workflow)
    record.status, record.code = "blocked", "CLEANUP_UNCONFIRMED"
    record.active_handle = {"kind": "fixture", "pid": 123}
    ws.save("workflow", record)
    runner.cleanable = False
    recovered = WorkflowService(service.home, runner=runner, collector=collect)
    assert recovered.status(research_id, include_materials=False)["code"] == "CLEANUP_UNCONFIRMED"
    with pytest.raises(WorkflowError, match="cleanup"):
        recovered.close()
    assert ws.get("workflow", research_id, Workflow).active_handle["pid"] == 123
    runner.cleanable = True
    recovered.close()
    record = ws.get("workflow", research_id, Workflow)
    assert not record.active_handle and record.stage == "analyzed"


def test_material_reading_returns_declared_source_and_current_experiment(setup):
    service, runner, research_id = setup
    source = service.read_material(research_id, "source", "transform.js", limit=10)
    assert source["text"] == "export fun" and source["offset"] == 0 and source["next_offset"] == 10
    assert source["is_untrusted_data"] is True
    remainder = service.read_material(research_id, "source", "transform.js", offset=10, limit=32000)
    text = source["text"] + remainder["text"]
    assert hashlib.sha256(text.encode()).hexdigest() == source["sha256"] == remainder["sha256"]
    assert len(text) == source["total_chars"] and remainder["next_offset"] is None
    prepare(service, research_id)
    code = service.read_material(research_id, "experiment", "experiment.mjs")
    assert code["text"].replace("\r\n", "\n") == BUNDLE["files"][0]["content"]
    materials = service.status(research_id)["material_manifest"]
    assert materials["source"][0]["name"] == "transform.js"
    assert materials["experiment"][0]["name"] == "experiment.mjs"
    assert materials["experiment"][0]["sha256"] == code["sha256"]


def test_reviewer_can_read_bounded_hash_verified_runtime_and_numbered_code_review(setup):
    service, runner, research_id = setup
    prepare(service, research_id)
    ws = Workspace(service.root / research_id)
    record = ws.get("workflow", research_id, Workflow)
    path = ws.path("research/runtime-manifest.json")
    write_json(path, {"compiled_source": "trusted transformer fixture " * 200, "transformer_version": "fixture"})
    record.artifacts["runtime-manifest"] = FrozenArtifact(path=path.relative_to(ws.root).as_posix(),
                                                          sha256=digest_file(path), size=path.stat().st_size)
    ws.save("workflow", record)
    for name in ("runtime-manifest", "code-review-1"):
        expected = service.artifact_path(research_id, name)
        page = service.read_material(research_id, "evidence", name, limit=19)
        assert len(page["text"]) == 19 and page["sha256"] == digest_file(expected) and page["next_offset"] == 19
        assert page["is_untrusted_data"] is True
        entry = next(item for item in service.status(research_id)["material_manifest"]["evidence"] if item["name"] == name)
        assert entry["sha256"] == page["sha256"] and entry["size"] == expected.stat().st_size
    assert runner.calls == 0
    path.write_text("{}", encoding="utf-8")
    with pytest.raises(WorkflowError) as changed:
        service.read_material(research_id, "evidence", "runtime-manifest")
    assert changed.value.code == "ARTIFACT_CHANGED"


@pytest.mark.parametrize("name", ["code-review", "code-review-0", "code-review-01", "code-review-1/../plan", "code-review--1"])
def test_numbered_code_review_allowlist_rejects_other_keys(setup, name):
    service, runner, research_id = setup
    prepare(service, research_id)
    with pytest.raises(WorkflowError) as rejected:
        service.read_material(research_id, "evidence", name)
    assert rejected.value.code == "MATERIAL_NOT_DECLARED"


def test_material_reading_exposes_actual_runner_raw_evidence_for_fresh_review(setup):
    service, runner, research_id = setup
    prepare(service, research_id)
    service.start_experiment(research_id)
    assert finished(service, research_id)["stage"] == "analyzed"
    raw = service.read_material(research_id, "evidence", "observations")
    assert json.loads(raw["text"]) == runner.outputs
    assert raw["sha256"] == service.status(research_id)["artifacts"]["observations"]["sha256"]
    service.submit_manuscript(research_id, manuscript(), MANUSCRIPT_REVIEW)
    assert service.read_material(research_id, "evidence", "manuscript")["text"].startswith("# ")
    canonical = service.read_material(research_id, "evidence", "canonical")
    fragments = [canonical["text"]]
    original_hash = canonical["sha256"]
    while canonical["next_offset"] is not None:
        canonical = service.read_material(research_id, "evidence", "canonical", offset=canonical["next_offset"])
        assert canonical["sha256"] == original_hash
        fragments.append(canonical["text"])
    assert json.loads("".join(fragments))["title"] == manuscript()["title"]


@pytest.mark.parametrize("area,name", [
    ("source", "../records.sqlite3"), ("source", "C:/Users/private.txt"), ("source", "/private.txt"),
    ("source", ".env"), ("source", "absent.py"), ("experiment", "../protocol.json"),
    ("experiment", "bundle.json"), ("experiment", "review.json"),
    ("evidence", "execution"), ("evidence", "active_handle"), ("evidence", "../../records.sqlite3"),
    ("other", "transform.js"),
])
def test_material_reading_refuses_host_paths_and_undeclared_files(setup, area, name):
    service, runner, research_id = setup
    prepare(service, research_id)
    with pytest.raises(ValueError):
        service.read_material(research_id, area, name)


@pytest.mark.parametrize("offset,limit", [(-1, 10), (True, 10), (0, False), (0, 0), (0, 32001), (0.5, 10), (0, "10")])
def test_material_paging_has_exact_bounded_integer_parameters(setup, offset, limit):
    service, runner, research_id = setup
    with pytest.raises(ValueError, match="paging"):
        service.read_material(research_id, "source", "transform.js", offset=offset, limit=limit)


@pytest.mark.parametrize("data", [b"\xff\xfe", b"text\x00binary"])
def test_material_reader_refuses_binary_files(setup, data):
    service, runner, research_id = setup
    source = service.home.parent / "binary-input"
    source.mkdir()
    (source / "asset.bin").write_bytes(data)
    created = service.create(str(source), "Inspect the declared binary fixture without reading private host files.")
    with pytest.raises(WorkflowError, match="UTF-8|Binary"):
        service.read_material(created["id"], "source", "asset.bin")


def test_material_reader_checks_all_frozen_artifacts_before_reading(setup):
    service, runner, research_id = setup
    prepare(service, research_id)
    protocol_file = service.artifact_path(research_id, "plan")
    protocol_file.write_text("{}", encoding="utf-8")
    with pytest.raises(WorkflowError, match="changed"):
        service.read_material(research_id, "source", "transform.js")


def test_material_paging_counts_unicode_characters_not_bytes(setup):
    service, runner, research_id = setup
    source = service.home.parent / "unicode-input"
    source.mkdir()
    text = "한글🙂readable"
    (source / "note.txt").write_text(text, encoding="utf-8")
    created = service.create(str(source), "Inspect a declared Unicode text fixture with bounded character pages.")
    page = service.read_material(created["id"], "source", "note.txt", offset=1, limit=2)
    assert page["text"] == "글🙂" and page["next_offset"] == 3 and page["total_chars"] == len(text)
    assert page["sha256"] == hashlib.sha256(text.encode()).hexdigest()


def test_material_unicode_paging_preserves_characters_across_stream_chunks(setup):
    service, runner, research_id = setup
    source = service.home.parent / "chunked-unicode-input"
    source.mkdir()
    text = "a" * 65535 + "한글🙂" + "tail" * 10
    (source / "note.txt").write_text(text, encoding="utf-8")
    created = service.create(str(source), "Inspect bounded Unicode material across decoder chunk boundaries.")
    page = service.read_material(created["id"], "source", "note.txt", offset=65534, limit=4)
    assert page["text"] == "a한글🙂" and page["next_offset"] == 65538
    assert page["total_chars"] == len(text) and page["sha256"] == hashlib.sha256(text.encode()).hexdigest()
