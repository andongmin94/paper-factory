from pathlib import Path
import hashlib
import json
import subprocess
from types import SimpleNamespace
import zipfile

import pytest

from paper_factory import conversion
from paper_factory.workspace import digest_file


@pytest.fixture
def source(tmp_path):
    path = tmp_path / "article.md"
    path.write_text("# Export contract\n\n## Results\n\nThe document preserves its verified prose and visible section structure.\n", encoding="utf-8")
    return path




@pytest.mark.parametrize("enabled", [False, True])
def test_docx_export_uses_native_line_and_page_settings(source, pandoc, enabled):
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn
    output = source.with_suffix(".docx")
    report = conversion.convert(source, output, pandoc=pandoc, line_numbers=enabled, page_numbers=enabled)
    document = docx.Document(output)
    assert "Export contract" in " ".join(paragraph.text for paragraph in document.paragraphs)
    for section in document.sections:
        numbers = section._sectPr.findall(qn("w:lnNumType"))
        assert bool(numbers) is enabled
        if enabled:
            assert numbers[0].get(qn("w:restart")) == "continuous"
            assert numbers[0].get(qn("w:countBy")) == "1"
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            fields = [field for field in footer._element.iter(qn("w:fldSimple")) if field.get(qn("w:instr")) == "PAGE"]
            assert bool(fields) is enabled
    with zipfile.ZipFile(output) as archive:
        assert "word/document.xml" in archive.namelist()
    assert report["output_sha256"] == digest_file(output)


@pytest.mark.parametrize("enabled", [False, True])
def test_latex_export_uses_native_line_and_page_settings(source, pandoc, enabled):
    output = source.with_suffix(".tex")
    conversion.convert(source, output, pandoc=pandoc, line_numbers=enabled, page_numbers=enabled)
    text = output.read_text(encoding="utf-8")
    assert "Export contract" in text
    assert (r"\usepackage{lineno}" in text) is enabled
    assert (r"\linenumbers" in text) is enabled
    assert (r"\pagestyle{empty}" in text) is not enabled


@pytest.mark.parametrize("enabled", [False, True])
def test_pdf_export_native_numbering_matches_requested_flags(source, pandoc, enabled):
    pytest.importorskip("typst")
    from pypdf import PdfReader
    output = source.with_suffix(".pdf")
    conversion.convert(source, output, pandoc=pandoc, line_numbers=enabled, page_numbers=enabled)
    typst = output.with_suffix(".typ").read_text(encoding="utf-8")
    assert ('set par.line(numbering: "1"' in typst) is enabled
    assert ('set page(numbering: none)' in typst) is not enabled
    text = " ".join(page.extract_text() or "" for page in PdfReader(output).pages)
    assert "Export contract" in text and "visible section structure" in text
    if not enabled:
        assert not any(character.isdigit() for character in text)


def test_pdf_export_keeps_complete_long_title_in_one_typst_heading(source, pandoc):
    pytest.importorskip("typst")
    from pypdf import PdfReader

    title = "Denominator Weighting and Missing-Day Preservation in Controlled Analytics Ledgers"
    source.write_text(f"# {title}\n\n## Results\n\nThe verified ledger observations remain unchanged.\n", encoding="utf-8")
    original = source.read_bytes()
    output = source.with_suffix(".pdf")
    receipt = conversion.convert(source, output, pandoc=pandoc)

    assert "--wrap=none" in receipt["command"]
    assert source.read_bytes() == original
    generated = output.with_suffix(".typ").read_text(encoding="utf-8")
    assert "= " + title in generated.splitlines()
    # Typst reflows the complete heading visually. A wrapped Typst source line
    # instead turns the trailing title words into ordinary 11 pt body text.
    headings = []

    def collect_heading(text, _cm, _tm, _font, font_size):
        if font_size > 11:
            headings.append(text)

    for page in PdfReader(output).pages:
        page.extract_text(visitor_text=collect_heading)
    heading_text = " ".join("".join(headings).split())
    # Native line-end hyphenation can split other title words. Check the
    # trailing words that were demoted to body text in the original export.
    assert heading_text.startswith("Denominator Weighting")
    assert "Analytics Ledgers" in heading_text
    assert receipt["input_sha256"] == hashlib.sha256(original).hexdigest()
    assert receipt["output_sha256"] == digest_file(output)


@pytest.mark.parametrize("extra_fonts", [False, True])
def test_pdf_export_preserves_mixed_korean_text_with_bundled_fonts(source, pandoc, monkeypatch, extra_fonts):
    pytest.importorskip("typst")
    from pypdf import PdfReader

    if extra_fonts:
        additional_fonts = source.parent / "runtime-fonts"
        additional_fonts.mkdir()
        monkeypatch.setenv("TYPST_FONT_PATHS", str(additional_fonts))
    else:
        monkeypatch.delenv("TYPST_FONT_PATHS", raising=False)
    source.write_text("# 한국어 재생성 규칙\n\n"
                      "Each seed is initialized. 각 seed에서 초기화한다.\n\n"
                      "**한국어 글꼴 검증을 진행합니다.**\n", encoding="utf-8")
    original = source.read_bytes()
    output = source.with_suffix(".pdf")
    receipt = conversion.convert(source, output, pandoc=pandoc, metadata={"mainfont": "Libertinus Serif"})

    reader = PdfReader(output, strict=True)
    text = " ".join(" ".join(page.extract_text().split()) for page in reader.pages)
    assert "한국어 재생성 규칙" in text
    assert "각 seed에서 초기화한다." in text
    assert "한국어 글꼴 검증을 진행합니다." in text
    assert "Each seed is initialized." in text
    font_names = {str(font.get_object()["/BaseFont"]) for page in reader.pages
                  for font in page["/Resources"]["/Font"].get_object().values()}
    assert any("Pretendard-Regular" in name for name in font_names)
    assert any("Pretendard-Bold" in name for name in font_names)
    assert source.read_bytes() == original
    assert receipt["input_sha256"] == hashlib.sha256(original).hexdigest()
    assert receipt["output_sha256"] == digest_file(output)


@pytest.mark.parametrize("suffix", [".docx", ".tex", ".pdf"])
def test_pipe_tables_convert_to_native_tables_without_editing_source(source, pandoc, suffix):
    source.write_text("# Controlled results\n\nDescriptive statistics.\n\n"
                      "| Metric | Condition | Count | Mean |\n| --- | --- | ---: | ---: |\n"
                      "| accuracy | production | 12 | 1 |\n\nPaired differences.\n\n"
                      "| Metric | Baseline | Pairs | Mean delta |\n| --- | --- | ---: | ---: |\n"
                      "| accuracy | production | 12 | -1 |\n", encoding="utf-8")
    original = source.read_bytes()
    output = source.with_suffix(suffix)
    receipt = conversion.convert(source, output, pandoc=pandoc)
    assert source.read_bytes() == original
    assert receipt["input_sha256"] == hashlib.sha256(original).hexdigest()
    if suffix == ".docx":
        from docx import Document
        tables = Document(output).tables
        assert len(tables) == 2
        assert [cell.text for cell in tables[0].rows[1].cells] == ["accuracy", "production", "12", "1"]
        assert [cell.text for cell in tables[1].rows[1].cells] == ["accuracy", "production", "12", "-1"]
    elif suffix == ".tex":
        text = output.read_text(encoding="utf-8")
        assert text.count(r"\begin{longtable}") == 2
        assert "accuracy & production & 12 & -1" in text
    else:
        from pypdf import PdfReader
        generated = output.with_suffix(".typ").read_text(encoding="utf-8")
        assert generated.count("#table(") == 2
        text = " ".join(page.extract_text() or "" for page in PdfReader(output).pages)
        assert "accuracy" in text and "Mean delta" in text and "---" not in text and "| Metric" not in text
    receipts = source.with_name("receipts.json")
    receipts.write_text(json.dumps({suffix[1:]: receipt}), encoding="utf-8")
    assert conversion.verify_receipts(source, {suffix[1:]: output}, receipts) == []


def test_trusted_statistics_renderer_separates_tables_before_pandoc_parsing(pandoc):
    from paper_factory.autonomous import science
    summaries = [{"metric": "accuracy", "condition": "production", "unit": "fraction",
                  "count": 12, "mean": 1, "median": 1, "stdev": 0, "min": 1, "max": 1}]
    paired = [{"metric": "accuracy", "condition": "ablation", "baseline": "production",
               "count": 12, "mean": -1, "median": -1, "stdev": 0, "min": -1, "max": -1}]
    text = science._tables({"summaries": summaries, "paired_deltas": paired})
    parsed = subprocess.run([pandoc, "--from=markdown-smart-raw_tex-raw_html", "--to=json"],
                            input=text, text=True, capture_output=True, encoding="utf-8", check=True)
    blocks = json.loads(parsed.stdout)["blocks"]
    assert sum(block["t"] == "Table" for block in blocks) == 2
    for index, block in enumerate(blocks):
        if block["t"] == "Table":
            assert blocks[index - 1]["t"] == "Header"
            assert blocks[index - 1]["c"][0] == 3


@pytest.mark.parametrize("condition_count", [2, 8])
def test_statistics_export_keeps_long_metric_names_outside_narrow_numeric_tables(source, pandoc, condition_count):
    from docx import Document
    from paper_factory.autonomous import science

    summaries = [{"metric": "canonicalization_invariance_with_long_identifiers", "condition": f"condition_{index}",
                  "unit": "nanoseconds", "count": 36, "mean": 115465.25, "median": 67037.5,
                  "stdev": 284341.7243, "min": 40880, "max": 1770754} for index in range(condition_count)]
    paired = [{"metric": summaries[0]["metric"], "condition": "condition_1", "baseline": "condition_0",
               "count": 36, "mean": -114053.8611, "median": -65711, "stdev": 284346.453, "min": -1769431, "max": -39989}]
    source.write_text("# Controlled statistics\n\n" + science._tables({"summaries": summaries, "paired_deltas": paired}), encoding="utf-8")
    output = source.with_suffix(".docx")
    conversion.convert(source, output, pandoc=pandoc)
    document = Document(output)
    assert len(document.tables) == condition_count // 2 + 1
    assert all(len(table.columns) <= 3 for table in document.tables)
    assert summaries[0]["metric"] in " ".join(paragraph.text for paragraph in document.paragraphs)
    captions = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith("Metric:")]
    assert len(captions) == 2 and all(paragraph.style.name == "Heading 3" for paragraph in captions)
    assert all(paragraph.style.paragraph_format.keep_with_next for paragraph in captions)
    for table in document.tables[:-1]:
        assert [cell.text for cell in table.rows[2].cells] == ["Mean", "115465.25", "115465.25"]
    assert [cell.text for cell in document.tables[-1].rows[2].cells] == ["Mean delta", "-114053.8611"]


def test_missing_converter_cannot_leave_a_stale_successful_artifact(source, monkeypatch):
    output = source.with_suffix(".tex")
    output.write_text("stale successful export", encoding="utf-8")
    with pytest.raises(ValueError, match="verified absolute path"):
        conversion.convert(source, output, pandoc="missing-explicit-converter")
    assert not output.exists()


@pytest.mark.parametrize("suffix", [".tex", ".pdf"])
def test_timeout_removes_partial_export_and_stale_output(source, monkeypatch, suffix):
    output = source.with_suffix(suffix)
    output.write_bytes(b"stale output")
    monkeypatch.setattr(conversion, "pandoc_binary", lambda _: "test-pandoc")

    def timeout(command, **kwargs):
        Path(command[command.index("-o") + 1]).write_text("partial output", encoding="utf-8")
        raise subprocess.TimeoutExpired(command, 120)

    monkeypatch.setattr(conversion.subprocess, "run", timeout)
    with pytest.raises(ValueError, match="timed out"):
        conversion.convert(source, output)
    assert not output.exists()
    assert not (output.with_suffix(".typ") if suffix == ".pdf" else output).exists()


def test_nonzero_converter_failure_removes_partial_export_and_keeps_diagnostics(source, monkeypatch):
    output = source.with_suffix(".tex")
    monkeypatch.setattr(conversion, "pandoc_binary", lambda _: "test-pandoc")

    def failed(command, **kwargs):
        Path(command[command.index("-o") + 1]).write_text("partial output", encoding="utf-8")
        return SimpleNamespace(returncode=9, stderr="actual parser diagnostic")

    monkeypatch.setattr(conversion.subprocess, "run", failed)
    with pytest.raises(ValueError, match=r"Pandoc failed \(9\): actual parser diagnostic"):
        conversion.convert(source, output)
    assert not output.exists()


def test_invalid_docx_export_is_removed_after_independent_reopen(source, monkeypatch):
    pytest.importorskip("docx")
    output = source.with_suffix(".docx")
    monkeypatch.setattr(conversion, "pandoc_binary", lambda _: "test-pandoc")

    def broken(command, **kwargs):
        Path(command[command.index("-o") + 1]).write_text("successful exit with invalid DOCX", encoding="utf-8")
        return SimpleNamespace(returncode=0, stderr="")

    monkeypatch.setattr(conversion.subprocess, "run", broken)
    with pytest.raises(ValueError, match="Word document formatting failed"):
        conversion.convert(source, output)
    assert not output.exists()


def test_pdf_engine_failure_removes_partial_pdf_and_retains_debug_source(source, monkeypatch):
    typst = pytest.importorskip("typst")
    output = source.with_suffix(".pdf")
    monkeypatch.setattr(conversion, "pandoc_binary", lambda _: "test-pandoc")

    def generated(command, **kwargs):
        Path(command[command.index("-o") + 1]).write_text("Generated debug source", encoding="utf-8")
        return SimpleNamespace(returncode=0, stderr="")

    def failed(path, **kwargs):
        Path(kwargs["output"]).write_bytes(b"partial PDF")
        raise RuntimeError("actual engine diagnostic")

    monkeypatch.setattr(conversion.subprocess, "run", generated)
    monkeypatch.setattr(typst, "compile", failed)
    with pytest.raises(ValueError, match="actual engine diagnostic"):
        conversion.convert(source, output)
    assert not output.exists()
    assert "Generated debug source" in output.with_suffix(".typ").read_text(encoding="utf-8")








def test_converter_requires_explicit_absolute_path_without_discovery(source, monkeypatch, pandoc):
    from docx import Document
    monkeypatch.delenv("PYPANDOC_PANDOC", raising=False)
    with pytest.raises(ValueError, match="verified absolute path"):
        conversion.convert(source, source.with_suffix(".docx"))
    with pytest.raises(ValueError, match="verified absolute path"):
        conversion.pandoc_binary("pandoc")
    output = source.with_suffix(".docx")
    conversion.convert(source, output, pandoc=pandoc)
    assert "Export contract" in " ".join(p.text for p in Document(output).paragraphs)
