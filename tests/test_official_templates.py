from pathlib import Path
import json
import os

import pytest

from paper_factory import conversion, official_templates
from paper_factory.workspace import digest_file


@pytest.mark.parametrize("venue", ["ieee-software", "ieee-security-privacy"])
def test_publisher_template_hashes_and_actual_template_structure(venue):
    from docx import Document
    path, record = official_templates.official_template(venue)
    assert digest_file(path) == record["sha256"]
    assert record["download_url"].startswith("https://template-selector.ieee.org/api/ieee-template-selector/template/")
    assert record["publication_type_id"] == 2
    assert record["template_id"] in {549, 405}
    document = Document(path)
    assert {"Title_document", "Authors", "Abstract", "Para", "Bib_entry"} <= {style.name for style in document.styles}
    assert len(document.sections) >= 4


def test_publisher_template_rejects_corrupt_or_unknown_assets(tmp_path, monkeypatch):
    with pytest.raises(ValueError, match="only for IEEE"):
        official_templates.official_template("conference")
    manifest = json.loads((official_templates.ROOT / "provenance.json").read_bytes())
    (tmp_path / "provenance.json").write_text(json.dumps(manifest))
    (tmp_path / manifest["templates"]["ieee-software"]["filename"]).write_bytes(b"not the publisher's bytes")
    monkeypatch.setattr(official_templates, "ROOT", tmp_path)
    with pytest.raises(ValueError, match="provenance SHA"):
        official_templates.official_template("ieee-software")


@pytest.mark.parametrize("venue", ["ieee-software", "ieee-security-privacy"])
def test_native_export_keeps_official_styles_sections_and_complete_contents(tmp_path, pandoc, venue):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    source = tmp_path / "manuscript.md"
    source.write_text("# Template content preservation\n\nA Real Author\n\nA Real University\n\n"
                      "## Abstract\n\nA bounded abstract with exact evidence.\n\n"
                      "## Decision evidence\n\nThe complete conclusion is qualified.\n\n"
                      "Table 1. Exact retained decision table.\n\n"
                      "| Decision | Evidence |\n|---|---|\n| Retain | Exact source witness |\n\n"
                      "## References\n\n1. Actual source record.\n", encoding="utf-8")
    before = source.read_bytes()
    output = source.with_suffix(".docx")
    first = conversion.convert(source, output, pandoc=pandoc, venue_id=venue)
    document = Document(output)
    text = " ".join(p.text for p in document.paragraphs)
    assert "Template content preservation" in text and "Abstract—A bounded abstract" in text
    assert "A Real Author" in text and "The complete conclusion is qualified." in text
    assert document.paragraphs[0].style.name == "Title_document"
    assert document.tables[0].cell(1, 1).text == "Exact source witness"
    assert len(document.tables[0].rows) == 2
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [
        ["Decision", "Evidence"], ["Retain", "Exact source witness"]]
    assert all(row._tr.trPr.find(qn("w:cantSplit")) is not None for row in document.tables[0].rows)
    assert all(paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
               for row in document.tables[0].rows for cell in row.cells for paragraph in cell.paragraphs)
    assert document.tables[0].rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert all(paragraph.paragraph_format.keep_with_next is True
               for cell in document.tables[0].rows[0].cells for paragraph in cell.paragraphs)
    assert all(paragraph.paragraph_format.keep_with_next is not True
               for cell in document.tables[0].rows[1].cells for paragraph in cell.paragraphs)
    assert next(paragraph for paragraph in document.paragraphs if paragraph.text == "Table 1. Exact retained decision table.").paragraph_format.keep_with_next is True
    assert next(paragraph for paragraph in document.paragraphs if paragraph.text == "The complete conclusion is qualified.").paragraph_format.keep_with_next is not True
    columns = {s._sectPr.find(qn("w:cols")).get(qn("w:num"), "1") for s in document.sections}
    assert columns == {"1", "2"}
    assert first["official_template"]["venue_id"] == venue
    assert first["official_template"]["sha256"] == digest_file(official_templates.official_template(venue)[0])
    again = conversion.convert(source, output, pandoc=pandoc, venue_id=venue)
    assert again["output_sha256"] == first["output_sha256"]
    assert source.read_bytes() == before
    assert "--reference-doc" in first["command"]


def test_pdf_renders_bound_native_template_bytes_without_typst(tmp_path, pandoc, monkeypatch):
    from pypdf import PdfWriter
    source = tmp_path / "manuscript.md"
    source.write_text("# Native PDF contract\n\n## Abstract\n\nThe exact native source is rendered.\n\n## Results\n\nPreserved findings.\n")
    calls = []
    def render(document, output):
        calls.append((document, digest_file(document)))
        writer = PdfWriter()
        writer.add_blank_page(width=500, height=700)
        writer.write(output)
        return {"source_docx_sha256": digest_file(document), "engine": "test-only native renderer"}
    monkeypatch.setattr(official_templates, "word_pdf", render)
    report = conversion.convert(source, source.with_suffix(".pdf"), pandoc=pandoc, venue_id="ieee-software")
    assert calls == [(source.with_suffix(".docx"), report["source_docx_sha256"])]
    assert not source.with_suffix(".typ").exists()
    assert report["output_sha256"] == digest_file(source.with_suffix(".pdf"))
    assert report["official_template"]["template_id"] == 549


def test_template_pdf_renderer_failure_cannot_leave_stale_success(tmp_path, pandoc, monkeypatch):
    source = tmp_path / "manuscript.md"
    source.write_text("# Failed renderer\n\n## Abstract\n\nVerified text.\n\n## Result\n\nPreserve native Word.\n")
    output = source.with_suffix(".pdf")
    output.write_bytes(b"old PDF")
    def unavailable(*_):
        raise ValueError("Microsoft Word is unavailable")
    monkeypatch.setattr(official_templates, "word_pdf", unavailable)
    with pytest.raises(ValueError, match="Microsoft Word is unavailable"):
        conversion.convert(source, output, pandoc=pandoc, venue_id="ieee-software")
    assert not output.exists()
    assert source.with_suffix(".docx").is_file()


def test_pdf_never_replaces_an_existing_word_companion_on_success_or_failure(tmp_path, pandoc, monkeypatch):
    source = tmp_path / "manuscript.md"
    source.write_text("# Immutable Word companion\n\n## Abstract\n\nThe source is bounded.\n\n## Result\n\nAn immutable export.\n")
    document = source.with_suffix(".docx")
    conversion.convert(source, document, pandoc=pandoc, venue_id="ieee-software")
    original = document.read_bytes()
    timestamp = document.stat().st_mtime_ns
    def fail(*_):
        raise ValueError("Explicit renderer failure")
    monkeypatch.setattr(official_templates, "word_pdf", fail)
    with pytest.raises(ValueError, match="Explicit renderer failure"):
        conversion.convert(source, source.with_suffix(".pdf"), pandoc=pandoc, venue_id="ieee-software")
    assert document.read_bytes() == original and document.stat().st_mtime_ns == timestamp
    document.write_bytes(b"Different immutable existing Word bytes")
    changed = document.read_bytes()
    with pytest.raises(ValueError, match="existing Word companion differs"):
        conversion.convert(source, source.with_suffix(".pdf"), pandoc=pandoc, venue_id="ieee-software")
    assert document.read_bytes() == changed


@pytest.mark.skipif(os.name != "nt", reason="Actual Windows Word COM profile isolation regression")
def test_word_pdf_opens_exact_native_manuscript_under_isolated_engine_profile(tmp_path, pandoc, monkeypatch):
    import winreg
    from pypdf import PdfReader
    try:
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"Word.Application\CLSID"):
            pass
    except FileNotFoundError:
        pytest.skip("Microsoft Word is not installed on this test host")
    source = tmp_path / "source.md"
    source.write_text("# Isolated engine profile\n\n## Abstract\n\nAn exact native document must open.\n\nKeywords: isolation; source; native rendering\n\n## Result\n\nThe renderer preserves the independent source bytes.\n")
    document = source.with_suffix(".docx")
    conversion.convert(source, document, pandoc=pandoc, venue_id="ieee-software")
    original = document.read_bytes()
    private = tmp_path / "private-engine-profile"
    private.mkdir()
    for name in ("HOME", "USERPROFILE", "APPDATA", "LOCALAPPDATA"):
        monkeypatch.setenv(name, str(private))
    monkeypatch.setenv("PATH", "")
    output = source.with_suffix(".pdf")
    report = official_templates.word_pdf(document, output)
    assert report["word"]["read_only"] is True and report["visible"] is False
    assert report["source_docx_sha256"] == digest_file(document)
    assert document.read_bytes() == original
    text = " ".join(page.extract_text() or "" for page in PdfReader(output).pages)
    assert "Isolated engine profile" in text
    measured = official_templates.measure_word(document)
    assert measured["source_docx_sha256"] == digest_file(document) and measured["source_unchanged"]
    assert measured["visible"] is False and measured["metrics"]["read_only"] is True
    assert measured["metrics"]["abstract_words"] == 6


@pytest.mark.parametrize("venue", ["ieee-software", "ieee-security-privacy"])
def test_actual_template_citation_rules_preserve_prose_code_links_and_known_metadata(tmp_path, pandoc, venue):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    source = tmp_path / "citations.md"
    source.write_text('# Publisher citation rules\n\nAuthor [1]\n\n## Abstract\n\nAn exact bounded abstract ends normally.\n\n'
                      'Keywords: architecture; security; retained evidence\n\n'
                      '## Evidence\n\n**A supported finding**. **[1]** [2]\n\nA second finding [1].\n\n'
                      'The code `arr[1]` is unchanged. [2]\n\nA [named link](https://example.org/retained) remains. [1]\n\n'
                      '“Quoted content.” [1]\n\n'
                      'An unknown citation stays [99].\n\n## References\n\n'
                      '[1] A. Author, “Known title,” 2024. [Retained source](https://example.org/one).\n\n'
                      '[2] B. Author, “Other title,” 2023. Accessed 2026-10-18.\n', encoding='utf-8')
    original = source.read_bytes()
    output = source.with_suffix('.docx')
    receipt = conversion.convert(source, output, pandoc=pandoc, venue_id=venue)
    document = Document(output)
    paragraphs = document.paragraphs
    text = '\n'.join(p.text for p in paragraphs)
    assert 'Author [1]' in text and 'arr[1]' in text and 'An unknown citation stays [99].' in text
    assert '1. A. Author, “Known title,” 2024. Retained source.' in text
    assert '2. B. Author, “Other title,” 2023. Accessed 2026-10-18.' in text
    assert all(p.alignment == WD_ALIGN_PARAGRAPH.LEFT for p in paragraphs if p.style.name == 'Bib_entry')
    assert any(run.bold and run.text == 'A supported finding' for p in paragraphs for run in p.runs)
    assert any('https://example.org/retained' == relationship.target_ref for relationship in document.part.rels.values())
    assert next(p for p in paragraphs if p.style.name == 'Abstract').alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert next(p for p in paragraphs if p.text.startswith('Keywords:')).style.name != 'Abstract'
    superscripts = [run.text for p in paragraphs for run in p.runs if run.font.superscript]
    if venue == 'ieee-software':
        assert 'A supported finding.1,2' in text and 'A second finding.1' in text
        assert '“Quoted content.”1' in text
        assert superscripts == ['1,2', '1', '2', '1', '1']
    else:
        assert 'A supported finding [1], [2].' in text and 'A second finding [1].' in text
        assert '“Quoted content” [1].' in text
        assert not superscripts
        reference = next(p for p in paragraphs if p.style.name == 'Bib_entry')
        assert reference.paragraph_format.first_line_indent < 0
    assert source.read_bytes() == original
    assert receipt['official_template']['citation_typography']['callout_groups_formatted'] == 5
    assert receipt['official_template']['citation_typography']['reference_labels_formatted'] == 2
    assert receipt['official_template']['citation_typography']['unknown_reference_numbers_preserved'] == ['99']
