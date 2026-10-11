"""Real native caption conversion; synthetic content establishes no research findings."""
import hashlib
import subprocess
import pytest

from paper_factory import conversion
from paper_factory.workspace import digest_file


def captioned_table(caption, columns=2):
    headers = [f"Dimension {number}" for number in range(1, columns + 1)]
    values = [f"Literal value {number}" for number in range(1, columns + 1)]
    return "| " + " | ".join(headers) + " |\n|" + "---|" * columns + "\n| " + " | ".join(values) + " |\n\nTable: " + caption + "\n\n"


@pytest.mark.parametrize("explicit", ["Table 7:", "TABLE 7.", "Table A1:"])
def test_actual_docx_numbers_narrow_and_wide_captions_without_repeating_explicit_labels(tmp_path, pandoc, explicit):
    from docx import Document
    from docx.oxml.ns import qn

    source = tmp_path / "captioned.md"
    source.write_text("# Caption fidelity\n\nThe unchanged source path is `create-frontron/__tests__/security.test.ts` and citation [1].\n\n" +
        captioned_table("Resource *designation* and `src/preload.ts`.") +
        captioned_table(explicit + " Explicit numbered caption.") +
        captioned_table("Complete witness *coverage*.", columns=6) +
        captioned_table("Final record comparison.") +
        "## References\n\n[1] Canonical retained reference. https://example.test/primary.\n\n" +
        "## Acknowledgments\n\nThis synthetic converter fixture has no model call or human study.\n", encoding="utf-8")
    original = source.read_bytes()
    output = source.with_suffix(".docx")
    receipt = conversion.convert(source, output, pandoc=pandoc)
    document = Document(output)
    labels = [paragraph.text for paragraph in document.paragraphs if paragraph.text.startswith(("Table ", "TABLE "))]
    next_number = 8 if explicit != "Table A1:" else 3
    assert labels == ["Table 1: Resource designation and src/preload.ts.", explicit + " Explicit numbered caption.",
                      f"Table {next_number}: Complete witness coverage.", f"Table {next_number + 1}: Final record comparison."]
    assert len(document.tables) == 3
    caption_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.style.style_id == "TableCaption"]
    assert [paragraph.text for paragraph in caption_paragraphs] == [labels[0], labels[1], labels[3]]
    assert any(run.italic and run.text == "designation" for run in caption_paragraphs[0].runs)
    assert all(not paragraph._p.findall(qn("w:numPr")) for paragraph in caption_paragraphs)
    visible = " ".join(paragraph.text for paragraph in document.paragraphs)
    assert "create-frontron/__tests__/security.test.ts" in visible
    assert "[1] Canonical retained reference. https://example.test/primary." in visible
    assert "This synthetic converter fixture has no model call or human study." in visible
    assert "Record 1" in visible
    for number in range(1, 7):
        assert f"Dimension {number}: Literal value {number}" in visible
    assert source.read_bytes() == original
    assert receipt["input_sha256"] == hashlib.sha256(original).hexdigest()
    assert receipt["output_sha256"] == digest_file(output)


def test_pdf_caption_numbering_and_typst_ast_remain_native_and_unchanged(tmp_path, pandoc):
    pytest.importorskip("typst")
    from pypdf import PdfReader

    source = tmp_path / "native-pdf-captions.md"
    source.write_text("# Native PDF captions\n\n" + captioned_table("First retained caption.") +
                      captioned_table("Second retained caption."), encoding="utf-8")
    original = source.read_bytes()
    lua = tmp_path / "filter.lua"
    lua.write_text(conversion.WIDE_TABLE_FILTER, encoding="utf-8")
    base = [pandoc, str(source), "--from=markdown-smart-raw_tex-raw_html", "--to=typst"]
    options = {"capture_output": True, "text": True, "encoding": "utf-8", "check": True}
    native = subprocess.run(base, **options).stdout
    filtered = subprocess.run([*base, "--lua-filter", str(lua)], **options).stdout
    assert filtered == native
    output = source.with_suffix(".pdf")
    receipt = conversion.convert(source, output, pandoc=pandoc)
    text = " ".join(" ".join(page.extract_text().split()) for page in PdfReader(output).pages)
    assert "Table 1: First retained caption." in text and "Table 2: Second retained caption." in text
    assert "Table 1: Table 1:" not in text and "Table 2: Table 2:" not in text
    assert source.read_bytes() == original
    assert receipt["input_sha256"] == hashlib.sha256(original).hexdigest()
    assert receipt["output_sha256"] == digest_file(output)
