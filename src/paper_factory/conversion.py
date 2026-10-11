"""Convert using the explicit bundled Pandoc executable and trusted Typst PDF compiler."""

import json
import os
import subprocess
import hashlib
import struct
from datetime import datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory

from .workspace import digest_file, write_json, ensure_unlinked


# Keep the source document unchanged. Pandoc's parsed cells retain their inline
# formatting and complete contents while wide rows become readable records in
# every native output format, rather than seven competing narrow columns.
WIDE_TABLE_FILTER = r'''
local table_number = 0
function Table(tbl)
  local numbered_caption = false
  if FORMAT == 'docx' and tbl.caption and tbl.caption.long and #tbl.caption.long > 0 then
    local first = tbl.caption.long[1]
    if first.t == 'Plain' or first.t == 'Para' then
      local explicit, tail = pandoc.utils.stringify(tbl.caption.long):lower():match('^%s*table%s+([a-z]?%d+)(.*)$')
      if explicit and not tail:match('^%w') then
        table_number = math.max(table_number, tonumber(explicit) or table_number + 1)
      else
        table_number = table_number + 1
        first.content:insert(1, pandoc.Space())
        first.content:insert(1, pandoc.Str('Table ' .. table_number .. ':'))
        numbered_caption = true
      end
    end
  end
  if #tbl.colspecs < 6 then return numbered_caption and tbl or nil end
  local result = pandoc.List()
  if tbl.caption and tbl.caption.long then
    result:extend(tbl.caption.long)
  end
  local headers = {}
  local head_spans = {}
  for _, row in ipairs(tbl.head.rows) do
    local column = 1
    for _, cell in ipairs(row.cells) do
      while (head_spans[column] or 0) > 0 do column = column + 1 end
      local label = pandoc.utils.blocks_to_inlines(cell.contents)
      for offset = 0, cell.col_span - 1 do
        local index = column + offset
        if headers[index] then
          headers[index]:extend({pandoc.Space(), pandoc.Str('/'), pandoc.Space()})
          headers[index]:extend(label)
        else
          headers[index] = label:clone()
        end
        head_spans[index] = cell.row_span
      end
      column = column + cell.col_span
    end
    for index, span in pairs(head_spans) do head_spans[index] = span - 1 end
  end
  local record = 0
  local function rows(items)
    local spans = {}
    for _, row in ipairs(items) do
      record = record + 1
      -- Native headings keep each record label with its first field across pages.
      result:insert(pandoc.Header(5, {pandoc.Str('Record ' .. record)}))
      local fields = pandoc.List()
      local column = 1
      for _, cell in ipairs(row.cells) do
        while (spans[column] or 0) > 0 do column = column + 1 end
        local names = pandoc.List()
        for offset = 0, cell.col_span - 1 do
          if offset > 0 then names:extend({pandoc.Space(), pandoc.Str('/'), pandoc.Space()}) end
          names:extend(headers[column + offset] or {pandoc.Str('Column ' .. (column + offset))})
          spans[column + offset] = cell.row_span
        end
        names:insert(pandoc.Str(':'))
        local label = pandoc.Strong(names)
        local blocks = pandoc.List()
        if #cell.contents > 0 and (cell.contents[1].t == 'Plain' or cell.contents[1].t == 'Para') then
          local inlines = pandoc.List({label, pandoc.Space()})
          inlines:extend(cell.contents[1].content)
          blocks:insert(pandoc.Para(inlines))
          for index = 2, #cell.contents do blocks:insert(cell.contents[index]) end
        else
          blocks:insert(pandoc.Para({label}))
          blocks:extend(cell.contents)
        end
        fields:insert(blocks)
        column = column + cell.col_span
      end
      result:insert(pandoc.BulletList(fields))
      for index, span in pairs(spans) do spans[index] = span - 1 end
    end
  end
  for _, body in ipairs(tbl.bodies) do rows(body.head); rows(body.body) end
  rows(tbl.foot.rows)
  if record == 0 then rows(tbl.head.rows) end
  return result
end
'''


def verify_receipts(markdown: Path, outputs: dict[str, Path], receipt_path: Path) -> list[str]:
    """Bind native exports to their recorded source and conversion output."""
    reports = json.loads(receipt_path.read_bytes())
    if not isinstance(reports, dict) or set(reports) != set(outputs):
        return ["Conversion receipt format set differs from the required exports"]
    source_digest = digest_file(markdown)
    errors = []
    for format, output in outputs.items():
        report = reports[format]
        if (not isinstance(report, dict) or report.get("input_sha256") != source_digest
                or report.get("output_sha256") != digest_file(output)):
            errors.append(f"Conversion receipt differs from its source or native output: {format}")
    return errors


def _format_docx(output: Path, *, line_numbers: bool, page_numbers: bool) -> None:
    """Use python-docx for native Word settings, keeping Pandoc's content."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise ValueError("Install the PDF/package extra for Word formatting: pip install -e '.[pdf]'") from exc
    document = Document(output)
    for section in document.sections:
        for old in section._sectPr.findall(qn("w:lnNumType")):
            section._sectPr.remove(old)
        if line_numbers:
            numbers = OxmlElement("w:lnNumType")
            numbers.set(qn("w:countBy"), "1")
            numbers.set(qn("w:restart"), "continuous")
            section._sectPr.insert_element_before(numbers, "w:pgNumType", "w:cols", "w:docGrid")
        section.different_first_page_header_footer = False
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            footer.is_linked_to_previous = False
            for child in list(footer._element):
                footer._element.remove(child)
            paragraph = footer.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if page_numbers:
                field = OxmlElement("w:fldSimple")
                field.set(qn("w:instr"), "PAGE")
                run = OxmlElement("w:r")
                text = OxmlElement("w:t")
                text.text = "1"
                run.append(text)
                field.append(run)
                paragraph._p.append(field)
    document.save(output)
    # Independently reopen the output rather than accepting a successful save.
    checked = Document(output)
    if not checked.paragraphs:
        raise ValueError("Compiled Word document has no paragraphs")
    for section in checked.sections:
        if bool(section._sectPr.findall(qn("w:lnNumType"))) != line_numbers:
            raise ValueError("Word line-number setting differs from policy")
        actual = any(field.get(qn("w:instr")) == "PAGE" for field in section.footer._element.iter(qn("w:fldSimple")))
        if actual != page_numbers:
            raise ValueError("Word page-number field differs from policy")


def pandoc_binary(explicit: str | None = None) -> str:
    selected = explicit or os.environ.get("PYPANDOC_PANDOC")
    binary = Path(selected) if selected else None
    if binary is None or not binary.is_absolute() or not binary.is_file():
        raise ValueError("The bundled Pandoc executable requires its verified absolute path")
    return str(binary.resolve())


def convert(markdown: Path, output: Path, *, pandoc: str | None = None, bibliography: Path | None = None, csl: Path | None = None, metadata: dict | None = None, line_numbers: bool = False, page_numbers: bool = True) -> dict:
    suffix = output.suffix.lower()
    if suffix not in {".pdf", ".tex", ".docx"}:
        raise ValueError("Supported conversion outputs: PDF, LaTeX and DOCX")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.unlink(missing_ok=True)
    target = output.with_suffix(".typ") if suffix == ".pdf" else output
    target.unlink(missing_ok=True)
    binary = pandoc_binary(pandoc)
    working_directory = markdown.parent.resolve()
    source_digest = digest_file(markdown)
    format = "typst" if suffix == ".pdf" else {".tex": "latex", ".docx": "docx"}[suffix]
    command = [binary, str(markdown.resolve()), "--from=markdown-smart-raw_tex-raw_html", "--standalone", "--to=" + format, "-o", str(target.resolve())]
    table_filter = output.with_name(output.stem + "-table-layout.lua")
    table_filter.write_text(WIDE_TABLE_FILTER, encoding="utf-8")
    generated_inputs = {"--lua-filter": table_filter}
    if bibliography:
        command += ["--citeproc", "--bibliography", str(bibliography.resolve())]
    if csl:
        command += ["--csl", str(csl.resolve())]
    if suffix == ".pdf":
        metadata = {**(metadata or {}), "mainfont": "Pretendard", "margin": {"x": "1in", "y": "1in"}, "page-numbering": "1" if page_numbers else None}
    elif format == "latex":
        headers = []
        if line_numbers:
            headers += [r"\usepackage{lineno}", r"\linenumbers"]
        if not page_numbers:
            headers.append(r"\pagestyle{empty}")
        if headers:
            # The source reader disables raw TeX; Pandoc's explicit header input
            # inserts these trusted native settings without treating them as prose.
            header_path = output.with_name(f"{output.stem}-formatting.tex")
            header_path.write_text("\n".join(headers) + "\n", encoding="utf-8")
            generated_inputs["--include-in-header"] = header_path
    if metadata:
        metadata_path = output.parent / f"{output.stem}-pandoc-metadata.json"
        write_json(metadata_path, metadata)
        generated_inputs["--metadata-file"] = metadata_path
    if suffix == ".pdf":
        # Typst headings end at a source newline. Let Typst perform visual
        # wrapping without turning Pandoc's continued title into body text.
        command += ["--wrap=none", "-V", "papersize:a4", "-V", "fontsize:11pt"]
    else:
        command += ["-V", "geometry:margin=1in"]
    try:
        # Lua's Windows file opening rejects long paths even when relative to
        # cwd. Preserve generated inputs beside the output, but run Pandoc with
        # their exact bytes at short temporary paths; keep original I/O and cwd.
        with TemporaryDirectory(prefix="paperfactory-convert-") as temporary:
            for option, retained in generated_inputs.items():
                staged = Path(temporary) / retained.name
                staged.write_bytes(retained.read_bytes())
                command += [option, str(staged.resolve())]
            result = subprocess.run(command, cwd=working_directory, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120, check=False,
                                    creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
    except subprocess.TimeoutExpired as exc:
        target.unlink(missing_ok=True)
        raise ValueError("Pandoc conversion timed out after 120 seconds") from exc
    if result.returncode != 0 or not target.is_file():
        target.unlink(missing_ok=True)
        raise ValueError(f"Pandoc failed ({result.returncode}): {result.stderr[:2000]}")
    engine = "Pandoc"
    if suffix == ".docx":
        try:
            _format_docx(output, line_numbers=line_numbers, page_numbers=page_numbers)
        except Exception as exc:
            output.unlink(missing_ok=True)
            raise ValueError(f"Word document formatting failed: {exc}") from exc
        engine = "Pandoc → python-docx"
    if suffix == ".pdf":
        try:
            import typst
        except ImportError as exc:
            raise ValueError("The bundled Typst PDF compiler is unavailable") from exc
        try:
            typst_source = target.read_text(encoding="utf-8")
            settings = '\nset page(numbering: "1")\n' if page_numbers else '\nset page(numbering: none)\n'
            if line_numbers:
                settings += 'set par.line(numbering: "1", numbering-scope: "document")\n'
            # Scope wrapping to table cells; retain every printed identifier character.
            # Typst table.cell show rules and grapheme clusters keep content out of
            # neighboring columns without rewriting the retained Markdown or data.
            settings += r'''
show table.cell: cell => {
  set text(size: 9pt, hyphenate: false)
  show regex("[A-Za-z0-9_]{20,}"): word => text(word.text.clusters().join("\u{200b}"))
  cell
}
'''
            # Native rules apply to the complete generated document.
            target.write_text('#show: body => {\n' + settings + 'body\n}\n' + typst_source, encoding="utf-8")
            # Pandoc generates the source; Typst's sandbox root is this directory.
            fonts = os.environ.get("TYPST_FONT_PATHS")
            font_paths = [str(Path(__file__).with_name("fonts").resolve())]
            if fonts:
                font_paths.append(fonts)
            typst.compile(str(target), output=str(output), root=str(output.parent),
                          timestamp=datetime(2000, 1, 1, tzinfo=timezone.utc),
                          font_paths=font_paths, ignore_system_fonts=True)
        except Exception as exc:
            output.unlink(missing_ok=True)
            raise ValueError(f"Typst PDF compilation failed: {exc}") from exc
        engine = "Pandoc → Typst"
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(output)
            if not output.read_bytes().startswith(b"%PDF-") or not reader.pages or reader.is_encrypted:
                raise ValueError("Compiled PDF is empty, encrypted or invalid")
        except Exception as exc:
            output.unlink(missing_ok=True)
            raise ValueError(f"Compiled PDF validation failed: {exc}") from exc
    if digest_file(markdown) != source_digest:
        output.unlink(missing_ok=True)
        raise ValueError("Manuscript changed during conversion")
    receipt = {"command": command, "engine": engine, "input_sha256": source_digest,
               "output_sha256": digest_file(output), "diagnostics": result.stderr,
               "working_directory": str(working_directory),
               "generated_inputs": {option: {"path": str(path.resolve()), "sha256": digest_file(path)}
                                    for option, path in generated_inputs.items()},
               "table_layout_sha256": digest_file(table_filter),
               "table_layout": "Tables with six or more columns render as complete labeled row records."}
    if bibliography is not None:
        receipt["bibliography_sha256"] = digest_file(bibliography)
    if csl is not None:
        receipt["csl_sha256"] = digest_file(csl)
    if suffix == ".pdf":
        receipt["typst_sha256"] = digest_file(target)
    return receipt


def inspect_pdf_quality(pdf: Path, output_root: Path, *, typst_source: Path | None = None) -> dict:
    """Render every bound page; structural checks never attest visual approval.

    Bundled Typst provides rasterization of its retained source. Recompiling the
    PDF must produce the exact original hash before those images can attest its
    pages. This avoids requiring a separately installed PDF renderer.
    """
    from pypdf import PdfReader
    ensure_unlinked(pdf)
    ensure_unlinked(output_root)
    if pdf.stat().st_size > 32 * 1024 * 1024:
        raise ValueError("PDF quality review exceeds its bounded file size")
    reader = PdfReader(pdf)
    if reader.is_encrypted or not 1 <= len(reader.pages) <= 100:
        raise ValueError("PDF quality review requires one to one hundred readable pages")
    receipt = {"pdf_sha256": digest_file(pdf), "page_count": len(reader.pages), "pages": [],
               "automatic_checks": False, "visual_review_status": "pending", "status": "render_unavailable",
               "issues": [], "scope": "Every-page structural/content and bound raster checks; clipping, overlap, formula correctness, missing glyphs and journal compliance still require actual visual review."}
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        receipt["pages"].append({"number": number, "text_sha256": hashlib.sha256(text.encode()).hexdigest(),
                                 "text_chars": len(text), "width_points": width, "height_points": height})
        if width <= 0 or height <= 0 or width > 2000 or height > 2000:
            receipt["issues"].append(f"Page {number} has unsupported dimensions.")
        if not text.strip():
            receipt["issues"].append(f"Page {number} has no extractable text; image-only content requires explicit visual verification.")
        if "\ufffd" in text or "\u0000" in text:
            receipt["issues"].append(f"Page {number} contains replacement or null text characters.")
    if typst_source is None or not typst_source.is_file():
        receipt["issues"].append("The exact retained Typst rendering source is unavailable; every-page rasterization is unverified.")
        return receipt
    ensure_unlinked(typst_source)
    if typst_source.parent != pdf.parent:
        raise ValueError("Quality rendering source must belong to the PDF export directory")
    receipt["typst_sha256"] = digest_file(typst_source)
    import typst
    font_paths = [str(Path(__file__).with_name("fonts").resolve())]
    fonts = os.environ.get("TYPST_FONT_PATHS")
    if fonts:
        font_paths.append(fonts)
    options = {"root": str(pdf.parent), "timestamp": datetime(2000, 1, 1, tzinfo=timezone.utc),
               "font_paths": font_paths, "ignore_system_fonts": True}
    recomputed = typst.compile(str(typst_source), format="pdf", **options)
    if hashlib.sha256(recomputed).hexdigest() != receipt["pdf_sha256"]:
        receipt.update(status="failed")
        receipt["issues"].append("Raster source recompilation differs from the actual exported PDF bytes.")
        return receipt
    rendered = typst.compile(str(typst_source), format="png", ppi=96, **options)
    images = [rendered] if isinstance(rendered, bytes) else list(rendered)
    if len(images) != len(reader.pages):
        raise ValueError("Rendered images do not cover every actual PDF page")
    output_root.mkdir(parents=True, exist_ok=True)
    for page, content in zip(receipt["pages"], images):
        if not isinstance(content, bytes) or content[:8] != b"\x89PNG\r\n\x1a\n" or content[12:16] != b"IHDR":
            raise ValueError("Page renderer did not return a PNG image")
        width, height = struct.unpack(">II", content[16:24])
        if not 1 <= width <= 3000 or not 1 <= height <= 3000 or len(content) > 16 * 1024 * 1024:
            raise ValueError("Page image exceeds bounded review dimensions")
        path = output_root / f"page-{page['number']:03}.png"
        ensure_unlinked(path)
        if path.exists():
            if path.read_bytes() != content:
                raise ValueError("Page review cannot replace different existing rendering bytes")
        else:
            with path.open("xb") as stream:
                stream.write(content)
        page.update(image_path=str(path.absolute()), image_sha256=hashlib.sha256(content).hexdigest(),
                    width_pixels=width, height_pixels=height)
    if digest_file(pdf) != receipt["pdf_sha256"] or digest_file(typst_source) != receipt["typst_sha256"]:
        raise ValueError("Export changed during its every-page quality inspection")
    receipt.update(automatic_checks=not receipt["issues"], status="pending_visual_review" if not receipt["issues"] else "failed")
    return receipt
