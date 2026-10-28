"""Convert authored Markdown to the retained publisher Word template only."""

from datetime import datetime, timezone
from io import BytesIO
import hashlib
import os
from pathlib import Path
import shutil
import subprocess
from tempfile import TemporaryDirectory
import zipfile

from _publication_files import digest_file, ensure_unlinked
from _publication_templates import official_template, apply_official_styles


AUTHOR_MARKER_FILTER = """-- Remove only a dedicated pair of author-control paragraphs before Abstract.
-- Code nodes and prose containing markers are never control paragraphs.
local function marker(block)
  if block.t ~= 'Para' then return nil end
  for _, inline in ipairs(block.content) do
    if inline.t ~= 'Str' and inline.t ~= 'Space' then return nil end
  end
  local value = pandoc.utils.stringify(block.content)
  if value == '<!-- paper-factory:authors:start -->' then return 'start' end
  if value == '<!-- paper-factory:authors:end -->' then return 'end' end
  return nil
end

function Pandoc(document)
  local starts, ends, abstract = {}, {}, nil
  for index, block in ipairs(document.blocks) do
    if block.t == 'Header' and block.level == 2
       and pandoc.utils.stringify(block.content):lower() == 'abstract' then
      abstract = index
      break
    end
    local kind = marker(block)
    if kind == 'start' then table.insert(starts, index) end
    if kind == 'end' then table.insert(ends, index) end
  end
  if not abstract or #starts ~= 1 or #ends ~= 1 or starts[1] >= ends[1] then
    return document
  end
  local blocks = pandoc.List()
  for index, block in ipairs(document.blocks) do
    if index ~= starts[1] and index ~= ends[1] then blocks:insert(block) end
  end
  return pandoc.Pandoc(blocks, document.meta)
end
"""


def _format_docx(output: Path) -> None:
    """Use python-docx for native Word settings, keeping Pandoc's content."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise ValueError("Install the skill requirements for native Word formatting") from exc
    document = Document(output)
    for section in document.sections:
        for old in section._sectPr.findall(qn("w:lnNumType")):
            section._sectPr.remove(old)
        section.different_first_page_header_footer = False
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            footer.is_linked_to_previous = False
            for child in list(footer._element):
                footer._element.remove(child)
            paragraph = footer.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "1"
            run.append(text)
            field.append(run)
            paragraph._p.append(field)
    document.save(output)
    # Independently reopen the output rather than accepting a successful save.
    checked = Document(output)
    if not checked.paragraphs:
        raise ValueError("Compiled Word document has no paragraphs")
    for section in checked.sections:
        if section._sectPr.findall(qn("w:lnNumType")):
            raise ValueError("Word line-number setting differs from policy")
        actual = any(field.get(qn("w:instr")) == "PAGE" for field in section.footer._element.iter(qn("w:fldSimple")))
        if not actual:
            raise ValueError("Word page-number field differs from policy")


def pandoc_binary(explicit: str | None = None) -> Path:
    """Use an explicit executable or existing installation; never download."""
    selected = explicit or os.environ.get("PYPANDOC_PANDOC")
    if selected:
        candidate = Path(selected).expanduser()
        if not candidate.is_absolute() or not candidate.is_file():
            raise ValueError("Pandoc must be an existing absolute executable path")
        return candidate.resolve()
    candidate = shutil.which("pandoc")
    if candidate:
        return Path(candidate).resolve()
    try:
        import pypandoc
        candidate = Path(pypandoc.get_pandoc_path())
    except (ImportError, OSError):
        candidate = None
    # Some Windows pypandoc builds return the executable stem without .exe.
    if candidate is not None and os.name == "nt" and not candidate.suffix and not candidate.is_file():
        executable = candidate.with_suffix(".exe")
        if executable.is_file():
            candidate = executable
    if candidate is None or not candidate.is_file():
        raise ValueError("Pandoc is unavailable; supply --pandoc or install the skill requirements")
    return candidate.resolve()


def convert_docx(markdown: Path, output: Path, *, venue_id: str,
                 pandoc: str | None = None) -> dict:
    """Use the actual publisher reference, citation rules and section masters."""
    ensure_unlinked(markdown)
    ensure_unlinked(output)
    if output.suffix.lower() != ".docx" or output.exists():
        raise ValueError("Native Word export requires a new .docx output path")
    binary = pandoc_binary(pandoc)
    template_path, template = official_template(venue_id)
    original_sha = digest_file(markdown)
    output.parent.mkdir(parents=True, exist_ok=True)
    command = [str(binary), str(markdown.resolve()), "--from=markdown-smart-raw_tex-raw_html",
               "--standalone", "--to=docx", "-o", str(output.resolve())]
    # A short reference path avoids native Windows path-length restrictions.
    with TemporaryDirectory(prefix="paper-factory-convert-") as temporary:
        staged_reference = Path(temporary) / template_path.name
        staged_reference.write_bytes(template_path.read_bytes())
        staged_filter = Path(temporary) / "author-control-markers.lua"
        filter_bytes = AUTHOR_MARKER_FILTER.encode("utf-8")
        staged_filter.write_bytes(filter_bytes)
        command += ["--reference-doc", str(staged_reference), "--lua-filter", str(staged_filter)]
        try:
            result = subprocess.run(command, cwd=markdown.parent.resolve(), capture_output=True,
                                    text=True, encoding="utf-8", errors="replace", timeout=120,
                                    creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        except subprocess.TimeoutExpired as exc:
            raise ValueError("Pandoc conversion timed out after 120 seconds") from exc
    if result.returncode or not output.is_file():
        raise ValueError(f"Pandoc failed ({result.returncode}): {result.stderr[:2000]}")
    proof = {**template, **apply_official_styles(output, template_path)}
    _format_docx(output)
    # Retain the existing converter's deterministic native ZIP bytes.
    from docx import Document
    document = Document(output)
    document.core_properties.created = datetime(2000, 1, 1, tzinfo=timezone.utc)
    document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    document.save(output)
    with zipfile.ZipFile(BytesIO(output.read_bytes())) as source_zip:
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as final_zip:
            for name in source_zip.namelist():
                info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                final_zip.writestr(info, source_zip.read(name))
    if digest_file(markdown) != original_sha:
        raise ValueError("The manuscript changed during native conversion")
    if digest_file(template_path) != template["sha256"]:
        raise ValueError("The publisher reference changed during conversion")
    return {"command": command, "engine": "Pandoc → official publisher Word styles",
            "input_sha256": original_sha, "output_sha256": digest_file(output),
            "diagnostics": result.stderr, "official_template": proof,
            "author_control_markers": {"scope": "Only one ordered pair of exact top-level author marker paragraphs before Abstract is omitted from native output; canonical Markdown and code examples are retained.",
                                       "filter_sha256": hashlib.sha256(filter_bytes).hexdigest(),
                                       "filter_source": AUTHOR_MARKER_FILTER},
            "working_directory": str(markdown.parent.resolve())}
