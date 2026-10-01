"""Export measured case-study manuscripts and an auditable reproducibility bundle.

The study owner supplies manuscript.md and study-info.json. This exporter does
not invent results, author approval, publication status, or bibliography entries.
Run from the installed Paper Factory development environment.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import zipfile

from paper_factory.conversion import convert
from paper_factory.author import load_author
from paper_factory.config import load_env_file


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def safe_file(root: Path, relative: str) -> Path:
    path = root / relative
    if path.is_symlink() or not path.resolve().is_relative_to(root.resolve()) or not path.is_file():
        raise ValueError(f"Missing or unsafe study file: {relative}")
    return path


def word_tokens(value: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", value.casefold())


def validate_native(root: Path, title: str) -> dict:
    from docx import Document
    from pypdf import PdfReader

    pdf = PdfReader(root / "paper.pdf")
    text = " ".join(page.extract_text() or "" for page in pdf.pages)
    docx = Document(root / "paper.docx")
    docx_text = " ".join(p.text for p in docx.paragraphs)
    required = word_tokens(title)[:6]
    for name, content in (("PDF", text), ("DOCX", docx_text)):
        tokens = set(word_tokens(content))
        if not set(required).issubset(tokens):
            raise ValueError(f"{name} failed independent title/content validation")
        if "abstract" not in tokens or "conclusion" not in tokens:
            raise ValueError(f"{name} is missing scholarly article sections")
    if len(pdf.pages) < 2:
        raise ValueError("A substantive case-study article must have multiple pages")
    if not (root / "paper.tex").read_text(encoding="utf-8").strip().endswith(r"\end{document}"):
        raise ValueError("LaTeX export is not standalone")
    return {"pdf_pages": len(pdf.pages), "pdf_text_characters": len(text),
            "docx_paragraphs": len(docx.paragraphs), "docx_tables": len(docx.tables),
            "independent_reopen_passed": True}


def reproduce_bundle(root: Path, info: dict, *, private_files: tuple[Path, ...] = ()) -> dict:
    """Keep scripts/data; exclude environments and regenerable large case trees."""
    excluded_dirs = {".venv", "node_modules", ".git", "__pycache__", ".pytest_cache", ".mpl", ".mpl-cache", ".tex-build"}
    excluded_prefixes = ("pilot/", "run/fixtures/", "compiled/", "build/", "diagnostics/")
    allowed_suffixes = {".py", ".mjs", ".cjs", ".js", ".ts", ".json", ".jsonl", ".csv", ".md", ".txt", ".bib", ".png", ".svg", ".pdf", ".docx", ".tex", ".typ", ".log"}
    private_paths = {path.resolve() for path in private_files}
    inventory = []
    for path in sorted(root.rglob("*")):
        relative = path.relative_to(root)
        if path.resolve() in private_paths:
            continue
        if any(part in excluded_dirs or part.startswith(".") for part in relative.parts) or relative.as_posix().startswith(excluded_prefixes):
            continue
        # Keep seeded definitions and actual termination/journal records while
        # omitting regenerable binary target trees from the recovery study.
        if relative.as_posix().startswith("experiment/cases/") and path.name not in {"fixture.json", "kill.json", "journal-before.jsonl"}:
            continue
        if path.is_symlink() or not path.is_file() or path.suffix.lower() not in allowed_suffixes:
            continue
        if relative.as_posix() in {"manifest.json", "bundle-inventory.json"}:
            continue
        if path.stat().st_size > 32 * 1024 * 1024:
            raise ValueError(f"Artifact too large for the review bundle: {relative}")
        inventory.append({"path": relative.as_posix(), "bytes": path.stat().st_size, "sha256": digest(path)})
    write_json(root / "bundle-inventory.json", {
        "repository": info["repository"], "files": inventory,
        "excluded_regenerable_material": ["virtual environments", "node_modules", "large seeded per-case file/MIDI trees", "compiled dependency outputs"],
        "note": "All measured CSV/JSON summaries and experiment generators are included. Full generated case trees remain in the cloud study directory and can be regenerated using the study protocol.",
    })
    target = root / "reproducibility.zip"
    temporary = root / ".reproducibility.zip.tmp"
    with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for item in inventory:
            archive.write(safe_file(root, item["path"]), item["path"])
        archive.write(root / "bundle-inventory.json", "bundle-inventory.json")
    if temporary.stat().st_size > 64 * 1024 * 1024:
        temporary.unlink()
        raise ValueError("Reproducibility bundle exceeds the web artifact size limit")
    temporary.replace(target)
    return {"path": target.name, "bytes": target.stat().st_size, "sha256": digest(target), "members": len(inventory) + 1}


def build(root: Path, author: dict, pandoc: str, bibliography: Path | None, compile_tex: bool,
          *, private_files: tuple[Path, ...] = ()) -> dict:
    info = json.loads(safe_file(root, "study-info.json").read_text(encoding="utf-8"))
    original = safe_file(root, "manuscript.md").read_text(encoding="utf-8")
    match = re.match(r"^#\s+(.+)\n", original)
    if not match:
        raise ValueError("manuscript.md must begin with its article title")
    title = info.get("title") or match.group(1).strip()
    body = original[match.end():].lstrip()
    cited = set(re.findall(r"@([A-Za-z][A-Za-z0-9_:-]*)", body))
    local_bibliography = None
    if bibliography is not None:
        records = json.loads(bibliography.read_text(encoding="utf-8"))
        ids = {record["id"] for record in records}
        if missing := cited - ids:
            raise ValueError("Unresolved bibliography entries: " + ", ".join(sorted(missing)))
        selected_ids = cited | set(info.get("bibliography_ids", []))
        if missing := selected_ids - ids:
            raise ValueError("Unresolved study bibliography entries: " + ", ".join(sorted(missing)))
        chosen = [record for record in records if record["id"] in selected_ids]
        if chosen:
            local_bibliography = root / "bibliography.json"
            write_json(local_bibliography, chosen)
    elif cited:
        raise ValueError("Citation keys require a verified bibliography")
    write_json(root / "author.json", author)
    # JSON quoted strings are valid YAML scalar values and preserve literal text.
    metadata = "---\n" + "\n".join(f"{key}: {json.dumps(value, ensure_ascii=False)}" for key, value in {
        "title": title, "author": [author["display_name"]], "lang": "en-US",
        "date": "1 October 2026", "link-citations": True,
        "reference-section-title": "References",
    }.items()) + "\n---\n\n"
    affiliation = f"{author['affiliation']}, {author['city']}, {author['country']}"
    header = f"{affiliation}\n\nORCID: [{author['orcid']}](https://orcid.org/{author['orcid']}) · Correspondence: {author['email']}\n\n"
    article = root / "paper.md"
    article.write_text(metadata + header + body, encoding="utf-8")
    reports = {}
    for suffix in (".pdf", ".docx", ".tex"):
        reports[suffix] = convert(article, root / ("paper" + suffix), pandoc=pandoc,
                                  bibliography=local_bibliography, line_numbers=True, page_numbers=True)
    native = validate_native(root, title)
    if compile_tex:
        engine = shutil.which("xelatex")
        if not engine:
            raise ValueError("XeLaTeX is required for the requested standalone TeX validation")
        target = root / ".tex-build"
        target.mkdir(exist_ok=True)
        command = [engine, "-interaction=nonstopmode", "-halt-on-error", "-output-directory=" + str(target), "paper.tex"]
        result = subprocess.run(command, cwd=root, capture_output=True, text=True, timeout=120)
        (root / "latex-validation.log").write_text(result.stdout + result.stderr, encoding="utf-8")
        if result.returncode:
            raise ValueError("Standalone XeLaTeX validation failed; inspect latex-validation.log")
        native["standalone_xelatex_passed"] = True
    write_json(root / "export-validation.json", {"native": native, "conversions": reports,
               "author_approval_asserted": False, "external_submission_performed": False})
    bundle = reproduce_bundle(root, info, private_files=private_files)
    files = [
        {"path": "paper.pdf", "label": "논문 PDF", "role": "paper"},
        {"path": "paper.docx", "label": "Word 원고", "role": "paper"},
        {"path": "paper.tex", "label": "LaTeX 원고", "role": "source"},
        {"path": "paper.md", "label": "Markdown 원고", "role": "source"},
        {"path": "export-validation.json", "label": "산출물 검증", "role": "validation"},
        {"path": "bundle-inventory.json", "label": "재현 자료 목록", "role": "provenance"},
        {"path": "reproducibility.zip", "label": "실험·원고 재현 자료", "role": "bundle"},
    ]
    for path in info.get("result_files", ["results.csv", "summary.json"]):
        safe_file(root, path)
        files.append({"path": path, "label": Path(path).name, "role": "results"})
    if local_bibliography:
        files.append({"path": "bibliography.json", "label": "검증된 참고문헌", "role": "bibliography"})
    figures = info.get("figures", [])
    for figure in figures:
        safe_file(root, figure["path"])
    manifest = {**info, "id": root.name, "title": title, "author": author, "status": info.get("status", "complete"),
                "files": files, "figures": figures, "reproducibility_bundle": "reproducibility.zip",
                "validation": native, "bundle": bundle,
                "completed_at": datetime.now(timezone.utc).isoformat(),
                "publication_status": "manuscript_artifacts_complete_not_submitted", "author_approval_asserted": False}
    write_json(root / "manifest.json", manifest)
    return {"study": root.name, "title": title, "native": native, "bundle": bundle}


def load_export_author(author_path: Path | None = None, env_file: Path | None = None) -> dict:
    """Read an explicit local configuration and retain the author API precedence."""
    if env_file is not None:
        load_env_file(env_file)
    explicit = json.loads(author_path.read_text(encoding="utf-8")) if author_path is not None else None
    profile = load_author(explicit=explicit, require=True)
    author = profile.model_dump(exclude_defaults=True)
    required = {"display_name", "affiliation", "city", "country", "orcid", "email"}
    if any(not author.get(field) for field in required):
        raise ValueError("Complete user-confirmed author metadata is required for final artifacts")
    return author


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("root", type=Path, help="One study directory, or a directory of studies")
    parser.add_argument("--author", type=Path, help="Explicit author JSON; overrides environment fields")
    parser.add_argument("--env-file", type=Path, help="Explicit local .env configuration; never included in bundles")
    parser.add_argument("--pandoc", default="pandoc")
    parser.add_argument("--bibliography", type=Path)
    parser.add_argument("--compile-tex", action="store_true")
    args = parser.parse_args()
    author = load_export_author(args.author, args.env_file)
    roots = [args.root] if (args.root / "study-info.json").is_file() else sorted(path for path in args.root.iterdir() if (path / "study-info.json").is_file())
    if not roots:
        raise ValueError("No study-info.json found")
    private_files = tuple(path for path in (args.env_file, args.author) if path is not None)
    results = [build(root.resolve(), author, args.pandoc, args.bibliography, args.compile_tex,
                     private_files=private_files) for root in roots]
    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
