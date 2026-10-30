"""Journal requirements, deterministic editorial preparation and bound native QA.

Numeric format, publication merit and human submission declarations are distinct.
No function in this module starts research, executes the imported repository or
infers an author's identity, biography, photograph or approval.
"""

from datetime import datetime, timezone

from contextlib import closing

import hashlib

import re

from pathlib import Path

from _publication_files import digest_file, ensure_unlinked
from _author_profile import profile_completeness
from _publication_text import mask_code, mask_fenced_code, section_headers

VENUE_LIMITS = {'ieee-software': {'word_count_scope': 'body_excluding_references_and_bios', 'max_words': 4200, 'max_abstract_words': 150, 'max_references': 15, 'visual_word_cost': 250, 'insight_count': 3}, 'ieee-security-privacy': {'word_count_scope': 'whole_article', 'min_words': 4900, 'max_words': 7200, 'max_abstract_words': 50, 'max_references': 15, 'visual_word_cost': 250}}

CS_GUIDANCE = "https://www.computer.org/publications/author-resources"

SOFTWARE_GUIDANCE = "https://www.computer.org/digital-library/magazines/so/cfp-ieee-software"

SP_GUIDANCE = "https://www.computer.org/digital-library/magazines/sp/cfp-real-world-security-impact"

AI_GUIDANCE = "https://open.ieee.org/author-guidelines-for-artificial-intelligence-ai-generated-text/"

CHECKED_AT = "2026-10-18"

SELF_COUNTS = re.compile(r"\b(?:the|this|our)\s+(?:manuscript|article|paper)\s+(?:contains?|has|includes?|is\s+within|meets?|complies?)\b.*(?:\breference\s+entries\b|\bword\s+limit\b|\bnumeric\s+(?:limits?|requirements?)\b|\babstract\s+(?:below|under|within)\b)", re.I)

SELF_COUNT_SENTENCE = re.compile(r"\b(?:the|this|our)\s+(?:manuscript|article|paper)\s+(?:contains?|has|includes?|is\s+within|meets?|complies?)\b[^.!?\n]*(?:\breference\s+entries\b|\bword\s+limit\b|\bnumeric\s+(?:limits?|requirements?)\b|\babstract\s+(?:below|under|within)\b)[^.!?\n]*[.!?]?", re.I)

EDITORIAL_MARKERS = re.compile(r"(?i:\b(?:TODO|TBD|FIXME)\b|(?:\[|<|\{\{?)\s*placeholder\b|\{\{(?:source|document|citation):)|\bPLACEHOLDER\b")

def format_guidance(venue_id: str, author_metadata: dict | None = None) -> dict:
    try:
        limits = VENUE_LIMITS[venue_id].copy()
    except KeyError as exc:
        raise ValueError("Only IEEE Software and IEEE Security & Privacy are supported") from exc
    return {"venue_id": venue_id, "checked_at": CHECKED_AT, "limits": limits,
            "source_urls": [SOFTWARE_GUIDANCE if venue_id == "ieee-software" else SP_GUIDANCE, CS_GUIDANCE, AI_GUIDANCE],
            "source_scope": "Retained venue guidance, not a fresh publisher check; recheck the current article category and rules before submission.",
            "writer_targets": {"abstract_words": 130 if venue_id == "ieee-software" else 45,
                               "whole_words": [3200, 3800] if venue_id == "ieee-software" else [5400, 6500]},
            "instructions": ["Use actual retained references with verified metadata; do not add a model-authored bibliography or numeric compliance self-attestation.",
                             "The abstract is one paragraph with no reference citations. Use unnumbered headings without ending punctuation. Current verified venue limits override the older template's sample numeric instructions.",
                             "Include a separate Keywords: term; term; term paragraph immediately after the abstract. Three to eight keywords is this skill's preparation policy, not a claimed universal journal limit.",
                             "Keep literal code escapes in code; prose must contain actual paragraph breaks, never escaped backslash-n separators.",
                             "Do not pad prose to reach a word limit. Extend supported analysis, worked decisions, comparisons, counterexamples or limitations.",
                             "Include an Acknowledgments paragraph naming the actual AI system/model, its actual use and affected sections from retained provenance; do not claim unrecorded human review, responsibility or approval.",
                             "Use only supplied, approved author facts. Missing biography, photo and declarations are pending human preparation, never invented."],
            "human_information_available": [{"id": entry["id"], **{key: bool(entry["author"].get(key)) for key in ("display_name", "affiliation", "email", "orcid", "biography")}}
                                            for entry in (author_metadata or {}).get("authors", [])]}

def _parts(markdown: str) -> dict[str, str]:
    masked = mask_fenced_code(markdown)
    matches = section_headers(markdown)
    return {match[1].strip().casefold(): masked[match.end():matches[i + 1].start() if i + 1 < len(matches) else len(markdown)].strip()
            for i, match in enumerate(matches)}

def _words(text: str) -> int:
    text = re.sub(r"<!--[\s\S]*?-->", "", text)
    text = re.sub(r"!?\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"[`*_#|]", " ", text)
    return len(text.split())

def _item(identifier, status, detail, *, stage="manuscript", source=CS_GUIDANCE):
    return {"id": identifier, "label": identifier.replace("_", " "), "status": status, "stage": stage, "reason": detail, "source_url": source}

def _summarize(items, metrics, venue_id):
    content = [item for item in items if item["stage"] in {"manuscript", "native"}]
    blockers = [item["reason"] for item in content if item["status"] == "failed"]
    pending = [item["reason"] for item in items if item["status"] == "pending"]
    return {"venue_id": venue_id, "checked_at": CHECKED_AT, "items": items, "metrics": metrics,
            "blockers": blockers, "pending": pending,
            "content_format_passed": bool(content) and all(item["status"] == "passed" for item in content),
            "submission_ready": all(item["status"] == "passed" for item in items),
            "prefilter_passed": not blockers, "scientific_merit_assessed": False, "author_approval_inferred": False}

def _insights(markdown: str) -> list[list[str]]:
    """Find explicitly labelled, contiguous insight lists, not arbitrary bullets."""
    prose = mask_code(markdown)
    lines, blocks = prose.splitlines(), {}
    label = re.compile(r"(?:three\s+)?actionable\s+insights(?:\s+follow)?\s*:", re.I)
    bullet = re.compile(r"^( {0,3})[-*+][ \t]+(.*)$")

    def explicit_label(line):
        text = line.strip()
        if text.startswith(("**", "__")) and text.endswith(text[:2]):
            text = text[2:-2].strip()
        return label.fullmatch(text) is not None

    for index, line in enumerate(lines):
        heading = re.fullmatch(r" {0,3}#{1,6}\s+(.+?)\s*#*", line)
        if not (explicit_label(line) or heading and re.search(r"\bactionable insights\b", heading[1], re.I)):
            continue
        start = index + 1
        # A heading and a following explicit introduction may label the same list.
        while start < len(lines) and (not lines[start].strip() or explicit_label(lines[start])):
            start += 1
        first = bullet.match(lines[start]) if start < len(lines) else None
        if first is None:
            blocks.setdefault(-index - 1, [])
            continue
        indent, cursor, items = len(first[1]), start, []
        while cursor < len(lines):
            current = lines[cursor]
            match = bullet.match(current)
            if match and len(match[1]) == indent:
                items.append(match[2].strip())
            elif current.strip() and len(current) - len(current.lstrip(" ")) < indent + 2:
                break
            cursor += 1
        blocks.setdefault(start, items)
    return list(blocks.values())

def _author_blocks(front: str, authors: list[dict]) -> tuple[dict[str, str], bool]:
    """Associate plain byline blocks with exact leading names, never global facts."""
    starts = list(re.finditer(r"(?m)^[ \t]*<!-- paper-factory:authors:start -->[ \t]*$", front))
    ends = list(re.finditer(r"(?m)^[ \t]*<!-- paper-factory:authors:end -->[ \t]*$", front))
    if starts or ends:
        if len(starts) != 1 or len(ends) != 1 or starts[0].start() >= ends[0].start():
            return {}, False
        front = front[starts[0].end():ends[0].start()]
    names = {}
    for author in authors:
        name = author["author"].get("display_name")
        if name:
            names.setdefault(name, []).append(author["id"])
    found, current, lines = [], None, []
    for line in front.splitlines():
        first = line.strip().split("|", 1)[0].strip()
        matched = names.get(first, [])
        if len(matched) == 1:
            if current is not None:
                found.append((current, "\n".join(lines)))
            current, lines = matched[0], [line]
        elif current is not None:
            lines.append(line)
    if current is not None:
        found.append((current, "\n".join(lines)))
    bindings = {identifier: text for identifier, text in found
                if sum(other == identifier for other, _ in found) == 1}
    return bindings, [identifier for identifier, _ in found] == [entry["id"] for entry in authors]


def validate_manuscript(markdown: str, venue_id: str, author_metadata: dict | None = None) -> dict:
    """Pure prefilter. Its word estimate never substitutes for native Word QA."""
    guidance = format_guidance(venue_id, author_metadata)
    limits, parts = guidance["limits"], _parts(markdown)
    abstract = parts.get("abstract", "")
    prose = mask_code(markdown)
    keyword_lines = re.findall(r"(?mi)^Keywords:\s*(.+)$", prose)
    keywords = [word.strip() for line in keyword_lines for word in line.split(";") if word.strip()]
    abstract = re.sub(r"(?mi)^Keywords:.*$", "", abstract).strip()
    references = parts.get("references", "")
    entries = list(re.finditer(r"(?ms)^\[(\d+)\]\s+(.+?)(?=\n\[\d+\]\s|\Z)", references))
    numbers = [int(entry[1]) for entry in entries]
    insight_blocks = _insights(markdown)
    acknowledgment = next((text for heading, text in parts.items()
                           if re.search(r"\backnowledg(?:e)?ments?\b", heading)), "")
    ai_disclosed = bool(re.search(r"\b(?:artificial intelligence|generative AI|AI[- ]assist\w*|language model|OpenAI|ChatGPT|GPT(?:[- ]\d+(?:\.\d+)*(?:-[a-z0-9]+)*)?)\b", acknowledgment, re.I)
                        and re.search(r"\b(?:draft\w*|edit\w*|analy\w*|research|writ\w*|generat\w*)\b", acknowledgment, re.I)
                        and re.search(r"\b(?:sections?|throughout|entire|abstract|introduction|manuscript|article|methods?|prose|writing)\b", acknowledgment, re.I))
    insights = sum(map(len, insight_blocks))
    reference_heading = next((match for match in section_headers(markdown) if match[1].strip().casefold() == "references"), None)
    body = markdown[:reference_heading.start()] if reference_heading else markdown
    body_prose = prose[:reference_heading.start()] if reference_heading else prose
    table_count = len(re.findall(r"(?m)^\|?\s*:?-{3,}.*\|.*$", prose))
    figure_count = len(re.findall(r"!\[[^\]]*\]\([^)]+\)", prose))
    metrics = {"abstract_words_estimate": _words(abstract), "reference_count": len(entries), "actionable_insights": insights,
               "actionable_insight_blocks": len(insight_blocks),
               "keywords": keywords, "whole_words_estimate": _words(markdown), "body_words_estimate": _words(body),
               "visual_count_estimate": table_count + figure_count,
               "word_count_status": "estimate_only_native_Word_measurement_required"}
    source = SOFTWARE_GUIDANCE if venue_id == "ieee-software" else SP_GUIDANCE
    items = [_item("abstract", "passed" if abstract and _words(abstract) <= limits["max_abstract_words"] else "failed",
                   f"Abstract estimate {_words(abstract)} words; maximum {limits['max_abstract_words']}.", source=source),
             _item("abstract_structure", "passed" if len(re.split(r"\n\s*\n", abstract)) == 1 and not re.search(r"\[\d+(?:[ \t]*[-–,;][ \t]*\d+)*[ \t]*\]|\{\{(?:source|document|citation):", abstract) else "failed",
                   "The publisher template requires a single-paragraph abstract without reference citations."),
             _item("reference_count", "passed" if entries and len(entries) <= limits["max_references"] and numbers == list(range(1, len(entries) + 1)) else "failed",
                   f"Canonical references: {len(entries)}; require sequential unique labels and at most {limits['max_references']}.", source=source),
             _item("keywords", "passed" if 3 <= len(keywords) <= 8 and len({k.casefold() for k in keywords}) == len(keywords) else "failed",
                   "Provide three to eight distinct article keywords as a Keywords: paragraph (skill preparation policy)."),
             _item("self_attestation", "failed" if SELF_COUNTS.search(body_prose) else "passed", "Numeric compliance claims and generated reference-count self-attestations must not remain in manuscript prose."),
             _item("escaped_prose", "failed" if re.search(r"\\n(?:\\n|## |[-*] |Keywords:)", prose) else "passed",
                   "Use real paragraph breaks in prose; code escapes are preserved."),
             _item("placeholders", "failed" if EDITORIAL_MARKERS.search(body_prose) else "passed",
                   "Resolve placeholders and evidence tokens before native format preflight."),
             _item("ai_disclosure", "passed" if ai_disclosed else "failed", "Acknowledge actual AI tools, their use and affected manuscript parts from retained provenance. Disclosure presence is checked; it never asserts author approval or verifies tool-name accuracy.", source=AI_GUIDANCE)]
    estimate = (_words(body) if venue_id == "ieee-software" else _words(markdown)) + limits["visual_word_cost"] * (table_count + figure_count)
    items.append(_item("word_count", "passed" if limits.get("min_words", 0) <= estimate <= limits["max_words"] else "failed",
                       f"Markdown estimate including visual allowances: {estimate}; native Word measurement is definitive.", source=source))
    if venue_id == "ieee-software":
        complete_insights = len(insight_blocks) == 1 and len(insight_blocks[0]) == 3 and all(insight_blocks[0])
        items.append(_item("actionable_insights", "passed" if complete_insights else "failed",
                           f"Provide one explicitly labelled block of exactly three actionable-insight bullets; found {insights} bullets in {len(insight_blocks)} blocks.", source=source))
    metadata_issues = []
    for entry in entries:
        value = entry[2]
        descriptive = re.sub(r"https?://[^\s)]+", "", value)
        descriptive = re.sub(r"\[[\s]*\]\([^)]*\)", "", descriptive)
        if not re.search(r"[A-Za-z]{3,}", descriptive) or re.match(r"\[https?://", value):
            metadata_issues.append(int(entry[1]))
    items.append(_item("reference_metadata", "failed" if metadata_issues else "passed",
                       "Use retained descriptive source titles and available author/year/publication/access metadata; raw-URL-only entries: " + str(metadata_issues) + "."))
    authors = profile_completeness(author_metadata)
    abstract_heading = next((match for match in section_headers(markdown) if match[1].strip().casefold() == "abstract"), None)
    byline = prose[:abstract_heading.start()] if abstract_heading else ""
    biography_text = "\n".join(text for heading, text in parts.items() if heading in {"author biography", "author biographies", "about the author", "about the authors"})
    ordered = (author_metadata or {}).get("authors", [])
    bindings, order_matches = _author_blocks(byline, ordered)
    biography_blocks = re.split(r"\n[ \t]*\n", biography_text)
    for row, entry in zip(authors, ordered):
        author = entry["author"]
        block = bindings.get(entry["id"], "")
        values = {cell.strip() for line in block.splitlines() for cell in line.split("|") if cell.strip()}
        row["missing_byline_values"] = [field for field in ("display_name", "affiliation", "email", "orcid")
                                        if author.get(field) and author[field] not in values
                                        and not (field == "orcid" and "ORCID " + author[field] in values)]
        row["byline_order_matches"] = order_matches
        row["supplied_biography_in_manuscript"] = bool(author.get("biography") and author.get("display_name") and any(
            paragraph.startswith(author["display_name"] + ". ") and author["biography"] in paragraph for paragraph in biography_blocks))
    metrics["authors"] = authors
    metrics["missing_biography_fields"] = {entry["id"]: [field for field in entry["missing_fields"] if field == "biography" or field.startswith("biography_details.")]
                                           for entry in authors}
    biographies_complete = bool(authors) and not any(metrics["missing_biography_fields"].values()) and all(row["supplied_biography_in_manuscript"] for row in authors)
    items.extend([_item("author_identity", "pending", "Author identity, affiliation, contact and ORCID require the author's actual confirmation; public ORCID data is not consent.", stage="submission"),
                  _item("author_profile_binding", "passed" if authors and order_matches and all(not row["missing_byline_values"] for row in authors) else "pending", "Every supplied author's identity fields must occur in that author's distinct byline block and in the supplied author order. Missing or unmatched fields are recorded per author; matching text is not consent.", stage="submission"),
                  _item("author_biography", "passed" if biographies_complete else "pending", "Biography text and all factual details are checked separately for every supplied author. Missing facts are listed in metrics.authors; factual accuracy still requires actual author confirmation.", stage="submission"),
                  _item("author_declarations", "pending", "The author must confirm actual originality, rights, AI disclosure, author approval and final submission declarations.", stage="submission")])
    if venue_id == "ieee-software":
        items.append(_item("author_photo", "pending", "IEEE Software requests a photograph of each author with submission; obtain the actual approved image.", stage="submission", source=source))
    return _summarize(items, metrics, venue_id)

def render_pdf_pages(pdf: Path, output_root: Path) -> dict:
    """Rasterize every actual PDF page, bound to immutable output/image hashes.

    Automatic character, crop-boundary and raster checks expose basic defects;
    they cannot establish that layout, figures, equations or wording look right.
    """
    import pypdfium2 as pdfium
    from PIL import Image, ImageChops
    from io import BytesIO

    ensure_unlinked(pdf)
    ensure_unlinked(output_root)
    if pdf.stat().st_size > 32 * 1024 * 1024:
        raise ValueError("Journal PDF QA exceeds its 32 MiB bounded file size")
    source_sha = digest_file(pdf)
    pages, issues = [], []
    # In-memory opening avoids PDFium's native Windows long-path restrictions.
    with pdfium.PdfDocument(pdf.read_bytes()) as document:
        if not 1 <= len(document) <= 100:
            raise ValueError("Journal PDF QA requires one to one hundred pages")
        output_root.mkdir(parents=True, exist_ok=True)
        for index in range(len(document)):
            with closing(document[index]) as page:
                width, height = page.get_size()
                if not 0 < width <= 2000 or not 0 < height <= 2000:
                    raise ValueError("Journal PDF page dimensions exceed bounded review limits")
                with closing(page.get_textpage()) as textpage:
                    text = textpage.get_text_range()
                    outside = 0
                    for char in range(textpage.count_chars()):
                        if not textpage.get_text_range(char, 1).strip():
                            continue
                        left, bottom, right, top = textpage.get_charbox(char)
                        outside += int(left < -1 or bottom < -1 or right > width + 1 or top > height + 1)
                if not text.strip() or "\ufffd" in text or "\x00" in text:
                    issues.append(f"Page {index + 1}: empty/replacement/null extractable text requires correction or explicit visual verification.")
                if outside:
                    issues.append(f"Page {index + 1}: {outside} character boxes extend outside the page bounds.")
                bitmap = page.render(scale=1.5)
                try:
                    rendered = bitmap.to_pil().convert("RGB")
                    difference = ImageChops.difference(rendered, Image.new("RGB", rendered.size, "white"))
                    if difference.getbbox() is None:
                        issues.append(f"Page {index + 1}: raster is blank.")
                    buffer = BytesIO()
                    rendered.save(buffer, format="PNG")
                    content = buffer.getvalue()
                    image_path = output_root / f"page-{index + 1:03}.png"
                    ensure_unlinked(image_path)
                    if image_path.exists():
                        if image_path.read_bytes() != content:
                            raise ValueError("Page QA cannot replace different existing PNG bytes")
                    else:
                        with image_path.open("xb") as stream:
                            stream.write(content)
                    pages.append({"number": index + 1, "text_chars": len(text), "text_sha256": hashlib.sha256(text.encode()).hexdigest(),
                                  "width_points": width, "height_points": height, "outside_page_characters": outside,
                                  "image_path": str(image_path.resolve()), "image_sha256": hashlib.sha256(content).hexdigest(),
                                  "width_pixels": rendered.width, "height_pixels": rendered.height})
                finally:
                    bitmap.close()
    if digest_file(pdf) != source_sha:
        raise ValueError("The actual PDF changed during whole-page inspection")
    return {"engine": "PDFium rendering of the actual exported PDF", "pdf_sha256": source_sha,
            "page_count": len(pages), "pages": pages, "issues": issues, "automatic_checks": not issues,
            "status": "failed" if issues else "pending_visual_review", "visual_review_status": "pending",
            "scope": "All actual PDF pages rendered; text/glyph/page-boundary checks are structural evidence, never human visual approval."}

def _native_structure(docx: Path, software: bool) -> dict:
    from docx import Document
    from docx.oxml.ns import qn
    document = Document(docx)
    paragraphs = document.paragraphs
    required = {"Title_document", "Abstract", "ReferenceHead", "Bib_entry"}
    styles = {paragraph.style.name for paragraph in paragraphs}
    bibliography = [paragraph for paragraph in paragraphs if paragraph.style.name == "Bib_entry"]
    labels = [int(match[1]) for paragraph in bibliography if (match := re.match(r"^(\d+)\.\s", paragraph.text))]
    known = {str(number) for number in labels}
    issues = []
    if not required <= styles or not paragraphs or paragraphs[0].style.name != "Title_document":
        issues.append("Native publisher title, abstract and reference styles are missing.")
    if not any(section._sectPr.find(qn("w:cols")).get(qn("w:num")) == "2" for section in document.sections):
        issues.append("Publisher two-column body section is missing.")
    if labels != list(range(1, len(bibliography) + 1)):
        issues.append("Native bibliography labels must be sequential numbered-period entries.")
    before_refs = []
    for paragraph in paragraphs:
        if paragraph.style.name == "ReferenceHead":
            break
        if paragraph.style.name not in {"Authors", "Title_document"}:
            before_refs.append(paragraph)
    for paragraph in before_refs:
        if paragraph.style.name.casefold() in {"source code", "code"}:
            continue
        for run in paragraph.runs:
            properties = run._r.rPr
            inline_style = properties.find(qn("w:rStyle")) if properties is not None else None
            code = inline_style is not None and "verbatim" in inline_style.get(qn("w:val"), "").casefold()
            if software and not code and re.search(r"(?<![\w])\[(?:" + "|".join(known or {"0"}) + r")\]", run.text):
                issues.append("Software retains a bracketed body citation instead of publisher superscript typography.")
        if not software and re.search(r"[.!?]\s+\[\d+\]", paragraph.text):
            issues.append("S&P body citation appears after terminal punctuation instead of inside it.")
    keywords = [paragraph.text for paragraph in paragraphs if paragraph.text.casefold().startswith("keywords:")]
    if any(paragraph.style.name == "Abstract" and paragraph.text.casefold().startswith("keywords:") for paragraph in paragraphs):
        issues.append("Keywords must not be counted as part of the Abstract style.")
    return {"issues": sorted(set(issues)), "reference_count": len(bibliography), "keywords": keywords,
            "citation_style": "superscript_outside_punctuation" if software else "bracketed_inside_punctuation"}

def compliance_report(markdown: Path, docx: Path | None, venue_id: str, *, pdf: Path | None = None,
                      qa_root: Path | None = None, author_metadata: dict | None = None,
                      conversion_receipts: dict | None = None, visual_review: dict | None = None) -> dict:
    """Definitive numeric/native preflight plus explicit human preparation items.

    The caller retains this result with the exact reviewed manuscript. Optional
    receipts are required for a passed binding, not guessed from file names.
    Page approval must identify the exact PDF and each raster reviewed.
    """
    from _publication_templates import official_template, measure_word
    ensure_unlinked(markdown)
    initial_sha = digest_file(markdown)
    base = validate_manuscript(markdown.read_text(encoding="utf-8"), venue_id, author_metadata)
    items, metrics = base["items"], base["metrics"]
    limits = format_guidance(venue_id)["limits"]
    source = SOFTWARE_GUIDANCE if venue_id == "ieee-software" else SP_GUIDANCE
    _, template = official_template(venue_id)
    binding = {"markdown_sha256": initial_sha, "template_sha256": template["sha256"], "template_provenance_sha256": template["provenance_sha256"]}
    native, pages = None, None
    if docx is None or not docx.is_file():
        items.append(_item("native_word", "pending", "Create the official publisher-template Word file, then measure its actual read-only Word counts.", stage="native", source=source))
    else:
        ensure_unlinked(docx)
        binding["docx_sha256"] = digest_file(docx)
        try:
            native = measure_word(docx)
            actual = native["metrics"]
            metrics["native_word"] = actual
            metrics["word_count_status"] = "measured_native_Word"
            effective = (actual["body_words"] - actual["body_table_words"] + 250 * (actual["body_tables"] + actual["body_figures"])) if venue_id == "ieee-software" else (actual["whole_words"] + 250 * (actual["tables"] + actual["figures"]))
            metrics["effective_words"] = effective
            metrics["word_count_formula"] = "Body Word words minus body table-text words plus 250 per body visual" if venue_id == "ieee-software" else "Whole Word words including table text plus 250 per table/figure"
            replacements = {"abstract": (0 < actual["abstract_words"] <= limits["max_abstract_words"], f"Actual Word abstract: {actual['abstract_words']}; maximum {limits['max_abstract_words']}."),
                            "word_count": (limits.get("min_words", 0) <= effective <= limits["max_words"], f"Actual effective Word count: {effective}; allowed {limits.get('min_words', 0)}–{limits['max_words']}."),
                            "reference_count": (actual["reference_heading_found"] and 0 < actual["reference_count"] <= limits["max_references"], f"Native reference entries: {actual['reference_count']}; maximum {limits['max_references']}.")}
            for item in items:
                if item["id"] in replacements:
                    passed, item["reason"] = replacements[item["id"]]
                    item.update(status="passed" if passed else "failed", stage="native")
            structure = _native_structure(docx, venue_id == "ieee-software")
            items.append(_item("native_word", "passed" if not structure["issues"] else "failed", "; ".join(structure["issues"]) or "Actual publisher styles, columns, citation conventions and numbered bibliography retained.", stage="native", source=template["download_url"]))
            metrics["native_structure"] = structure
            metrics["missing_template_fonts"] = actual.get("missing_template_fonts", [])
        except Exception as exc:
            items.append(_item("native_word", "failed", str(exc), stage="native", source=source))
    receipts = conversion_receipts or {}
    binding_issues = []
    for name, path in (("docx", docx), ("pdf", pdf)):
        if path is None or not path.is_file():
            binding_issues.append(f"Actual {name.upper()} export is absent.")
            continue
        ensure_unlinked(path)
        binding[name + "_sha256"] = digest_file(path)
        receipt = receipts.get(name, {})
        official = receipt.get("official_template", {})
        if receipt.get("input_sha256") != initial_sha or receipt.get("output_sha256") != binding[name + "_sha256"] or official.get("sha256") != template["sha256"] or official.get("venue_id") != venue_id:
            binding_issues.append(f"{name.upper()} conversion receipt is absent or differs from actual manuscript/output/publisher-template bytes.")
    if pdf is not None and pdf.is_file() and docx is not None and docx.is_file():
        pdf_receipt = receipts.get("pdf", {})
        if pdf_receipt.get("source_docx_sha256") != binding["docx_sha256"] or pdf_receipt.get("word_rendering", {}).get("source_docx_sha256") != binding["docx_sha256"]:
            binding_issues.append("PDF is not bound to the exact saved official Word file.")
        if qa_root is not None:
            try:
                pages = render_pdf_pages(pdf, qa_root)
                items.append(_item("all_pages", "passed" if pages["automatic_checks"] else "failed", "; ".join(pages["issues"]) or f"All {pages['page_count']} actual PDF pages rendered and structurally checked; human visual review remains separate.", stage="native"))
            except Exception as exc:
                items.append(_item("all_pages", "failed", str(exc), stage="native"))
    if pages is None:
        items.append(_item("all_pages", "pending", "Render and structurally check every actual exported PDF page.", stage="native"))
    items.append(_item("export_binding", "failed" if binding_issues else "passed", "; ".join(binding_issues) or "Actual Markdown, publisher template, Word and Word-native PDF hashes match their conversion receipts.", stage="native", source=template["download_url"]))
    review = visual_review or {}
    reviewed = bool(pages and review.get("pdf_sha256") == binding.get("pdf_sha256") and review.get("approved") is True and review.get("reviewer") and review.get("reviewed_at") and review.get("page_image_sha256") == [page["image_sha256"] for page in pages["pages"]])
    items.append(_item("visual_review", "passed" if reviewed else "pending", "An actual reviewer must inspect every bound PDF page; automatic bounds, glyph and PNG checks are not approval.", stage="submission"))
    if digest_file(markdown) != initial_sha or docx is not None and docx.is_file() and digest_file(docx) != binding["docx_sha256"] or pdf is not None and pdf.is_file() and digest_file(pdf) != binding["pdf_sha256"]:
        raise ValueError("The manuscript or native exports changed during journal compliance inspection")
    result = _summarize(items, metrics, venue_id)
    return {**result, "binding": binding, "word_measurement": native, "page_qa": pages,
            "checked_at_utc": datetime.now(timezone.utc).isoformat(), "guidance": format_guidance(venue_id, author_metadata)}
