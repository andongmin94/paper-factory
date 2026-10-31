"""Verified publisher Word templates and their native Windows PDF rendering."""

from copy import deepcopy
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
from tempfile import TemporaryDirectory

from _publication_files import digest_file, ensure_unlinked


ROOT = Path(__file__).resolve().parents[1] / "assets" / "templates"


def _replace_native_text(paragraph, start, end, pieces):
    """Replace a visible text span without flattening links or inline styles."""
    from docx.oxml import OxmlElement
    from docx.text.run import Run
    runs, position = [], 0
    for element in paragraph._p.xpath(".//w:r"):
        run = Run(element, paragraph)
        text = run.text
        runs.append((position, position + len(text), run))
        position += len(text)
    affected = [(a, b, run) for a, b, run in runs if a < end and b > start]
    if not affected:
        return
    a, _, first = affected[0]
    last_start, _, last = affected[-1]
    prefix, suffix = first.text[:start - a], last.text[end - last_start:]
    first_properties = deepcopy(first._r.rPr)
    last_properties = deepcopy(last._r.rPr)
    first.text = prefix
    for _, _, run in affected[1:]:
        run.text = ""
    anchor = first._r
    for text, superscript in pieces:
        if not text:
            continue
        element = OxmlElement("w:r")
        if first_properties is not None:
            element.append(deepcopy(first_properties))
        run = Run(element, paragraph)
        run.text = text
        if superscript is not None:
            run.font.superscript = superscript
        anchor.addnext(element)
        anchor = element
    if suffix:
        element = OxmlElement("w:r")
        if last_properties is not None:
            element.append(last_properties)
        Run(element, paragraph).text = suffix
        anchor.addnext(element)


def _citation_typography(document, reference, *, software):
    """Apply the different citation instructions in the two acquired templates."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.text.run import Run

    bibliography = [p for p in document.paragraphs if p.style.name == "Bib_entry"]
    known = set()
    labels = 0
    reference_sample = next((p for p in reference.paragraphs if p.style.name == "Bib_entry"), None)
    for paragraph in bibliography:
        match = re.match(r"^\[(\d+)\]\s*", paragraph.text)
        if match:
            known.add(match[1])
            _replace_native_text(paragraph, 0, match.end(), [(match[1] + ". ", False)])
            labels += 1
        else:
            numbered = re.match(r"^(\d+)\.\s", paragraph.text)
            if numbered:
                known.add(numbered[1])
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if reference_sample is not None and reference_sample._p.pPr is not None:
            # The S&P template's hanging indent is a direct paragraph setting;
            # Software supplies it through Bib_entry. Preserve each authority.
            for name in ("ind", "tabs"):
                source = reference_sample._p.pPr.find(qn("w:" + name))
                if source is not None:
                    properties = paragraph._p.get_or_add_pPr()
                    old = properties.find(qn("w:" + name))
                    if old is not None:
                        properties.remove(old)
                    properties.append(deepcopy(source))
    callout = re.compile(r'''(?:(?P<before>[.,;:!?])(?P<close>["'”’)]*)[ \t]*)?[ \t]*(?P<cites>\[\d+\](?:[ \t]*(?:,[ \t]*)?\[\d+\])*)(?P<after>[.,;:!?])?(?P<afterclose>["'”’)]*)''')
    changed, unknown = 0, set()
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        paragraphs.extend(p for row in table.rows for cell in row.cells for p in cell.paragraphs)
    for paragraph in paragraphs:
        if paragraph.style.name in {"Title_document", "Authors", "Bib_entry", "ReferenceHead", "AuthorBio", "AuthorBioHead"}:
            continue
        text, protected = "", []
        for element in paragraph._p.xpath(".//w:r"):
            run = Run(element, paragraph)
            content = run.text
            style = " ".join(element.xpath("./w:rPr/w:rStyle/@w:val")).casefold()
            if run.style is not None:
                style += " " + run.style.name.casefold()
            guarded = any(term in style for term in ("verbatim", "code")) or bool(element.xpath("ancestor::w:hyperlink"))
            text += content
            protected.extend([guarded] * len(content))
        if paragraph.style.name.casefold() in {"source code", "code"}:
            continue
        for match in reversed(list(callout.finditer(text))):
            if any(protected[match.start():match.end()]):
                continue
            numbers = re.findall(r"\d+", match["cites"])
            if not set(numbers).issubset(known):
                unknown.update(set(numbers) - known)
                continue
            punctuation = (match["before"] or "") + (match["after"] or "")
            closing = (match["close"] or "") + (match["afterclose"] or "")
            if software:
                pieces = [(punctuation + closing, False), (",".join(numbers), True)]
            else:
                space = " " if closing or match.start() and not text[match.start() - 1].isspace() else ""
                pieces = [(closing + space + ", ".join("[" + n + "]" for n in numbers), False), (punctuation, False)]
            _replace_native_text(paragraph, match.start(), match.end(), pieces)
            changed += 1
    return {"callout_groups_formatted": changed, "reference_labels_formatted": labels,
            "unknown_reference_numbers_preserved": sorted(unknown),
            "callout_style": "Superscript numbers outside punctuation" if software else "Bracketed numbers inside punctuation",
            "bibliography_metadata": "Supplied names, descriptions, titles, URLs and dates preserved; unavailable bibliographic metadata is not invented and still requires editorial verification."}


def official_template(venue_id: str) -> tuple[Path, dict]:
    manifest = json.loads((ROOT / "provenance.json").read_bytes())
    try:
        record = manifest["templates"][venue_id]
    except KeyError as exc:
        raise ValueError("An official template is configured only for IEEE Software and IEEE Security & Privacy") from exc
    path = ROOT / record["filename"]
    if not path.is_file() or digest_file(path) != record["sha256"]:
        raise ValueError("The bundled official publisher template is missing or differs from its provenance SHA")
    return path, {**record, "venue_id": venue_id, "authority": manifest["authority"],
                  "provenance_sha256": digest_file(ROOT / "provenance.json")}


def apply_official_styles(output: Path, template: Path) -> dict:
    """Map parsed content to the publisher's actual named styles and sections.

    Pandoc imports the reference package, including its styles. Its final section
    alone cannot represent the template's one-column front and two-column body;
    copy those section masters explicitly, removing sample headers and artwork.
    """
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shape import InlineShape
    from docx.shared import Twips
    from docx.text.paragraph import Paragraph

    document, reference = Document(output), Document(template)
    paragraphs = document.paragraphs
    if not paragraphs or paragraphs[0].style.name != "Heading 1":
        raise ValueError("The official magazine export requires one leading Markdown title")
    title = paragraphs[0]
    title.style = document.styles["Title_document"]
    abstract_heading = next((p for p in paragraphs[1:] if p.style.name == "Heading 2" and p.text.strip().casefold() == "abstract"), None)
    if abstract_heading is None:
        raise ValueError("The official magazine export requires a distinct Abstract section")
    abstract_index = paragraphs.index(abstract_heading)
    body_heading = next((p for p in paragraphs[abstract_index + 1:] if p.style.name.startswith("Heading ")), None)
    if body_heading is None:
        raise ValueError("The official magazine export requires body sections after its abstract")
    body_index = paragraphs.index(body_heading)
    for paragraph in paragraphs[1:abstract_index]:
        paragraph.style = document.styles["Authors"]
    for paragraph in paragraphs[abstract_index + 1:body_index]:
        paragraph.style = document.styles["Para" if paragraph.text.lstrip().casefold().startswith("keywords:") else "Abstract"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # The template prints the abstract label inline, rather than as a section.
    abstract_text = paragraphs[abstract_index + 1]
    if not abstract_text.text.strip():
        raise ValueError("The official magazine abstract is empty")
    label = OxmlElement("w:r")
    label_text = OxmlElement("w:t")
    label_text.text = "Abstract—"
    label.append(label_text)
    abstract_text._p.insert(1 if abstract_text._p.pPr is not None else 0, label)
    abstract_heading._p.getparent().remove(abstract_heading._p)
    section_style = "Para"
    for paragraph in paragraphs[body_index:]:
        name = paragraph.style.name
        if name.startswith("Heading "):
            level = int(name.split()[-1])
            paragraph.style = document.styles[f"Heading {max(1, level - 1)}"]
            heading = paragraph.text.strip().casefold()
            section_style = "Para"
            if heading == "references":
                paragraph.style = document.styles["ReferenceHead"]
                section_style = "Bib_entry"
            elif heading in {"author biography", "author biographies", "about the author", "about the authors"}:
                paragraph.style = document.styles["AuthorBioHead"]
                section_style = "AuthorBio"
            elif heading in {"acknowledgments", "acknowledgements", "acknowledgment", "acknowledgement"}:
                paragraph.style = document.styles["AckHead"]
                section_style = "AckPara"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.keep_with_next = True
        elif name in {"Normal", "Body Text", "First Paragraph", "Compact"}:
            paragraph.style = document.styles[section_style]
        elif name in {"Table Caption", "Image Caption", "Figure Caption"}:
            paragraph.style = document.styles["TableCaption" if "Table" in name else "FigureCaption"]

    def master(section):
        element = deepcopy(section._sectPr)
        # Sample running heads, placeholder first-page labels, and page restart
        # are publisher instructions, not author content.
        for name in ("headerReference", "footerReference", "pgNumType", "titlePg"):
            for child in element.findall(qn("w:" + name)):
                element.remove(child)
        kind = element.find(qn("w:type"))
        if kind is None:
            kind = OxmlElement("w:type")
            element.insert(0, kind)
        kind.set(qn("w:val"), "continuous")
        return element

    front = master(reference.sections[0])
    body_section = next(s for s in reversed(list(reference.sections))
                        if s._sectPr.find(qn("w:cols")).get(qn("w:num")) == "2")
    body = master(body_section)
    body_element = document._element.body
    for old in list(body_element.iter(qn("w:sectPr"))):
        old.getparent().remove(old)
    preceding = body_heading._p.xpath("preceding-sibling::w:p[1]")
    if not preceding:
        raise ValueError("The native manuscript has no abstract paragraph before its body")
    before_body = preceding[0]
    before_body.get_or_add_pPr().append(front)
    body_element.append(body)

    # Pandoc sizes images before the publisher's two-column body is restored.
    # Fit body images to one column and page without changing their media.
    columns = body.find(qn("w:cols"))
    column_count = int(columns.get(qn("w:num")))
    column_gap = Twips(int(columns.get(qn("w:space"), "0")))
    column_width = int((body_section.page_width - body_section.left_margin - body_section.right_margin
                        - column_gap * (column_count - 1)) // column_count)
    page_height = int(body_section.page_height - body_section.top_margin - body_section.bottom_margin)
    for paragraph in paragraphs[body_index:]:
        inlines = paragraph._p.xpath(".//wp:inline")
        if inlines and not paragraph.text.strip():
            paragraph.paragraph_format.first_line_indent = 0
        for inline in inlines:
            shape = InlineShape(inline)
            scale = min(column_width / shape.width if shape.width > column_width else 1,
                        page_height / shape.height if shape.height > page_height else 1)
            if scale < 1:
                width, height = round(shape.width * scale), round(shape.height * scale)
                shape.width, shape.height = width, height

    # Tables need the full printable width. Retain exact content and surround
    # them with the template's own one-column section, then return to its body.
    def table_caption(paragraph):
        return (paragraph.style.name in {"TableCaption", "Table Caption"} or
                re.match(r"^Table\s+\d+[.:]\s*\S", paragraph.text.strip(), re.I))

    for table in document.tables:
        previous = table._tbl.getprevious()
        preceding = []
        candidate = previous
        while candidate is not None and candidate.tag == qn("w:p"):
            paragraph = Paragraph(candidate, document._body)
            preceding.append(paragraph)
            if paragraph.text.strip():
                break
            candidate = candidate.getprevious()
        if not preceding or not table_caption(preceding[-1]):
            following = table._tbl.getnext()
            while following is not None and following.tag == qn("w:p"):
                paragraph = Paragraph(following, document._body)
                if paragraph.text.strip():
                    if table_caption(paragraph):
                        # Markdown may place the caption after the table. The
                        # publisher's source places it before; move the exact
                        # XML so text, callouts and runs remain intact.
                        table._tbl.addprevious(following)
                        preceding = [paragraph]
                    break
                following = following.getnext()
        if preceding and table_caption(preceding[-1]):
            for paragraph in preceding:
                paragraph.paragraph_format.keep_with_next = True
            caption = preceding[-1]
            caption.style = document.styles["TableCaption"]
            caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
            caption.paragraph_format.keep_together = True
            # A sectPr ends the section containing its paragraph. Put the
            # two-column ending BEFORE the caption, so caption and table share
            # the following one-column section instead of splitting widths.
            anchor = caption._p
        else:
            anchor = table._tbl
        # Keep the section break off scientific prose. Word can otherwise
        # distribute the last line of a justified paragraph before the table.
        boundary = OxmlElement("w:p")
        properties = boundary.get_or_add_pPr()
        spacing = OxmlElement("w:spacing")
        for name, value in (("before", "0"), ("after", "0"), ("line", "1"), ("lineRule", "exact")):
            spacing.set(qn("w:" + name), value)
        properties.append(spacing)
        properties.append(deepcopy(body))
        anchor.addprevious(boundary)
        after = OxmlElement("w:p")
        after.get_or_add_pPr().append(deepcopy(front))
        table._tbl.addnext(after)
        width = int(reference.sections[0].page_width - reference.sections[0].left_margin - reference.sections[0].right_margin)
        for column in table.columns:
            column.width = width // len(table.columns)
        for index, row in enumerate(table.rows):
            properties = row._tr.get_or_add_trPr()
            if index == 0 and properties.find(qn("w:tblHeader")) is None:
                properties.append(OxmlElement("w:tblHeader"))
            if properties.find(qn("w:cantSplit")) is None:
                properties.append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                cell.width = width // len(row.cells)
                for paragraph in cell.paragraphs:
                    paragraph.style = document.styles["Para"]
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    paragraph.paragraph_format.first_line_indent = 0
                    if properties.find(qn("w:tblHeader")) is not None:
                        paragraph.paragraph_format.keep_with_next = True
                        for run in paragraph.runs:
                            run.bold = True
    # These reference headers include literal HEAD / editor / sample copyright.
    # Populate the author's actual submission kind without importing samples.
    for section in document.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            for child in list(header._element):
                header._element.remove(child)
            header.add_paragraph("Feature Article", style="Header")
    citation_proof = _citation_typography(document, reference, software=template.name == "CS_Mag_7_22.docx")
    document.save(output)
    checked = Document(output)
    if checked.paragraphs[0].style.name != "Title_document" or not any(
            s._sectPr.find(qn("w:cols")).get(qn("w:num")) == "2" for s in checked.sections):
        raise ValueError("Native magazine template structure was not retained")
    return {"style_mapping": "Publisher Title_document, Authors, Abstract, Para, numbered heading styles and Bib_entry",
            "layout": "Publisher one-column front matter and two-column body; tables span the printable width",
            "adaptations": ["Removed publisher sample artwork, editor placeholders and sample text",
                            "Feature Article running header; requested native page and line numbering",
                            "Left-aligned section headings stay with the following paragraph",
                            "Table captions share the table's full-width section; first header rows repeat with bold labels",
                            "Following table captions retain exact content above the table; section boundaries do not modify prose paragraphs",
                            "Abstract text aligned left to avoid distributed final-line spacing",
                            "Oversized body images shrink proportionally to the publisher column width and printable page height",
                            "Publisher-specific native citation typography and numbered-period reference labels"],
            "citation_typography": citation_proof,
            "font_policy": "Official font names retained; Microsoft Word substitutes locally unavailable fonts. Exact proprietary font fidelity is not attested."}


def word_pdf(docx: Path, output: Path) -> dict:
    """Render the exact exported Word bytes, never a generic PDF substitute."""
    if os.name != "nt":
        raise ValueError("Official-template PDF export currently requires Windows with Microsoft Word installed; the official Word file remains available")
    ensure_unlinked(docx)
    script = ROOT / "render-word.ps1"
    shell = Path(os.environ.get("SystemRoot", r"C:\Windows")) / "System32/WindowsPowerShell/v1.0/powershell.exe"
    ensure_unlinked(output)
    if output.exists():
        raise ValueError("PDF export refuses an existing output file")
    source_sha = digest_file(docx)
    # Word's COM path handling is more restrictive than the workflow's native
    # artifact paths. Render an exact short-path copy, then retain those bytes.
    with TemporaryDirectory(prefix="paperfactory-word-") as temporary:
        staged_docx, staged_pdf = Path(temporary) / "source.docx", Path(temporary) / "rendered.pdf"
        shutil.copyfile(docx, staged_docx)
        if digest_file(staged_docx) != source_sha:
            raise ValueError("The short-path Word rendering source differs from the native export")
        command = [str(shell), "-NoLogo", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(script.resolve()),
                   "-SourceDocx", str(staged_docx), "-OutputPdf", str(staged_pdf)]
        try:
            result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180,
                                    creationflags=subprocess.CREATE_NO_WINDOW)
        except subprocess.TimeoutExpired as exc:
            output.unlink(missing_ok=True)
            raise ValueError("Hidden Microsoft Word PDF export timed out after 180 seconds") from exc
        if result.returncode or not staged_pdf.is_file():
            output.unlink(missing_ok=True)
            raise ValueError("Official Word-template PDF rendering failed; Microsoft Word must be installed: " + result.stderr[:1500])
        if digest_file(staged_docx) != source_sha:
            raise ValueError("The short-path Word document changed during its read-only rendering")
        shutil.copyfile(staged_pdf, output)
    if digest_file(docx) != source_sha:
        output.unlink(missing_ok=True)
        raise ValueError("The official Word document changed during its read-only PDF rendering")
    return {"engine": "Microsoft Word COM ExportAsFixedFormat", "command": command,
            "source_docx_sha256": source_sha, "renderer_script_sha256": digest_file(script),
            "source_docx_path": str(docx.resolve()), "exact_short_path_copy": True,
            "read_only": True, "visible": False, "macros_disabled": True,
            "diagnostics": result.stderr, "word": json.loads(result.stdout.lstrip("\ufeff").strip())}


def measure_word(docx: Path) -> dict:
    """Measure a short exact copy with the same hidden, read-only Word process.

    This opens neither the original nor an interactive Word window and never
    saves a document. It uses the renderer's native-profile isolation fix.
    """
    if os.name != "nt":
        raise ValueError("Definitive magazine word counts require Windows with Microsoft Word installed")
    ensure_unlinked(docx)
    script = ROOT / "render-word.ps1"
    shell = Path(os.environ.get("SystemRoot", r"C:\Windows")) / "System32/WindowsPowerShell/v1.0/powershell.exe"
    source_sha = digest_file(docx)
    with TemporaryDirectory(prefix="paperfactory-wordmetrics-") as temporary:
        staged = Path(temporary) / "source.docx"
        shutil.copyfile(docx, staged)
        command = [str(shell), "-NoLogo", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(script.resolve()),
                   "-SourceDocx", str(staged), "-MetricsOnly"]
        try:
            result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180,
                                    creationflags=subprocess.CREATE_NO_WINDOW)
        except subprocess.TimeoutExpired as exc:
            raise ValueError("Hidden read-only Microsoft Word statistics timed out after 180 seconds") from exc
        if result.returncode:
            raise ValueError("Native Word statistics failed: " + result.stderr[:1500])
        if digest_file(staged) != source_sha or digest_file(docx) != source_sha:
            raise ValueError("The official Word source or its exact copy changed during read-only statistics")
        metrics = json.loads(result.stdout.lstrip("\ufeff").strip())
        if metrics.get("read_only") is not True:
            raise ValueError("Native Word statistics did not open the exact manuscript read-only")
    return {"engine": "Microsoft Word COM ComputeStatistics", "source_docx_sha256": source_sha,
            "source_unchanged": True, "visible": False, "macros_disabled": True,
            "renderer_script_sha256": digest_file(script), "metrics": metrics}
